from openpyxl import load_workbook

from excel_template_factory import ExcelTemplateFactory, create_template_workbook


def simple_spec() -> dict:
    return {
        "slug": "test-planner",
        "title": "Test planner",
        "subtitle": "A compact workbook",
        "palette": {
            "background": "#FAF7F2",
            "ink": "#263A35",
            "accent": "#477A67",
            "secondary": "#7A9B89",
            "highlight": "#D0A542",
        },
        "tabs": ["Start here", "Dashboard", "Items", "Settings"],
        "screen_contracts": {
            "Start here": {"template": "guide"},
            "Dashboard": {"template": "dashboard"},
            "Items": {"template": "table"},
            "Settings": {"template": "settings"},
        },
    }


def test_factory_creates_replicable_native_grid(tmp_path) -> None:
    book = create_template_workbook(simple_spec(), {
        "Items": {
            "headers": ["Item", "Status", "Value"],
            "rows": [["Example", "Open", 10]],
            "editable_columns": [0],
        }
    })
    path = tmp_path / "template.xlsx"
    book.save(path)

    loaded = load_workbook(path)
    assert loaded.sheetnames == ["Start here", "Dashboard", "Items", "Settings"]
    assert loaded["Items"]["B6"].value == "Item"
    assert loaded["Items"]["B7"].value == "Example"
    assert loaded["Items"]["B7"].alignment.indent == 1
    assert loaded["Dashboard"]["B32"].value is None
    assert loaded["Settings"]["B6"].value == "Setting"
    loaded.close()


def test_navigation_uses_native_excel_sheet_tabs() -> None:
    book = create_template_workbook(simple_spec())
    factory = ExcelTemplateFactory(book, simple_spec())
    assert factory.library["grid"]["first_content_column"] == 2
    assert factory.library["grid"]["last_content_column"] == 13
    assert book.sheetnames == ["Start here", "Dashboard", "Items", "Settings"]
    assert all(book["Dashboard"].cell(32, column).value is None for column in range(2, 6))


def test_factory_adds_declared_native_excel_chart() -> None:
    spec = simple_spec()
    spec["screen_contracts"]["Dashboard"] = {
        "template": "dashboard",
        "components": ["header", "summary_row", "table", "chart", "navigation"],
        "chart": {
            "type": "column",
            "title": "Values by item",
            "anchor": "H9",
            "categories": {"min_col": 2, "min_row": 10, "max_col": 2, "max_row": 11},
            "values": {"min_col": 4, "min_row": 9, "max_col": 4, "max_row": 11},
        },
    }
    book = create_template_workbook(spec, {
        "Dashboard": {
            "headers": ["Item", "Status", "Value"],
            "rows": [["A", "Open", 10], ["B", "Done", 20]],
        }
    })
    assert len(book["Dashboard"]._charts) == 1
    assert book["Dashboard"]._charts[0].title.tx.rich.p[0].r[0].t == "Values by item"
