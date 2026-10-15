from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.table import Table

from excel_template_factory import create_template_workbook
from model_quality import validate_model_workbook


def test_model_quality_rejects_generic_headers_percent_counts_and_internal_totals() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Guests"
    sheet.append(["Columna1", "Invitados"])
    sheet.append(["TOTAL", 2])
    sheet["B2"].number_format = "0%"
    sheet.add_table(Table(displayName="tblGuests", ref="A1:B2"))

    report = validate_model_workbook(workbook, ["Guests"])

    assert report["passed"] is False
    assert any("encabezados genéricos" in error for error in report["errors"])
    assert any("totales dentro" in error for error in report["errors"])
    assert any("no puede usar porcentaje" in error for error in report["errors"])


def test_model_quality_rejects_missing_sheet_formula_and_accepts_semantic_formats() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["Categoría", "Presupuesto", "Progreso"])
    sheet.append(["Casa", 100, 0.5])
    sheet["B2"].number_format = "#,##0.00"
    sheet["C2"].number_format = "0%"
    sheet["D2"] = "=Missing!A1"
    sheet.add_table(Table(displayName="tblBudget", ref="A1:C2"))

    report = validate_model_workbook(workbook, ["Budget"])

    assert report["passed"] is False
    assert any("hoja inexistente" in error for error in report["errors"])
    assert not any("Presupuesto" in error for error in report["errors"])
    assert not any("Progreso" in error for error in report["errors"])


def test_factory_declares_native_table_and_semantic_number_formats() -> None:
    spec = {
        "slug": "semantic-test",
        "palette": {
            "background": "#FAF7F2",
            "ink": "#263A35",
            "accent": "#477A67",
            "secondary": "#7A9B89",
            "highlight": "#D0A542",
        },
        "tabs": ["Items"],
        "screen_contracts": {"Items": {"template": "table"}},
    }
    workbook = create_template_workbook(
        spec,
        {"Items": {"headers": ["Invitados", "Progreso", "Presupuesto"], "rows": [[2, 0.5, 100]]}},
    )
    sheet = workbook["Items"]

    assert len(sheet.tables) == 1
    assert sheet["B7"].number_format == "0"
    assert sheet["C7"].number_format == "0%"
    assert sheet["D7"].number_format == "#,##0.00"


def test_model_quality_rejects_total_without_fit_to_cell() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet["A1"] = "Area"
    sheet["A2"] = "TOTAL"
    sheet["B2"] = "=SUM(B3:B4)"
    sheet.add_table(Table(displayName="tblBudget", ref="A1:B1"))
    contract = {
        "contract_id": "fit-v1", "title": "Fit", "purpose": "Test total formatting.",
        "sheets": ["Budget"], "tables": [{
            "sheet": "Budget", "ref": "A1:B1", "headers": [{"name": "Area"}, {"name": "Amount"}],
            "total_ref": "A2:B2",
        }],
        "formula_rules": [], "relationships": [], "cases": [], "failure_criteria": [],
    }
    from model_quality import validate_model_contract
    report = validate_model_contract(workbook, contract)
    assert any("totales sin ajuste" in error for error in report["errors"])
    sheet["B2"].alignment = Alignment(shrinkToFit=True)
    report = validate_model_contract(workbook, contract)
    assert not any("totales sin ajuste" in error for error in report["errors"])
