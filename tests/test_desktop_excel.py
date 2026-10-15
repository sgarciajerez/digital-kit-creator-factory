"""Opt-in tests against the installed Excel calculation engine."""
import os
import pytest
from openpyxl import Workbook, load_workbook
from model_quality import run_functional_contract
from workbook_views import apply_native_charts, apply_workbook_views

pytestmark = pytest.mark.skipif(os.getenv("RUN_EXCEL_TESTS") != "1", reason="Set RUN_EXCEL_TESTS=1 to use desktop Excel")


def test_cases_are_independent_and_excel_errors_are_detected(tmp_path):
    path = tmp_path / "functional.xlsx"
    book = Workbook()
    book.active.title = "Data"
    book.active["A1"] = 1
    book.active["B1"] = "=A1*2"
    book.active["C1"] = "=1/0"
    book.save(path)
    def check(address, value):
        return {"sheet": "Data", "cell": address, "operator": "equals", "value": value}
    report = run_functional_contract(path, {"cases": [
        {"name": "change", "inputs": [{"sheet": "Data", "cell": "A1", "value": 10}], "assertions": [check("B1", 20)]},
        {"name": "fresh", "inputs": [], "assertions": [check("B1", 2)]},
    ]})
    assert report["passed"], report
    error = run_functional_contract(path, {"cases": [{"name": "error", "inputs": [], "assertions": [{"sheet": "Data", "cell": "C1", "operator": "not_error"}]}]})
    assert not error["passed"]
    assert error["status"] == "executed", error


def test_declared_native_chart_and_frozen_header(tmp_path):
    path = tmp_path / "views.xlsx"
    book = Workbook()
    book.active.title = "Data"
    for row in [["Category", "Amount"], ["One", 1], ["Two", 2]]:
        book.active.append(row)
    book.save(path)
    spec = {"freeze_panes": {"Data": "B2"}, "screen_contracts": {"Data": {"components": ["chart"], "chart": {
        "type": "column", "title": "Example", "values": {"min_row": 1, "max_row": 3, "min_col": 2},
        "categories": {"min_row": 2, "max_row": 3, "min_col": 1}}}}}
    apply_native_charts(path, spec)
    apply_workbook_views(path, spec)
    result = load_workbook(path)
    try:
        assert result["Data"].freeze_panes == "B2"
        assert len(result["Data"]._charts) == 1
        assert result["Data"]["B3"].value == 2
    finally:
        result.close()
