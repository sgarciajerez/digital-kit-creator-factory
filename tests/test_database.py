from pathlib import Path
from datetime import date
import sqlite3
from types import SimpleNamespace

import pytest
from pydantic import BaseModel, ValidationError

from ai_provider import (
    AIRequest,
    MockProvider,
    ProviderConfigurationError,
    build_provider,
    validate_structured_output,
)
from budget import BudgetExceededError, BudgetManager
from batch_generation import generate_small_batch
from analytics import calculate_listing_report, create_experiment, export_listing_report, import_metrics_csv
from autopilot import register_idea_candidates, run_autopilot, select_publishable_opportunities
from research import (
    DEFAULT_WEIGHTS,
    ScoreComponentInput,
    import_evidence_csv,
    list_evidence,
    list_score_components,
    score_opportunity,
)
from product_generation import (
    DEFAULT_BRIEF,
    ProductBrief,
    create_first_product,
    generate_budget_workbook,
    validate_workbook,
)
from listing_package import ListingDraft, create_listing_package
from etsy_draft import build_etsy_draft_request, write_etsy_dry_run

from factory_core import (
    create_opportunity,
    create_product,
    delete_opportunity,
    delete_product,
    export_table_csv,
    export_table_json,
    get_opportunity,
    get_product,
    initialize_database,
    list_expenses,
    list_opportunities,
    list_products,
    load_settings,
    record_audit_event,
    record_expense,
    record_run,
    scalar,
    update_opportunity,
    update_product,
)


def test_database_initializes_and_records_run(tmp_path: Path) -> None:
    path = tmp_path / "factory.sqlite3"
    initialize_database(path)
    run_id = record_run(path, source="test")
    assert run_id == 1
    assert scalar(path, "SELECT COUNT(*) FROM app_runs") == 1
    assert scalar(path, "SELECT COUNT(*) FROM audit_events") == 1


def test_database_creation_is_idempotent(tmp_path: Path) -> None:
    path = tmp_path / "factory.sqlite3"
    initialize_database(path)
    initialize_database(path)
    assert scalar(path, "SELECT COUNT(*) FROM sqlite_master WHERE type = 'table'") >= 5


def test_load_settings_reads_environment_without_exposing_api_key(
    tmp_path: Path, monkeypatch
) -> None:
    database_path = tmp_path / "configured.sqlite3"
    monkeypatch.setenv("DATABASE_PATH", str(database_path))
    monkeypatch.setenv("MONTHLY_BUDGET_EUR", "12.50")
    monkeypatch.setenv("AUTONOMY_MODE", "2")
    monkeypatch.setenv("AI_PROVIDER", "mock")
    monkeypatch.setenv("AI_MODEL", "test-model")
    monkeypatch.delenv("AI_API_KEY", raising=False)

    settings = load_settings()

    assert settings.database_path == database_path
    assert settings.monthly_budget_eur == 12.5
    assert settings.autonomy_mode == 2
    assert settings.ai_provider == "mock"
    assert settings.ai_model == "test-model"
    assert settings.ai_api_key_configured is False


def test_load_settings_rejects_invalid_autonomy_mode(monkeypatch) -> None:
    monkeypatch.setenv("AUTONOMY_MODE", "9")

    try:
        load_settings()
    except ValueError as error:
        assert "AUTONOMY_MODE" in str(error)
    else:
        raise AssertionError("Se esperaba un ValueError para AUTONOMY_MODE inválido")


def test_database_migrates_schema_and_creates_indexes(tmp_path: Path) -> None:
    path = tmp_path / "factory.sqlite3"

    initialize_database(path)

    assert scalar(path, "PRAGMA user_version") == 3
    tables = scalar(
        path,
        "SELECT COUNT(*) FROM sqlite_master WHERE type = 'table' AND name IN "
        "('research_runs', 'sources', 'evidence', 'score_components', 'listings', "
        "'metrics_daily', 'experiments', 'approvals')",
    )
    assert tables == 8
    assert scalar(path, "SELECT COUNT(*) FROM sqlite_master WHERE type = 'table' AND name = 'opportunity_evidence'") == 1
    assert scalar(
        path,
        "SELECT COUNT(*) FROM sqlite_master WHERE type = 'index' "
        "AND name = 'idx_opportunities_decision'",
    ) == 1


def test_database_v2_migrates_opportunity_evidence_table(tmp_path: Path) -> None:
    path = tmp_path / "factory-v2.sqlite3"
    initialize_database(path)
    with sqlite3.connect(path) as connection:
        connection.execute("DROP TABLE opportunity_evidence")
        connection.execute("PRAGMA user_version = 2")

    initialize_database(path)

    assert scalar(path, "PRAGMA user_version") == 3
    assert scalar(path, "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name='opportunity_evidence'") == 1


def test_opportunity_and_product_crud_is_audited(tmp_path: Path) -> None:
    path = tmp_path / "factory.sqlite3"
    opportunity_id = create_opportunity(
        path,
        "Plantilla de presupuesto",
        category="finanzas",
        score=82,
        confidence=74,
    )

    opportunity = get_opportunity(path, opportunity_id)
    assert opportunity is not None
    assert opportunity["decision"] == "review"
    assert list_opportunities(path, decision="review")[0]["id"] == opportunity_id

    assert update_opportunity(path, opportunity_id, decision="accept", decision_reason="Prueba")
    assert get_opportunity(path, opportunity_id)["decision"] == "accept"

    product_id = create_product(
        path,
        "Plantilla de presupuesto v1",
        "spreadsheet",
        opportunity_id=opportunity_id,
        quality_score=90,
    )
    assert get_product(path, product_id)["opportunity_id"] == opportunity_id
    assert list_products(path, status="draft")[0]["id"] == product_id
    assert update_product(path, product_id, status="approved", quality_report="OK")
    assert get_product(path, product_id)["status"] == "approved"

    assert delete_product(path, product_id)
    assert delete_opportunity(path, opportunity_id)
    assert scalar(path, "SELECT COUNT(*) FROM audit_events") >= 5


def test_expenses_and_exports_are_persistent(tmp_path: Path) -> None:
    path = tmp_path / "factory.sqlite3"
    record_expense(path, 1.25, "test", provider="mock", notes="sin secreto")
    record_expense(path, 2.0, "test", currency="USD")
    record_audit_event(path, "test_event", actor="test", details_redacted="safe")

    assert len(list_expenses(path, currency="EUR")) == 1
    csv_path = export_table_csv(path, "expenses", tmp_path / "exports" / "expenses.csv")
    json_path = export_table_json(path, "audit_events", tmp_path / "exports" / "audit.json")
    assert "amount" in csv_path.read_text(encoding="utf-8")
    assert "test_event" in json_path.read_text(encoding="utf-8")


def test_existing_legacy_schema_is_upgraded(tmp_path: Path) -> None:
    path = tmp_path / "legacy.sqlite3"
    with sqlite3.connect(path) as connection:
        connection.executescript(
            "CREATE TABLE opportunities ("
            "id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, "
            "decision TEXT NOT NULL DEFAULT 'review', created_at TEXT NOT NULL, "
            "updated_at TEXT NOT NULL);"
        )

    initialize_database(path)

    with sqlite3.connect(path) as connection:
        columns = {
            row[1]
            for row in connection.execute("PRAGMA table_info(opportunities)")
        }
    assert {"subcategory", "proposed_solution", "market_research_json", "next_review_at"} <= columns
    assert scalar(path, "PRAGMA user_version") == 3


def test_mock_provider_validates_structured_output_and_records_cost(tmp_path: Path) -> None:
    class MockOutput(BaseModel):
        status: str
        task: str
        prompt_received: bool

    path = tmp_path / "factory.sqlite3"
    budget = BudgetManager(path, monthly_limit_eur=1.0, task_limit_eur=0.5)
    provider = MockProvider(budget)
    request = AIRequest(
        task="classify",
        prompt="Clasifica esta idea",
        estimated_cost_eur=0.25,
        scope="task",
        scope_id="classification-1",
    )

    response = provider.complete(request)
    structured = validate_structured_output(response, MockOutput)

    assert structured.status == "simulated"
    assert response.actual_cost_eur == 0
    assert list_expenses(path, currency="EUR")[0]["category"] == "ai"
    assert budget.snapshot().spent_eur == 0


def test_budget_blocks_calls_over_monthly_limit(tmp_path: Path) -> None:
    path = tmp_path / "factory.sqlite3"
    provider = MockProvider(BudgetManager(path, monthly_limit_eur=0.10))
    request = AIRequest(task="blocked", prompt="No debe ejecutarse", estimated_cost_eur=0.11)

    with pytest.raises(BudgetExceededError):
        provider.complete(request)


def test_provider_factory_rejects_unimplemented_external_provider(tmp_path: Path) -> None:
    path = tmp_path / "factory.sqlite3"
    settings = SimpleNamespace(ai_provider="external", ai_model="", ai_api_key_configured=True)

    with pytest.raises(ProviderConfigurationError):
        build_provider(settings, BudgetManager(path, monthly_limit_eur=1.0))

    with pytest.raises(ValidationError):
        AIRequest(task="invalid", prompt="x", max_output_tokens=0)


def test_manual_csv_evidence_and_explainable_score(tmp_path: Path) -> None:
    path = tmp_path / "factory.sqlite3"
    csv_path = tmp_path / "evidence.csv"
    csv_path.write_text(
        "observed_at,query,market,language,metric_name,metric_value,unit,evidence_type,"
        "confidence,limitations,raw_reference\n"
        "2026-10-13,manual interview,ES,es,requests,3,count,observed,80,small sample,REF-1\n",
        encoding="utf-8",
    )

    result = import_evidence_csv(path, csv_path, "FICTITIOUS_EXAMPLE_MANUAL")
    assert result.evidence_count == 1
    assert list_evidence(path, result.research_run_id)[0]["evidence_type"] == "observed"

    opportunity_id = create_opportunity(path, "Idea ficticia para pruebas")
    components = [
        ScoreComponentInput(criterion=criterion, weight=weight, score=score, rationale="Ejemplo ficticio")
        for (criterion, weight), score in zip(
            DEFAULT_WEIGHTS.items(), (2, 4, 3, 4, 3, 4, 4)
        )
    ]
    total = score_opportunity(path, opportunity_id, components, confidence=60)
    scored = get_opportunity(path, opportunity_id)

    assert total == 64
    assert scored["decision"] == "review"
    assert scored["confidence"] == 60
    assert len(list_score_components(path, opportunity_id)) == 7


def test_first_product_generates_and_validates_xlsx(tmp_path: Path) -> None:
    path = tmp_path / "factory.sqlite3"
    output_directory = tmp_path / "output" / "products"

    product_id, report, output_path = create_first_product(path, output_directory, DEFAULT_BRIEF)

    assert product_id == 1
    assert report.passed is True
    assert report.score == 100
    assert output_path.exists()
    assert validate_workbook(output_path).passed is True
    product = get_product(path, product_id)
    assert product["status"] == "draft"
    assert product["quality_score"] == 100


def test_family_product_includes_shared_expense_workbook(tmp_path: Path) -> None:
    brief = ProductBrief(
        family="family_budget_planner",
        name="Planificador familiar",
        audience="Hogares que comparten gastos",
        problem="Necesitan saber quién paga y cómo repartir gastos",
        categories=["Vivienda", "Alimentación", "Otros"],
    )
    path = tmp_path / "family.xlsx"
    generate_budget_workbook(path, brief)
    report = validate_workbook(path, expected_family=brief.family)
    assert report.passed
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True)
    assert {"Configuración", "Gastos compartidos"} <= set(workbook.sheetnames)
    workbook.close()


def test_small_batch_is_measured_and_sent_to_human_review(tmp_path: Path) -> None:
    path = tmp_path / "factory.sqlite3"

    report = generate_small_batch(path, tmp_path / "output" / "products")

    assert len(report.product_ids) == 5
    assert len(report.review_request_ids) == 5
    assert report.mean_quality_score == 100
    assert report.report_path.exists()
    assert scalar(path, "SELECT COUNT(*) FROM approvals WHERE decision = 'pending'") == 5
    assert scalar(path, "SELECT COUNT(*) FROM products WHERE status = 'draft'") == 5


def test_listing_package_is_local_validated_and_pending_approval(tmp_path: Path) -> None:
    path = tmp_path / "factory.sqlite3"
    product_id, _, _ = create_first_product(path, tmp_path / "output" / "products")
    draft = ListingDraft(
        title="Planificador de presupuesto personal",
        description="Plantilla local para organizar presupuesto y gasto real por categorías.",
        tags=["presupuesto", "planificación", "gastos"],
        price=4.99,
        policy_reviewed=True,
        rights_confirmed=True,
    )

    listing_id, validation, package_directory = create_listing_package(
        path,
        tmp_path / "output" / "listings",
        product_id,
        "manual-upload",
        draft,
    )

    assert listing_id == 1
    assert validation.passed is True
    assert (package_directory / "listing.json").exists()
    assert (package_directory / "description.txt").exists()
    assert scalar(path, "SELECT COUNT(*) FROM listings WHERE status = 'pending_approval'") == 1
    assert scalar(path, "SELECT COUNT(*) FROM approvals WHERE target_type = 'listing'") == 1


def test_metrics_import_report_and_experiment(tmp_path: Path) -> None:
    path = tmp_path / "factory.sqlite3"
    product_id, _, _ = create_first_product(path, tmp_path / "output" / "products")
    draft = ListingDraft(
        title="Planificador de presupuesto personal",
        description="Plantilla local para organizar presupuesto y gasto real por categorías.",
        tags=["presupuesto"],
        price=4.99,
        policy_reviewed=True,
        rights_confirmed=True,
    )
    listing_id, _, _ = create_listing_package(
        path, tmp_path / "output" / "listings", product_id, "manual-upload", draft
    )
    metrics_path = tmp_path / "metrics.csv"
    metrics_path.write_text(
        "date,impressions,visits,favorites,orders,units_sold,gross_revenue,refunds,fees,"
        "ad_spend,source,data_quality_notes\n"
        "2026-10-13,200,100,12,10,10,100,5,10,3,manual,FICTIONAL_EXAMPLE_DATA\n",
        encoding="utf-8",
    )

    result = import_metrics_csv(path, listing_id, metrics_path)
    report = calculate_listing_report(path, listing_id)
    report_path = export_listing_report(path, listing_id, tmp_path / "reports" / "metrics.json")
    experiment_id = create_experiment(path, "Precio inicial", "Un precio bajo mejora conversión", "conversion_rate", 1.0)

    assert result.rows_imported == 1
    assert report.conversion_rate == 10
    assert report.contribution_margin_estimated == 82
    assert any("FICTIONAL_EXAMPLE" in warning for warning in report.warnings)
    assert report_path.exists()
    assert experiment_id == 1


def test_etsy_draft_is_prepared_without_network_or_activation(tmp_path: Path) -> None:
    draft = ListingDraft(
        title="Planificador de presupuesto personal",
        description="Plantilla local para organizar presupuesto y gasto real por categorías.",
        tags=["presupuesto"],
        price=4.99,
        policy_reviewed=True,
        rights_confirmed=True,
    )

    request = build_etsy_draft_request(draft, shop_id=123, taxonomy_id=456, file_name="product.xlsx")
    path = write_etsy_dry_run(request, tmp_path / "etsy_draft_request.json")

    assert request.form_data["type"] == "download"
    assert request.activation_requested is False
    assert path.exists()


def test_autopilot_registers_hypotheses_without_claiming_demand(tmp_path: Path) -> None:
    path = tmp_path / "factory.sqlite3"
    ids = register_idea_candidates(path)
    assert len(ids) == 5
    assert select_publishable_opportunities(path) == ()
    assert all(row["decision"] == "review" for row in list_opportunities(path))


def test_autopilot_generates_only_accepted_local_drafts(tmp_path: Path) -> None:
    path = tmp_path / "factory.sqlite3"
    accepted_id = create_opportunity(
        path,
        "Plantilla de presupuesto para opositores",
        category="estudio",
        subcategory="oposiciones",
        target_audience="Personas que preparan una oposición",
        problem="Necesitan separar gastos de estudio y vida diaria",
        decision="accept",
        score=88,
        confidence=80,
    )
    evidence_path = tmp_path / "accepted_market_evidence.csv"
    evidence_path.write_text(
        "observed_at,query,market,language,metric_name,metric_value,unit,evidence_type,"
        "confidence,limitations,raw_reference\n"
        f"{date.today().isoformat()},budget opposition,US,en,monthly searches,1200,count,observed,90,Search estimate varies,https://keywords.example/searches\n"
        f"{date.today().isoformat()},budget opposition,US,en,active listings,450,count,observed,85,Search results are supply,https://etsy.example/search\n"
        f"{date.today().isoformat()},budget opposition,US,en,buyer interviews,4,count,observed,70,Small exploratory sample,https://interviews.example/notes\n",
        encoding="utf-8",
    )
    evidence_import = import_evidence_csv(path, evidence_path, "selection_test_sources")
    evidence_ids = tuple(row["id"] for row in list_evidence(path, evidence_import.research_run_id))
    score_opportunity(
        path,
        accepted_id,
        [ScoreComponentInput("evidencia de demanda", 100, 5, "Test evidence-linked decision", evidence_ids)],
        confidence=80,
    )
    create_opportunity(path, "Candidato sin aceptación", decision="review", score=99, confidence=99)
    report = run_autopilot(
        path,
        tmp_path / "output" / "products",
        tmp_path / "output" / "listings",
        max_products=2,
    )
    assert report.selected_opportunity_ids == (accepted_id,)
    assert len(report.product_ids) == 1
    assert len(report.listing_ids) == 1
    assert scalar(path, "SELECT status FROM products WHERE id = ?", report.product_ids) == "draft"
    assert scalar(path, "SELECT status FROM listings WHERE id = ?", report.listing_ids) == "needs_review"
    assert scalar(path, "SELECT COUNT(*) FROM approvals WHERE target_type = 'product'") == 1
    assert scalar(path, "SELECT COUNT(*) FROM approvals WHERE target_type = 'listing'") == 1
    assert scalar(path, "SELECT COUNT(*) FROM opportunities WHERE decision = 'review'") == 1

