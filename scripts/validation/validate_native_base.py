from __future__ import annotations

"""Validate the native XLSX master used by new product designs."""

import argparse
import json
import warnings
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from model_quality import validate_model_contract, validate_model_workbook, validate_model_behavior_markdown
from model_contracts import family_base_contract


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DIR = ROOT / "templates" / "excel-first-base"
DEFAULT_WORKBOOK = DEFAULT_DIR / "master_workbook.xlsx"
DEFAULT_CONTRACT = DEFAULT_DIR / "native_base_contract.json"
DEFAULT_REPORT = DEFAULT_DIR / "native_base_report.json"


def validate(workbook_path: Path, contract_path: Path) -> dict:
    contract = json.loads(contract_path.read_text(encoding="utf-8"))
    expected_sheets = list(contract["sheet_geometry"])
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message=r"The image .*\.svg will be removed")
        warnings.filterwarnings("ignore", message=r"Data Validation extension is not supported")
        workbook = load_workbook(workbook_path, data_only=False)
    try:
        formulas = [
            cell.value
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
        row7_borders = [
            f"{sheet.title}!{cell.coordinate}"
            for sheet in workbook.worksheets
            for cell in sheet[7]
            if getattr(cell.border.bottom, "style", None) or getattr(cell.border.bottom, "color", None)
        ]
        generic_headers = []
        total_rows_inside_tables = []
        for sheet in workbook.worksheets:
            for table in sheet.tables.values():
                min_col, min_row, max_col, max_row = range_boundaries(table.ref)
                headers = [sheet.cell(min_row, col).value for col in range(min_col, max_col + 1)]
                if any(isinstance(header, str) and header.lower().startswith("columna") for header in headers):
                    generic_headers.append(f"{sheet.title}!{table.name}")
                for row in range(min_row + 1, max_row + 1):
                    values = [sheet.cell(row, col).value for col in range(min_col, min(max_col, min_col + 1) + 1)]
                    if any(isinstance(value, str) and value.strip().upper() == "TOTAL" for value in values):
                        total_rows_inside_tables.append(f"{sheet.title}!{table.name}!{row}")
        checks = {
            "sheet_order_matches_contract": workbook.sheetnames == expected_sheets,
            "freeze_panes_disabled": all(sheet.freeze_panes is None for sheet in workbook.worksheets),
            "row_7_has_no_bottom_borders": not row7_borders,
            "formula_contract": bool(formulas) and all("#REF!" not in formula for formula in formulas),
            "native_tables_present": sum(len(sheet.tables) for sheet in workbook.worksheets) >= 8,
            "table_headers_are_semantic": not generic_headers,
            "table_totals_are_outside_data_ranges": not total_rows_inside_tables,
            "native_charts_present": sum(len(sheet._charts) for sheet in workbook.worksheets) >= 2,
        }
        details = {
            "formula_count": len(formulas),
            "table_count": sum(len(sheet.tables) for sheet in workbook.worksheets),
            "chart_count": sum(len(sheet._charts) for sheet in workbook.worksheets),
            "row_7_borders": row7_borders,
            "generic_table_headers": generic_headers,
            "total_rows_inside_tables": total_rows_inside_tables,
        }
        model_report = validate_model_workbook(workbook, expected_sheets)
        model_contract = family_base_contract()
        behavior_path = DEFAULT_DIR / "model_behavior.md"
        contract_report = validate_model_contract(workbook, model_contract, behavior_path)
        model_report["model_contract"] = contract_report
        model_report["passed"] = model_report["passed"] and contract_report["passed"]
        model_report["errors"].extend(contract_report.get("errors", []))
        checks["model_quality_contract"] = model_report["passed"]
        details["model_quality_errors"] = model_report["errors"]
    finally:
        workbook.close()

    with ZipFile(workbook_path) as archive:
        package = {name: archive.read(name) for name in archive.namelist() if name.endswith((".xml", ".rels", ".svg"))}
    package_text = b"".join(package.values())
    checks.update(
        {
            "svg_asset_embedded": any(name.endswith(".svg") for name in package),
            "icon_on_every_sheet": package_text.count(b"FamilyBudgetIcon") >= len(expected_sheets),
            "rounded_navigation_on_every_sheet": package_text.count(b"Nav_") >= len(expected_sheets),
        }
    )
    return {
        "status": "passed" if all(checks.values()) else "failed",
        "workbook": str(workbook_path.relative_to(ROOT)),
        "contract": str(contract_path.relative_to(ROOT)),
        "checks": checks,
        "details": details,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Valida la base maestra XLSX nativa")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--contract", type=Path, default=DEFAULT_CONTRACT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    workbook = args.workbook if args.workbook.is_absolute() else ROOT / args.workbook
    contract = args.contract if args.contract.is_absolute() else ROOT / args.contract
    report_path = args.report if args.report.is_absolute() else ROOT / args.report
    report = validate(workbook, contract)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
