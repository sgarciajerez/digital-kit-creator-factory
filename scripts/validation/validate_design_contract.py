from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "outputs" / "listings" / "product-7" / "autopilot_3_Planificador_de_presupuesto_familiar_0.1.0.xlsx"
SPEC_PATH = ROOT / "outputs" / "listings" / "product-7" / "design_spec.json"


def fail(errors: list[str], message: str) -> None:
    errors.append(message)


def main() -> int:
    spec = json.loads(SPEC_PATH.read_text(encoding="utf-8"))
    wb = load_workbook(BOOK, data_only=False)
    errors: list[str] = []
    required_sheets = spec["acceptance"]["required_sheets"]
    for name in required_sheets:
        if name not in wb.sheetnames:
            fail(errors, f"Falta la hoja: {name}")

    if errors:
        raise SystemExit("\n".join(errors))

    for name in required_sheets:
        design = spec["sheets"][name]
        ws = wb[name]
        if ws["A1"].value != design["title"]:
            fail(errors, f"{name}!A1 no coincide con el diseño")
        if ws["A3"].value != design["subtitle"]:
            fail(errors, f"{name}!A3 no coincide con el diseño")

    budget = wb["Presupuesto"]
    budget_spec = spec["sheets"]["Presupuesto"]
    for column, header in enumerate(budget_spec["headers"], 1):
        if budget.cell(5, column).value != header:
            fail(errors, f"Presupuesto!{budget.cell(5, column).coordinate} no coincide")
    for row_number, expected in enumerate(budget_spec["rows"], 6):
        actual = [budget.cell(row_number, column).value for column in range(1, 4)]
        if actual != expected:
            fail(errors, f"Presupuesto fila {row_number}: {actual!r} != {expected!r}")
    expected_formulas = {
        "D6": "=B6-C6",
        "B13": "=SUM(B6:B12)",
        "C13": "=SUM(C6:C12)",
        "D13": "=B13-C13",
    }
    for cell, formula in expected_formulas.items():
        if budget[cell].value != formula:
            fail(errors, f"{budget.title}!{cell} no contiene {formula}")
    if budget["B6"].fill.fgColor.rgb != "00" + spec["theme"]["pale_yellow"]:
        fail(errors, "La columna editable del presupuesto no usa el color definido")

    preview = wb["Vista previa"]
    for cell, formula in {"A6": "=Presupuesto!B13", "D6": "=Presupuesto!C13", "G6": "=Presupuesto!D13"}.items():
        if preview[cell].value != formula:
            fail(errors, f"Vista previa!{cell} no contiene {formula}")
    if len(preview._charts) != 0:
        fail(errors, "Vista previa no debe depender de un gráfico clásico; debe usar barras en celdas")
    if len(preview.conditional_formatting) == 0:
        fail(errors, "Vista previa debe contener barras visuales en las celdas de gasto")

    settings = wb["Configuración"]
    for row, (key, value) in enumerate(spec["sheets"]["Configuración"]["values"].items(), 6):
        if settings.cell(row, 1).value != key or settings.cell(row, 2).value != value:
            fail(errors, f"Configuración fila {row} no coincide")
    if "Miembros" not in wb.defined_names:
        fail(errors, "Falta el rango definido Miembros")

    shared = wb["Gastos compartidos"]
    for row_number, expected in enumerate(spec["sheets"]["Gastos compartidos"]["rows"], 6):
        actual_date = shared.cell(row_number, 1).value
        actual_date = actual_date.date().isoformat() if isinstance(actual_date, datetime) else str(actual_date)
        actual = [actual_date] + [shared.cell(row_number, column).value for column in range(2, 5)] + [shared.cell(row_number, 5).value]
        if actual != expected:
            fail(errors, f"Gastos compartidos fila {row_number}: {actual!r} != {expected!r}")
        if shared.cell(row_number, 6).value != f"=E{row_number}/2" or shared.cell(row_number, 7).value != f"=E{row_number}/2":
            fail(errors, f"Fórmulas de reparto incorrectas en fila {row_number}")

    planned_total = sum(row[1] for row in budget_spec["rows"])
    actual_total = sum(row[2] for row in budget_spec["rows"])
    if planned_total != spec["acceptance"]["budget_total"] or actual_total != spec["acceptance"]["actual_total"]:
        fail(errors, "Los totales del contrato no coinciden con sus datos")
    if planned_total - actual_total != spec["acceptance"]["difference_total"]:
        fail(errors, "La diferencia total del contrato no coincide")

    if errors:
        print("DESIGN_CONTRACT=FAILED")
        print("\n".join(f"- {error}" for error in errors))
        return 1
    print("DESIGN_CONTRACT=PASSED")
    print(f"Sheets={len(required_sheets)} Charts={len(preview._charts)} DataBars={len(preview.conditional_formatting)}")
    print(f"Totals={planned_total}/{actual_total}/{planned_total - actual_total}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
