from __future__ import annotations

"""Validate the reusable native Excel master directly."""

import argparse
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from design_architecture import load_design_system, validate_design_spec


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_DIR = ROOT / "templates" / "excel-first-base"


def validate_base(spec_path: Path, master_path: Path, workbook_path: Path, report_path: Path) -> dict:
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    errors = validate_design_spec(spec, load_design_system())
    if not master_path.is_file():
        errors.append("master_workbook.xlsx no existe")
    if workbook_path.resolve() != master_path.resolve():
        errors.append("La validacion de base debe usar master_workbook.xlsx como libro de referencia")
    try:
        workbook = load_workbook(master_path, read_only=False, data_only=False)
    except Exception as exc:
        workbook = None
        errors.append(f"No se puede abrir master_workbook.xlsx: {exc}")
    formula_count = 0
    validation_count = 0
    chart_count = 0
    table_count = 0
    no_freeze_panes = True
    explicit_headers = True
    if workbook is not None:
        for ws in workbook.worksheets:
            formula_count += sum(1 for row in ws.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("="))
            validation_count += len(ws.data_validations.dataValidation)
            chart_count += len(ws._charts)
            table_count += len(ws.tables)
            no_freeze_panes = no_freeze_panes and ws.freeze_panes is None
            for table in ws.tables.values():
                min_col, min_row, max_col, _ = range_boundaries(table.ref)
                headers = [ws.cell(min_row, col).value for col in range(min_col, max_col + 1)]
                explicit_headers = explicit_headers and bool(headers) and all(value not in (None, "") for value in headers)
        sheet_names = list(workbook.sheetnames)
        workbook.close()
    else:
        sheet_names = []
    checks = {
        "master_exists": master_path.is_file(),
        "master_is_workbook": bool(sheet_names),
        "no_freeze_panes": no_freeze_panes,
        "explicit_table_headers": explicit_headers,
        "formula_count": formula_count >= 12,
        "validation_count": validation_count >= 1,
        "chart_contract": chart_count >= 1,
        "native_objects_present": table_count >= 1,
    }
    errors.extend(f"fallo: {name}" for name, passed in checks.items() if not passed)
    report = {
        "schema_version": "3.0.0",
        "status": "passed" if not errors else "failed",
        "source_of_truth": "master_workbook.xlsx",
        "master_workbook": str(master_path),
        "spec": str(spec_path),
        "checks": checks,
        "workbook": {"sheets": sheet_names, "formula_count": formula_count, "validation_count": validation_count, "chart_count": chart_count, "table_count": table_count},
        "errors": errors,
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="Valida directamente master_workbook.xlsx")
    parser.add_argument("--spec", type=Path, default=DEFAULT_DIR / "design_spec.json")
    parser.add_argument("--master", type=Path, default=DEFAULT_DIR / "master_workbook.xlsx")
    parser.add_argument("--workbook", type=Path, help="Alias de --master por compatibilidad")
    parser.add_argument("--report", type=Path, default=DEFAULT_DIR / "design_contract_report.json")
    args = parser.parse_args()
    master = args.workbook or args.master
    report = validate_base(args.spec, master, master, args.report)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
