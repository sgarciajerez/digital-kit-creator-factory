from __future__ import annotations

"""Fast semantic audit between the declared native model and the XLSX."""

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from model_contracts import contract_for_product


ROOT = Path(__file__).resolve().parents[2]


def main() -> int:
    parser = argparse.ArgumentParser(description="Audita el contrato nativo contra el XLSX")
    parser.add_argument("--product", type=int, required=True)
    args = parser.parse_args()
    product = ROOT / "outputs" / "listings" / f"product-{args.product}"
    spec = json.loads((product / "design_spec.json").read_text(encoding="utf-8"))
    candidates = sorted((product / "app").glob("*.xlsx"))
    if not candidates:
        raise FileNotFoundError(f"No hay XLSX generado en {product / 'app'}")
    book_path = product / "app" / f"{spec['slug']}-excel-template.xlsx"
    contract = contract_for_product(args.product)
    wb = load_workbook(book_path, data_only=False)
    results: dict[str, dict] = {}
    errors: list[str] = []
    for table in contract.get("tables", []):
        sheet = table["sheet"]
        ref = table["ref"]
        item_errors: list[str] = []
        if sheet not in wb.sheetnames:
            item_errors.append("missing_sheet")
        else:
            ws = wb[sheet]
            min_col, min_row, max_col, max_row = range_boundaries(ref)
            actual_headers = [ws.cell(min_row, col).value for col in range(min_col, max_col + 1)]
            expected_headers = [column["name"] for column in table["headers"]]
            if actual_headers != expected_headers:
                item_errors.append(f"headers:{actual_headers!r}")
            matching_tables = [obj for obj in ws.tables.values() if obj.ref == ref]
            if table.get("required", True) and not matching_tables:
                item_errors.append(f"missing_table:{ref}")
            if max_row <= min_row:
                item_errors.append("empty_data_range")
        results[f"{sheet}!{ref}"] = {"status": "passed" if not item_errors else "failed", "errors": item_errors, "expected_headers": [column["name"] for column in table["headers"]]}
        errors.extend(f"{sheet}!{ref}: {item}" for item in item_errors)

    formula_results: list[dict] = []
    for rule in contract.get("formula_rules", []):
        sheet = rule["sheet"]
        cell = rule["cell"]
        actual = wb[sheet][cell].value if sheet in wb.sheetnames else None
        passed = isinstance(actual, str) and actual.startswith("=") and rule["contains"].casefold() in actual.casefold()
        formula_results.append({"sheet": sheet, "cell": cell, "passed": passed, "expected_contains": rule["contains"], "actual": actual})
        if not passed:
            errors.append(f"{sheet}!{cell}: formula_missing_or_incoherent")

    for ws in wb.worksheets:
        if ws.freeze_panes != spec.get("freeze_panes", {}).get(ws.title):
            errors.append(f"{ws.title}: freeze_panes_not_declared")

    report = {
        "product": args.product,
        "workbook": str(book_path),
        "source_of_truth": "master_workbook_and_native_model_contract",
        "status": "passed" if not errors else "failed",
        "tables": results,
        "formula_rules": formula_results,
        "errors": errors,
    }
    (product / "design_contract_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
