from __future__ import annotations

import hashlib
import json
import re
import textwrap
import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont, ImageOps


ROOT = Path(__file__).resolve().parents[2]
CANVAS = (2000, 1500)
SCREEN_SIZE = (1920, 1080)
INK = "#243747"
MUTED = "#667781"
WHITE = "#FFFFFF"
GRID = "#DDE6E5"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = (
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    )
    for candidate in candidates:
        path = Path(candidate)
        if path.is_file():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def palette_for(spec: dict[str, Any]) -> dict[str, str]:
    declared = spec.get("palette") or {}
    result = {key: str(declared.get(key, fallback)) for key, fallback in {
        "background": "#F3F7F8", "ink": INK, "accent": "#2D8F8B",
        "secondary": "#E9826B", "highlight": "#D2AA4B",
    }.items()}
    for key, value in tuple(result.items()):
        if not re.fullmatch(r"#[0-9A-Fa-f]{6}", value):
            result[key] = "#2D8F8B" if key == "accent" else (INK if key == "ink" else WHITE)
    return result


def marketing_identity(spec: dict[str, Any], package: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    """Resolve the visual identity used by marketing, including its product art."""
    declared = spec.get("visual_identity") or {}
    title = str(spec.get("title", "Digital planning kit")).strip()
    name = str(declared.get("name") or title).strip()
    motif = str(declared.get("motif") or f"A distinctive visual for {title}").strip()
    art_ref = str(
        (declared.get("marketing_artwork_file") or manifest.get("marketing_artwork_file") or manifest.get("ebook_cover_image"))
        or ""
    ).strip()
    art_path = package / art_ref if art_ref else None
    art_sha256 = sha256(art_path) if art_path and art_path.is_file() else None
    colors = palette_for(spec)
    layout_id = str(declared.get("marketing_layout", ""))
    art_direction = str(declared.get("marketing_art_direction", ""))
    signature = hashlib.sha256(
        "|".join((name.casefold(), motif.casefold(), layout_id.casefold(), art_direction.casefold(), art_sha256 or "", *(colors[key].casefold() for key in ("background", "ink", "accent", "secondary", "highlight")))).encode("utf-8")
    ).hexdigest()
    return {
        "name": name,
        "motif": motif,
        "layout_id": layout_id,
        "art_direction": art_direction,
        "palette": colors,
        "signature_sha256": signature,
        "artwork_file": art_ref if art_sha256 else None,
        "artwork_sha256": art_sha256,
    }


def draw_identity_mark(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], identity: dict[str, Any]) -> None:
    """Draw a compact product-specific identity plaque, not a generic kit label."""
    left, top, right, bottom = box
    palette = identity["palette"]
    draw.rounded_rectangle(box, radius=14, fill=palette["ink"])
    draw.rounded_rectangle((left + 8, top + 8, left + 20, bottom - 8), radius=6, fill=palette["accent"])
    label = str(identity["name"]).upper()
    draw.text((left + 34, top + 10), label[:32], font=font(20, True), fill=WHITE)
    draw_wrapped(draw, (left + 34, top + 38), str(identity["motif"]), width=43, font_obj=font(15), fill=palette["background"], spacing=2)


def draw_identity_signature(draw: ImageDraw.ImageDraw, identity: dict[str, Any]) -> None:
    """Repeat a discreet brand cue across all mockups in the same product set."""
    left, top, right, bottom = 1510, 52, 1920, 142
    palette = identity["palette"]
    draw.rounded_rectangle((left, top, right, bottom), radius=18, fill=palette["ink"])
    draw.rounded_rectangle((left + 9, top + 9, left + 21, bottom - 9), radius=6, fill=palette["accent"])
    draw.text((left + 34, top + 10), str(identity["name"]).upper()[:30], font=font(18, True), fill=WHITE)
    motif = str(identity["motif"])
    draw.text((left + 34, top + 44), motif[:39] + ("…" if len(motif) > 39 else ""), font=font(13), fill=palette["background"])


def draw_identity_artwork(canvas: Image.Image, box: tuple[int, int, int, int], identity: dict[str, Any], package: Path) -> None:
    """Use the product's own e-book art, with a designed motif fallback."""
    left, top, right, bottom = box
    palette = identity["palette"]
    draw = ImageDraw.Draw(canvas)
    draw.rounded_rectangle(box, radius=22, fill=palette["background"], outline=palette["accent"], width=4)
    art_path = identity.get("artwork_file")
    if art_path:
        with Image.open(package / str(art_path)) as source:
            artwork = ImageOps.contain(source.convert("RGB"), (right - left - 28, bottom - top - 28), method=Image.Resampling.LANCZOS)
        x = left + (right - left - artwork.width) // 2
        y = top + (bottom - top - artwork.height) // 2
        canvas.paste(artwork, (x, y))
        return

    # E-book-free profiles still receive a topic-specific visual mark.
    text = f"{identity['name']} {identity['motif']}".casefold()
    cx, cy = (left + right) // 2, (top + bottom) // 2
    if any(word in text for word in ("garden", "harvest", "leaf", "plant")):
        draw.line((left + 75, bottom - 55, right - 75, top + 55), fill=palette["accent"], width=16)
        for x, y in ((left + 135, top + 85), (cx, cy), (right - 135, bottom - 85)):
            draw.ellipse((x - 52, y - 27, x + 52, y + 27), fill=palette["secondary"])
    elif any(word in text for word in ("move", "landing", "route", "box")):
        draw.line((left + 75, bottom - 75, cx, top + 75, right - 75, bottom - 95), fill=palette["accent"], width=14)
        for x, y in ((left + 120, bottom - 90), (cx, top + 95), (right - 115, bottom - 110)):
            draw.rectangle((x - 42, y - 42, x + 42, y + 42), fill=palette["secondary"], outline=palette["ink"], width=3)
    elif any(word in text for word in ("home", "renovation", "roof", "room")):
        draw.polygon(((left + 65, cy), (cx, top + 48), (right - 65, cy), (right - 65, bottom - 45), (left + 65, bottom - 45)), fill=palette["secondary"])
        draw.line((left + 65, cy, cx, top + 48, right - 65, cy), fill=palette["ink"], width=12)
        draw.rectangle((cx - 38, cy + 50, cx + 38, bottom - 45), fill=palette["accent"])
    elif any(word in text for word in ("wedding", "ring", "ribbon")):
        draw.ellipse((cx - 140, cy - 95, cx + 35, cy + 80), outline=palette["accent"], width=20)
        draw.ellipse((cx - 35, cy - 80, cx + 140, cy + 95), outline=palette["secondary"], width=20)
        draw.polygon(((cx, top + 40), (cx + 18, cy - 24), (right - 70, top + 70), (cx + 50, cy + 10)), fill=palette["highlight"])
    else:
        draw.ellipse((cx - 125, cy - 125, cx + 125, cy + 125), fill=palette["accent"])
        draw.text((cx - 34, cy - 62), str(identity["name"][:1]).upper(), font=font(108, True), fill=WHITE)


def assert_unique_active_identity(product_id: int, spec: dict[str, Any]) -> None:
    """Fail marketing generation when an active product copies another brand identity."""
    declared = spec.get("visual_identity") or {}
    name = str(declared.get("name", "")).strip().casefold()
    motif = str(declared.get("motif", "")).strip().casefold()
    if not name or not motif:
        raise ValueError(f"Product {product_id} needs a named visual_identity and a distinctive motif before marketing generation")
    if product_id != 1:
        layout = str(declared.get("marketing_layout", "")).strip()
        artwork_ref = str(declared.get("marketing_artwork_file", "")).strip()
        art_direction = str(declared.get("marketing_art_direction", "")).strip()
        artwork = ROOT / "outputs" / "listings" / f"product-{product_id}" / "draft_package" / artwork_ref
        if not layout or not artwork_ref or not art_direction or not artwork.is_file():
            raise ValueError(
                f"Product {product_id} needs a distinct marketing_layout, a product-specific generated artwork file, "
                "and a written marketing_art_direction before marketing generation"
            )
    try:
        from factory_config import active_product_ids

        product_ids = active_product_ids()
    except (ImportError, OSError, ValueError):
        product_ids = ()
    for other_id in product_ids:
        if other_id == product_id:
            continue
        other_spec_path = ROOT / "outputs" / "listings" / f"product-{other_id}" / "design_spec.json"
        if not other_spec_path.is_file():
            continue
        other_spec = json.loads(other_spec_path.read_text(encoding="utf-8"))
        other_identity = other_spec.get("visual_identity") or {}
        other_name = str(other_identity.get("name", "")).strip().casefold()
        other_motif = str(other_identity.get("motif", "")).strip().casefold()
        if name == other_name or motif == other_motif:
            raise ValueError(f"Product {product_id} reuses the visual identity of product {other_id}; choose a distinct name and motif")
        layout = str(declared.get("marketing_layout", "")).strip().casefold()
        other_layout = str(other_identity.get("marketing_layout", "")).strip().casefold()
        if layout and other_layout and layout == other_layout:
            raise ValueError(f"Product {product_id} reuses marketing layout {layout!r} from product {other_id}; choose a different composition")


def cell_text(value: Any, limit: int = 34) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        try:
            return value.strftime("%d %b %Y")
        except (ValueError, TypeError):
            pass
    if isinstance(value, float):
        return f"{value:,.2f}".rstrip("0").rstrip(".")
    text = str(value).replace("\n", " ").strip()
    if text.startswith("="):
        return ""
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def cell_fill(cell, fallback: str) -> str:
    try:
        color = cell.fill.fgColor
        if color.type == "rgb" and color.rgb:
            value = str(color.rgb)[-6:]
            if re.fullmatch(r"[0-9A-Fa-f]{6}", value) and value.upper() not in {"FFFFFF", "000000"}:
                return f"#{value}"
    except (AttributeError, TypeError):
        pass
    return fallback


def cell_color(cell, fallback: str) -> str:
    try:
        color = cell.font.color
        if color and color.type == "rgb" and color.rgb:
            value = str(color.rgb)[-6:]
            if re.fullmatch(r"[0-9A-Fa-f]{6}", value):
                return f"#{value}"
    except (AttributeError, TypeError):
        pass
    return fallback


def render_sheet(workbook_path: Path, sheet_name: str, output: Path, spec: dict[str, Any]) -> None:
    """Render genuine worksheet cells as a consistent, legible listing preview."""
    wb = load_workbook(workbook_path, data_only=True, read_only=False)
    try:
        ws = wb[sheet_name]
        max_row = min(max(ws.max_row, 1), 24)
        max_col = min(max(ws.max_column, 1), 20)
        source_rows = [row for row in range(1, max_row + 1) if any(ws.cell(row, col).value is not None for col in range(1, max_col + 1))]
        source_cols = [col for col in range(1, max_col + 1) if any(ws.cell(row, col).value is not None for row in source_rows)]
        source_rows = source_rows[:17]
        source_cols = source_cols[:10]
        max_row = len(source_rows)
        max_col = len(source_cols)
        rows = [[cell_text(ws.cell(row, col).value) for col in source_cols] for row in source_rows]
        if not rows:
            rows = [[sheet_name]]

        palette = palette_for(spec)
        image = Image.new("RGB", SCREEN_SIZE, palette["background"])
        draw = ImageDraw.Draw(image)
        draw.rectangle((0, 0, SCREEN_SIZE[0], 92), fill=palette["ink"])
        draw.text((42, 23), str(spec.get("title", "Excel workbook"))[:72], font=font(29, True), fill=WHITE)
        draw.text((48, 125), sheet_name, font=font(53, True), fill=palette["ink"])
        draw.text((50, 192), "WORKBOOK PREVIEW · SAMPLE DATA", font=font(20, True), fill=palette["accent"])

        left, top, right = 48, 250, SCREEN_SIZE[0] - 48
        table_width = right - left
        weights = []
        for col in range(max_col):
            longest = max((len(row[col]) for row in rows if col < len(row)), default=8)
            weights.append(max(0.75, min(2.4, longest / 18)))
        weight_sum = sum(weights) or 1
        widths = [int(table_width * weight / weight_sum) for weight in weights]
        widths[-1] += table_width - sum(widths)
        row_height = min(58, max(39, int((SCREEN_SIZE[1] - top - 115) / max(1, len(rows)))))
        body_font = font(19)
        header_font = font(18, True)
        for row_index, row in enumerate(rows):
            y0 = top + row_index * row_height
            x = left
            for col_index, width in enumerate(widths):
                cell = ws.cell(source_rows[row_index], source_cols[col_index])
                value = row[col_index] if col_index < len(row) else ""
                is_header = row_index == 0 or cell.font.bold
                background = cell_fill(cell, palette["accent"] if row_index == 0 else (WHITE if row_index % 2 else "#F8FAFA"))
                draw.rectangle((x, y0, x + width, y0 + row_height), fill=background, outline=GRID, width=1)
                text_color = cell_color(cell, WHITE if is_header and row_index == 0 else palette["ink"])
                use_font = header_font if is_header else body_font
                max_chars = max(6, int(width / (use_font.size * 0.56)))
                if len(value) > max_chars:
                    value = value[: max_chars - 1].rstrip() + "…"
                draw.text((x + 12, y0 + max(8, (row_height - use_font.size) // 2 - 2)), value, font=use_font, fill=text_color)
                x += width
        footer_y = SCREEN_SIZE[1] - 70
        draw.text((48, footer_y), "Illustrative preview rendered from the included workbook · Replace sample values with your own", font=font(17), fill=MUTED)
        output.parent.mkdir(parents=True, exist_ok=True)
        image.save(output, format="PNG", optimize=True)
    finally:
        wb.close()


def draw_wrapped(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, *, width: int, font_obj, fill: str, spacing: int = 8) -> int:
    lines: list[str] = []
    for paragraph in str(text).splitlines() or [""]:
        lines.extend(textwrap.wrap(paragraph, width=width) or [""])
    draw.multiline_text(xy, "\n".join(lines), font=font_obj, fill=fill, spacing=spacing)
    return len(lines)


def fit_preview(path: Path, size: tuple[int, int]) -> Image.Image:
    preview = Image.open(path).convert("RGB")
    return ImageOps.fit(preview, size, method=Image.Resampling.LANCZOS)


def device_frame(screen: Image.Image, label: str, palette: dict[str, str]) -> Image.Image:
    screen = screen.convert("RGB")
    frame = Image.new("RGBA", (screen.width + 48, screen.height + 108), (0, 0, 0, 0))
    draw = ImageDraw.Draw(frame)
    draw.rounded_rectangle((0, 0, frame.width - 1, frame.height - 1), radius=30, fill=palette["ink"])
    for x, color in ((22, palette["secondary"]), (47, palette["highlight"]), (72, palette["accent"])):
        draw.ellipse((x, 16, x + 13, 29), fill=color)
    draw.text((112, 14), label[:36], font=font(18, True), fill=WHITE)
    frame.paste(screen, (24, 52))
    return frame


def package_deliverables(manifest: dict[str, Any], package: Path, spec: dict[str, Any] | None = None) -> list[dict[str, str]]:
    """Describe only customer files that are actually present in the kit."""
    deliverables: list[dict[str, str]] = []
    for key, label, kind in (
        ("workbook_file", "Example-data workbook", "workbook"),
        ("blank_file", "Blank workbook", "workbook"),
        ("usage_guide", "Excel usage guide (Markdown)", "guide"),
        ("usage_guide_pdf", "Excel user guide (PDF)", "pdf"),
        ("ebook_file", "Topic e-book (PDF)", "ebook"),
    ):
        value = manifest.get(key)
        if value and (package / str(value)).is_file():
            deliverables.append({"name": Path(str(value)).name, "label": label, "kind": kind})

    archive_name = manifest.get("digital_kit_zip") or manifest.get("local_file")
    archive = package / str(archive_name or "")
    if archive.is_file():
        try:
            with zipfile.ZipFile(archive) as bundle:
                for name in bundle.namelist():
                    path = Path(name)
                    suffix = path.suffix.casefold()
                    if suffix not in {".pdf", ".html", ".htm"} or path.name.casefold() in {
                        Path(str(manifest.get("usage_guide_pdf") or "")).name.casefold(),
                        Path(str(manifest.get("ebook_file") or "")).name.casefold(),
                    }:
                        continue
                    if suffix in {".html", ".htm"}:
                        label, kind = "Interactive HTML tool (offline)", "html"
                    elif path.stem.casefold() == "kit-index":
                        label = "Digital kit contents index (PDF)"
                        kind = "pdf"
                    else:
                        label = f"{path.stem.replace('-', ' ').replace('_', ' ').title()} (PDF)"
                        kind = "pdf"
                    deliverables.append({"name": path.name, "label": label, "kind": kind})
        except (OSError, zipfile.BadZipFile):
            pass
    seen: set[str] = set()
    unique: list[dict[str, str]] = []
    copy = inventory_copy(spec) if spec else None
    for item in deliverables:
        key = item["name"].casefold()
        if key not in seen:
            seen.add(key)
            if copy:
                item["label"] = inventory_item_label(item, copy)
            unique.append(item)
    return unique


def marketing_images(
    spec: dict[str, Any], screenshots: list[Path], package: Path, manifest: dict[str, Any],
    calculator_preview: Path | None = None,
) -> list[Path]:
    declared_identity = spec.get("visual_identity") or {}
    layout_id = str(declared_identity.get("marketing_layout", "")).strip()
    artwork_ref = str(declared_identity.get("marketing_artwork_file") or manifest.get("marketing_artwork_file") or "").strip()
    if layout_id and artwork_ref:
        return campaign_marketing_images(spec, screenshots, package, manifest, layout_id, artwork_ref, calculator_preview)

    palette = palette_for(spec)
    title = str(spec.get("title", "Excel Planner"))
    subtitle = str(spec.get("subtitle", "A practical workbook designed to make everyday planning clearer."))
    scope = [str(item) for item in spec.get("functional_scope", []) if str(item).strip()]
    output_paths: list[Path] = []
    screen_by_name = {path.stem: path for path in screenshots}
    lead = screenshots[0]
    deliverables = package_deliverables(manifest, package, spec)
    identity = marketing_identity(spec, package, manifest)

    # Product-led cover: show the real workbook and the files the customer receives.
    cover = Image.new("RGB", CANVAS, palette["background"])
    draw = ImageDraw.Draw(cover)
    draw.rounded_rectangle((65, 65, 1935, 1435), radius=42, fill=WHITE, outline=GRID, width=3)
    draw.text((125, 125), "DIGITAL PLANNING KIT", font=font(25, True), fill=palette["accent"])
    draw_wrapped(draw, (120, 185), title, width=31, font_obj=font(61, True), fill=palette["ink"], spacing=3)
    draw_wrapped(draw, (125, 395), subtitle, width=56, font_obj=font(25), fill=MUTED, spacing=6)
    draw_identity_artwork(cover, (1425, 135, 1830, 455), identity, package)
    draw = ImageDraw.Draw(cover)
    draw_identity_mark(draw, (1425, 470, 1830, 555), identity)
    if calculator_preview and calculator_preview.is_file():
        workbook_frame = device_frame(fit_preview(lead, (880, 495)), "EDITABLE EXCEL WORKBOOK", palette)
        calculator_frame = device_frame(fit_preview(calculator_preview, (810, 456)), "INTERACTIVE HTML CALCULATOR", palette)
        cover.paste(workbook_frame, (75, 620), workbook_frame)
        cover.paste(calculator_frame, (1035, 650), calculator_frame)
        draw = ImageDraw.Draw(cover)
        draw.rounded_rectangle((805, 1215, 1800, 1325), radius=22, fill=palette["ink"])
        draw.text((845, 1252), "THE EXCEL PLAN + THE INTERACTIVE TOOL", font=font(24, True), fill=WHITE)
    else:
        cover_frame = device_frame(fit_preview(lead, (1040, 610)), "EDITABLE EXCEL WORKBOOK", palette)
        cover.paste(cover_frame, (105, 590), cover_frame)
        card_x = 1325
        draw.rounded_rectangle((card_x, 610, 1845, 1265), radius=28, fill=palette["ink"])
        draw.text((1370, 655), "INCLUDED IN YOUR KIT", font=font(22, True), fill=WHITE)
        labels = [item["label"] for item in deliverables]
        if not labels:
            labels = ["Excel workbook and digital resources"]
        for index, label in enumerate(labels[:7]):
            y = 720 + index * 70
            draw.ellipse((1370, y + 6, 1390, y + 26), fill=palette["secondary"] if index % 2 else palette["accent"])
            draw_wrapped(draw, (1410, y), label, width=27, font_obj=font(20, True), fill=WHITE, spacing=3)
    draw = ImageDraw.Draw(cover)
    draw.text((125, 1335), "INSTANT DIGITAL DOWNLOAD - NO PHYSICAL ITEM", font=font(19, True), fill=palette["ink"])
    cover_path = package / "marketing_mockups" / "00-product-cover.png"
    cover_path.parent.mkdir(parents=True, exist_ok=True)
    cover.save(cover_path, optimize=True)
    output_paths.append(cover_path)

    # Hero mockup.
    hero = Image.new("RGBA", CANVAS, palette["background"])
    draw = ImageDraw.Draw(hero)
    draw.text((115, 88), "EDITABLE EXCEL WORKBOOK", font=font(24, True), fill=palette["accent"])
    draw_wrapped(draw, (110, 150), title, width=29, font_obj=font(68, True), fill=palette["ink"], spacing=2)
    draw_wrapped(draw, (118, 360), subtitle, width=55, font_obj=font(27), fill=MUTED, spacing=8)
    hero_frame = device_frame(fit_preview(lead, (1190, 690)), "REAL WORKBOOK PREVIEW", palette)
    hero.alpha_composite(hero_frame, (625, 570))
    draw = ImageDraw.Draw(hero)
    draw.rounded_rectangle((115, 590, 590, 1040), radius=26, fill=WHITE, outline=GRID, width=3)
    draw.text((160, 650), "INCLUDED", font=font(24, True), fill=palette["accent"])
    for index, line in enumerate(scope[:5]):
        y = 715 + index * 58
        draw.ellipse((160, y + 7, 178, y + 25), fill=palette["secondary"] if index % 2 else palette["accent"])
        draw_wrapped(draw, (196, y), line, width=29, font_obj=font(21), fill=palette["ink"], spacing=3)
    draw.text((115, 1375), "DIGITAL DOWNLOAD · INSTANT ACCESS", font=font(20, True), fill=palette["ink"])
    hero_path = package / "marketing_mockups" / "01-hero-mockup.png"
    hero_path.parent.mkdir(parents=True, exist_ok=True)
    hero.convert("RGB").save(hero_path, optimize=True)
    output_paths.append(hero_path)

    # Feature grid: three genuine workbook sheet previews.
    grid = Image.new("RGB", CANVAS, palette["background"])
    draw = ImageDraw.Draw(grid)
    draw.text((100, 75), "ONE WORKBOOK, CONNECTED VIEWS", font=font(23, True), fill=palette["accent"])
    draw_wrapped(draw, (96, 125), "A clear view of the work", width=35, font_obj=font(58, True), fill=palette["ink"])
    cards = screenshots[:3]
    card_width = 570
    for index, path in enumerate(cards):
        x = 90 + index * 630
        draw.rounded_rectangle((x, 330, x + card_width, 1280), radius=25, fill=WHITE, outline=GRID, width=2)
        grid.paste(fit_preview(path, (card_width - 36, 410)), (x + 18, 350))
        draw.text((x + 26, 800), path.stem.replace("-", " ").replace("workbook preview", "" ).strip().title(), font=font(28, True), fill=palette["ink"])
        feature = scope[index] if index < len(scope) else "A focused worksheet with room for your own information."
        draw_wrapped(draw, (x + 27, 870), feature, width=36, font_obj=font(22), fill=MUTED)
    grid_path = package / "marketing_mockups" / "02-feature-grid-mockup.png"
    grid_path.parent.mkdir(parents=True, exist_ok=True)
    grid.save(grid_path, optimize=True)
    output_paths.append(grid_path)

    # What's-included and simple workflow cards.
    includes = Image.new("RGB", CANVAS, palette["background"])
    draw = ImageDraw.Draw(includes)
    draw.text((105, 95), "WHAT'S INSIDE", font=font(23, True), fill=palette["accent"])
    draw_wrapped(draw, (100, 150), "Everything starts in one place", width=34, font_obj=font(58, True), fill=palette["ink"])
    frame = device_frame(fit_preview(screenshots[min(1, len(screenshots) - 1)], (1040, 610)), "EDITABLE EXCEL", palette)
    includes.paste(frame, (75, 490), frame)
    draw = ImageDraw.Draw(includes)
    draw.rounded_rectangle((1210, 505, 1885, 1210), radius=28, fill=palette["ink"])
    declared_sheets = [str(value) for value in spec.get("tabs", [])][:8]
    included_lines = [item["label"] for item in deliverables]
    if not included_lines:
        included_lines = ["Digital files listed in the product download"]
    included_lines.extend(f"Workbook tab: {name}" for name in declared_sheets[:5])
    for index, line in enumerate(included_lines[:9]):
        y = 565 + index * 65
        draw.ellipse((1260, y + 5, 1279, y + 24), fill=palette["secondary"] if index % 2 else palette["accent"])
        draw_wrapped(draw, (1300, y), line, width=40, font_obj=font(22, True), fill=WHITE, spacing=2)
    included_path = package / "marketing_mockups" / "03-whats-included-mockup.png"
    included_path.parent.mkdir(parents=True, exist_ok=True)
    includes.save(included_path, optimize=True)
    output_paths.append(included_path)

    how = Image.new("RGB", CANVAS, palette["background"])
    draw = ImageDraw.Draw(how)
    draw.text((105, 95), "A SIMPLE WORKFLOW", font=font(23, True), fill=palette["accent"])
    draw_wrapped(draw, (100, 150), "From first entry to useful overview", width=30, font_obj=font(58, True), fill=palette["ink"])
    steps = [
        ("01", "Open the guide", "Start with the included PDF instructions and the workbook's opening sheet."),
        ("02", "Add your details", "Replace the sample information with the inputs relevant to your situation."),
        ("03", "Review the connected views", "Use the workbook's summaries and worksheets to keep the plan up to date."),
    ]
    for index, (number, heading, body) in enumerate(steps):
        x = 100 + index * 615
        draw.rounded_rectangle((x, 460, x + 550, 1195), radius=30, fill=WHITE, outline=GRID, width=2)
        draw.ellipse((x + 45, 520, x + 145, 620), fill=palette["accent"] if index != 1 else palette["secondary"])
        draw.text((x + 65, 545), number, font=font(28, True), fill=WHITE)
        draw_wrapped(draw, (x + 42, 690), heading, width=25, font_obj=font(34, True), fill=palette["ink"])
        draw_wrapped(draw, (x + 42, 790), body, width=34, font_obj=font(25), fill=MUTED, spacing=10)
    how_path = package / "marketing_mockups" / "04-how-it-works-mockup.png"
    how_path.parent.mkdir(parents=True, exist_ok=True)
    how.save(how_path, optimize=True)
    output_paths.append(how_path)

    # Dedicated workbook-templates mockup: make the two editions unmistakable.
    templates = Image.new("RGB", CANVAS, palette["background"])
    draw = ImageDraw.Draw(templates)
    draw.text((105, 95), "TWO EDITIONS INCLUDED", font=font(23, True), fill=palette["accent"])
    draw_wrapped(draw, (100, 150), "Start with an example or a clean copy", width=34, font_obj=font(54, True), fill=palette["ink"])
    workbook_items = [item for item in deliverables if item["kind"] == "workbook"][:2]
    if not workbook_items:
        workbook_items = [
            {"name": str(manifest.get("workbook_file") or "Example workbook.xlsx"), "label": "Example-data edition"},
            {"name": str(manifest.get("blank_file") or "Blank workbook.xlsx"), "label": "Blank edition"},
        ]
    for index, item in enumerate(workbook_items):
        x = 105 + index * 905
        preview = fit_preview(screenshots[min(index, len(screenshots) - 1)], (800, 480))
        frame = device_frame(preview, "EXCEL TEMPLATE", palette)
        templates.paste(frame, (x + 28, 370), frame)
        draw = ImageDraw.Draw(templates)
        draw.rounded_rectangle((x, 1020, x + 820, 1320), radius=24, fill=WHITE, outline=GRID, width=2)
        draw.text((x + 35, 1060), item["label"], font=font(28, True), fill=palette["ink"])
        draw_wrapped(draw, (x + 35, 1110), item["name"], width=49, font_obj=font(19), fill=MUTED)
    templates_path = package / "marketing_mockups" / "05-excel-templates-mockup.png"
    templates_path.parent.mkdir(parents=True, exist_ok=True)
    templates.save(templates_path, optimize=True)
    output_paths.append(templates_path)

    # Dedicated PDF-resources mockup with the real e-book cover, if the kit has one.
    pdf_items = [item for item in deliverables if item["kind"] in {"ebook", "pdf"}]
    if pdf_items:
        resources = Image.new("RGB", CANVAS, palette["background"])
        draw = ImageDraw.Draw(resources)
        draw.text((105, 95), "DIGITAL READING & REFERENCE", font=font(23, True), fill=palette["accent"])
        draw_wrapped(draw, (100, 150), "Practical guidance, ready to open", width=34, font_obj=font(54, True), fill=palette["ink"])
        ebook_cover = package / str(manifest.get("ebook_cover_image") or "")
        if not ebook_cover.is_file() and manifest.get("cover_image"):
            candidate = package / str(manifest.get("cover_image"))
            if candidate.is_file() and "cover" in candidate.name.casefold():
                ebook_cover = candidate
        cards = pdf_items[:4]
        for index, item in enumerate(cards):
            x = 125 + index * 455
            y = 410
            draw.rounded_rectangle((x, y, x + 385, 1260), radius=24, fill=WHITE, outline=GRID, width=2)
            art_box = (x + 65, y + 65, x + 320, y + 440)
            if item["kind"] == "ebook" and ebook_cover.is_file():
                try:
                    with Image.open(ebook_cover) as source:
                        art = ImageOps.contain(source.convert("RGB"), (250, 375))
                    resources.paste(art, (x + (385 - art.width) // 2, y + 65))
                except OSError:
                    draw.rounded_rectangle(art_box, radius=15, fill=palette["ink"])
            else:
                draw.rounded_rectangle(art_box, radius=15, fill=palette["ink"])
                draw.text((x + 155, y + 155), "PDF", font=font(52, True), fill=WHITE)
                draw.rectangle((x + 120, y + 265, x + 300, y + 275), fill=palette["accent"])
                draw.rectangle((x + 120, y + 295, x + 280, y + 305), fill=palette["secondary"])
            draw_wrapped(draw, (x + 28, y + 500), item["label"], width=25, font_obj=font(23, True), fill=palette["ink"])
            draw_wrapped(draw, (x + 28, y + 590), item["name"], width=31, font_obj=font(17), fill=MUTED)
        resources_path = package / "marketing_mockups" / "06-pdf-resources-mockup.png"
        resources_path.parent.mkdir(parents=True, exist_ok=True)
        resources.save(resources_path, optimize=True)
        output_paths.append(resources_path)
    for path in output_paths[1:]:
        with Image.open(path) as existing:
            branded = existing.convert("RGB")
        draw_identity_signature(ImageDraw.Draw(branded), identity)
        branded.save(path, optimize=True)
    inventory_path = render_deliverables_inventory(spec, package, manifest, identity)
    calculator_path = render_html_calculator_mockup(spec, package, manifest, identity, calculator_preview) if calculator_preview else None
    render_ebook_highlight_mockup(spec, package, manifest, identity)
    extras = [path for path in (inventory_path, calculator_path) if path]
    return [output_paths[0], *extras, *[path for path in output_paths[1:] if path not in extras]]


def campaign_marketing_images(
    spec: dict[str, Any], screenshots: list[Path], package: Path, manifest: dict[str, Any],
    layout_id: str, artwork_ref: str, calculator_preview: Path | None = None,
) -> list[Path]:
    """Compose product-specific campaign mockups around real workbook captures."""
    layouts = {"moving_atlas", "interiors_editorial", "aquarium_fieldguide", "puppy_firstyear_editorial", "garage_service_manual"}
    if layout_id not in layouts:
        raise ValueError(f"Unknown marketing_layout {layout_id!r}; add a distinct campaign composition before use")
    artwork_path = package / artwork_ref
    if not artwork_path.is_file():
        raise FileNotFoundError(f"Missing bespoke marketing artwork: {artwork_path}")

    palette = palette_for(spec)
    title = str(spec.get("title", "Digital planning kit"))
    subtitle = str(spec.get("subtitle", "A practical digital planning kit."))
    scope = [str(item) for item in spec.get("functional_scope", []) if str(item).strip()]
    deliverables = package_deliverables(manifest, package, spec)
    identity = marketing_identity(spec, package, manifest)
    output_dir = package / "marketing_mockups"
    output_dir.mkdir(parents=True, exist_ok=True)
    artwork = ImageOps.fit(Image.open(artwork_path).convert("RGB"), CANVAS, method=Image.Resampling.LANCZOS)
    previews = [fit_preview(path, (760, 430)) for path in screenshots]
    paths: list[Path] = []

    def save(canvas: Image.Image, name: str) -> Path:
        draw_identity_signature(ImageDraw.Draw(canvas), identity)
        path = output_dir / name
        canvas.convert("RGB").save(path, optimize=True)
        paths.append(path)
        return path

    # Each art direction uses a distinct campaign composition, not a reused mockup template.
    if layout_id == "moving_atlas":
        cover = artwork.copy()
        draw = ImageDraw.Draw(cover, "RGBA")
        draw.rounded_rectangle((62, 68, 975, 1428), radius=40, fill=(248, 247, 235, 238))
        draw.text((120, 125), "THE MOVE, MAPPED", font=font(24, True), fill=palette["accent"])
        draw_wrapped(draw, (115, 195), title, width=22, font_obj=font(62, True), fill=palette["ink"], spacing=4)
        draw_wrapped(draw, (122, 410), subtitle, width=46, font_obj=font(25), fill=palette["ink"], spacing=7)
        frame = device_frame(previews[0], "YOUR MOVE DASHBOARD", palette)
        cover.alpha_composite(frame, (190, 690)) if cover.mode == "RGBA" else cover.paste(frame, (190, 690), frame)
        if calculator_preview and calculator_preview.is_file():
            calc_frame = device_frame(fit_preview(calculator_preview, (650, 366)), "OFFLINE HTML CALCULATOR", palette)
            cover.paste(calc_frame, (1175, 680), calc_frame)
        draw = ImageDraw.Draw(cover)
        draw.rounded_rectangle((118, 1235, 884, 1368), radius=22, fill=palette["ink"])
        labels = [item["label"] for item in deliverables[:4]]
        draw_wrapped(draw, (150, 1260), "  /  ".join(labels) or "Excel workbook + digital resources", width=62, font_obj=font(19, True), fill=WHITE, spacing=4)
        save(cover, "00-product-cover.png")

        route = artwork.copy()
        draw = ImageDraw.Draw(route, "RGBA")
        draw.rectangle((0, 0, 2000, 365), fill=(36, 55, 71, 238))
        draw.text((100, 70), "01  /  THE ROUTE THROUGH YOUR MOVE", font=font(26, True), fill=palette["secondary"])
        draw_wrapped(draw, (95, 125), "Plan the milestones. Keep the details together.", width=45, font_obj=font(49, True), fill=WHITE)
        route.paste(previews[0], (100, 460))
        draw = ImageDraw.Draw(route)
        for i, step in enumerate(("PLAN", "COMPARE", "MOVE", "SETTLE")):
            x = 960 + i * 240
            draw.ellipse((x, 570, x + 62, 632), fill=palette["secondary"] if i % 2 == 0 else palette["accent"])
            draw.text((x + 18, 587), str(i + 1), font=font(22, True), fill=WHITE)
            draw.text((x - 20, 660), step, font=font(19, True), fill=palette["ink"])
            if i < 3:
                draw.line((x + 65, 601, x + 230, 601), fill=palette["secondary"], width=8)
        draw.rounded_rectangle((960, 805, 1810, 980), radius=26, fill=palette["ink"])
        draw.text((1005, 840), "THE MOVE-DAY PLANNER", font=font(22, True), fill=palette["secondary"])
        draw_wrapped(draw, (1005, 885), "Budget  /  quotes  /  purchases  /  dates", width=42, font_obj=font(25, True), fill=WHITE)
        save(route, "01-route-ledger-mockup.png")

        board = Image.new("RGB", CANVAS, palette["background"])
        board.paste(ImageOps.fit(artwork, (610, 1500), method=Image.Resampling.LANCZOS), (0, 0))
        draw = ImageDraw.Draw(board)
        draw.rectangle((610, 0, 2000, 1500), fill="#F8F7EB")
        draw.text((690, 85), "02  /  MOVE-DAY CONTROL BOARD", font=font(23, True), fill=palette["accent"])
        draw_wrapped(draw, (680, 145), "Budget, quotes and purchases in one view", width=35, font_obj=font(48, True), fill=palette["ink"])
        for i, preview in enumerate(previews[:3]):
            x = 690 + (i % 2) * 625
            y = 420 + (i // 2) * 430
            board.paste(preview.resize((560, 316)), (x, y))
            draw = ImageDraw.Draw(board)
            label = (scope[i] if i < len(scope) else ("Budget", "Quotes", "Purchases")[i])
            draw.text((x, y + 330), label[:40], font=font(25, True), fill=palette["ink"])
        save(board, "02-move-day-board-mockup.png")

        included = artwork.copy()
        draw = ImageDraw.Draw(included, "RGBA")
        draw.rectangle((0, 0, 2000, 1500), fill=(243, 247, 248, 215))
        draw.text((100, 90), "03  /  WHAT TRAVELS WITH YOU", font=font(23, True), fill=palette["accent"])
        draw_wrapped(draw, (95, 145), "The kit, laid out like a packing list", width=38, font_obj=font(50, True), fill=palette["ink"])
        for i, item in enumerate(deliverables[:6]):
            x = 110 + (i % 2) * 920
            y = 390 + (i // 2) * 275
            draw.rounded_rectangle((x, y, x + 820, y + 210), radius=28, fill=(255, 255, 255, 245), outline=palette["secondary"], width=3)
            draw.text((x + 28, y + 25), f"{i + 1:02d}  /  INCLUDED", font=font(18, True), fill=palette["accent"])
            draw_wrapped(draw, (x + 30, y + 72), item["label"], width=43, font_obj=font(28, True), fill=palette["ink"])
        save(included, "03-whats-included-mockup.png")

        process = Image.new("RGB", CANVAS, palette["ink"])
        process.paste(ImageOps.fit(artwork, (770, 1500), method=Image.Resampling.LANCZOS), (1230, 0))
        draw = ImageDraw.Draw(process)
        draw.text((100, 95), "04  /  FROM FIRST BOX TO NEW KEYS", font=font(23, True), fill=palette["secondary"])
        draw_wrapped(draw, (95, 155), "A practical sequence for the weeks around moving day", width=34, font_obj=font(48, True), fill=WHITE)
        for i, label in enumerate(("Gather the dates", "Compare the real quotes", "Track spend as it happens", "Close out the move")):
            y = 480 + i * 205
            draw.ellipse((110, y, 180, y + 70), fill=palette["secondary"] if i % 2 == 0 else palette["accent"])
            draw.text((133, y + 18), str(i + 1), font=font(25, True), fill=WHITE)
            draw.text((220, y + 12), label, font=font(28, True), fill=WHITE)
        save(process, "04-moving-sequence-mockup.png")
    elif layout_id == "aquarium_fieldguide":
        cover = artwork.copy()
        draw = ImageDraw.Draw(cover, "RGBA")
        draw.rounded_rectangle((70, 72, 880, 955), radius=34, fill=(242, 245, 235, 238))
        draw.text((120, 125), "STILLWATER  /  FIELD NOTES", font=font(23, True), fill=palette["accent"])
        draw_wrapped(draw, (112, 200), title, width=22, font_obj=font(59, True), fill=palette["ink"], spacing=4)
        draw_wrapped(draw, (120, 480), "Plan the habitat before the first purchase.", width=38, font_obj=font(26), fill=palette["ink"], spacing=7)
        draw.line((120, 635, 685, 635), fill=palette["secondary"], width=5)
        draw.text((120, 670), "SETUP  /  EQUIPMENT  /  READINESS", font=font(17, True), fill=palette["accent"])
        if calculator_preview and calculator_preview.is_file():
            # Give the shipped calculator and workbook equal, side-by-side
            # space on the cover rather than hiding the tool in a small callout.
            calc_frame = device_frame(fit_preview(calculator_preview, (820, 461)), "OFFLINE HTML CALCULATOR", palette)
            workbook_frame = device_frame(fit_preview(screenshots[0], (820, 461)), "EXCEL SETUP PLANNER", palette)
            cover.paste(workbook_frame, (60, 900), workbook_frame)
            cover.paste(calc_frame, (1060, 900), calc_frame)
        else:
            frame = device_frame(previews[0], "EXCEL SETUP PLANNER", palette)
            cover.paste(frame, (1030, 815), frame)
            draw = ImageDraw.Draw(cover)
            draw.rounded_rectangle((105, 1290, 1080, 1390), radius=18, fill=palette["ink"])
            draw_wrapped(draw, (140, 1318), "Includes: Excel setup planner  /  field worksheets  /  e-book", width=66, font_obj=font(18, True), fill=WHITE)
        save(cover, "00-product-cover.png")

        specimen = artwork.copy()
        draw = ImageDraw.Draw(specimen, "RGBA")
        draw.rectangle((0, 0, 2000, 285), fill=(32, 58, 59, 238))
        draw.text((95, 55), "01  /  THE SETUP SPECIMEN", font=font(22, True), fill=palette["highlight"])
        draw_wrapped(draw, (92, 108), "A clear plan for the things around the tank", width=42, font_obj=font(45, True), fill=WHITE)
        specimen.paste(previews[0], (95, 390))
        draw = ImageDraw.Draw(specimen)
        labels = ("MEASURE", "COMPARE", "SET UP", "REVIEW")
        for i, label in enumerate(labels):
            x = 1010 + i * 225
            draw.line((x, 500, x, 815), fill=palette["secondary"], width=3)
            draw.ellipse((x - 27, 530, x + 27, 584), fill=palette["accent"] if i % 2 == 0 else palette["secondary"])
            draw.text((x - 75, 610), label, font=font(17, True), fill=palette["ink"])
        draw_wrapped(draw, (1000, 880), scope[0] if scope else "Keep the room, equipment and next task in view.", width=40, font_obj=font(25, True), fill=palette["ink"])
        save(specimen, "01-tank-setup-specimen-mockup.png")

        ecosystem = Image.new("RGB", CANVAS, palette["ink"])
        ecosystem.paste(ImageOps.fit(artwork, (720, 1500), method=Image.Resampling.LANCZOS), (1280, 0))
        draw = ImageDraw.Draw(ecosystem)
        draw.text((95, 95), "02  /  A SYSTEM, NOT A SHOPPING LIST", font=font(22, True), fill=palette["highlight"])
        draw_wrapped(draw, (90, 155), "Room, support, equipment and care all connect", width=34, font_obj=font(47, True), fill=WHITE)
        for i, item in enumerate(("Measure the location", "Compare like-for-like equipment", "Keep setup costs together", "Record the questions to resolve")):
            y = 500 + i * 190
            draw.ellipse((105, y, 170, y + 65), fill=palette["secondary"] if i % 2 == 0 else palette["accent"])
            draw.text((127, y + 16), str(i + 1), font=font(23, True), fill=WHITE)
            draw.text((205, y + 10), item, font=font(25, True), fill=WHITE)
        save(ecosystem, "02-aquarium-setup-system-mockup.png")

        included = Image.new("RGB", CANVAS, palette["background"])
        included.paste(ImageOps.fit(artwork, (520, 1500), method=Image.Resampling.LANCZOS), (1480, 0))
        draw = ImageDraw.Draw(included)
        draw.text((90, 95), "03  /  WHAT'S IN THE KIT", font=font(22, True), fill=palette["accent"])
        draw_wrapped(draw, (85, 155), "A compact field kit for the setup stage", width=31, font_obj=font(45, True), fill=palette["ink"])
        for i, item in enumerate(deliverables[:6]):
            y = 430 + i * 162
            draw.line((105, y + 126, 1360, y + 126), fill=palette["secondary"], width=2)
            draw.text((110, y + 8), f"{i + 1:02d}", font=font(20, True), fill=palette["secondary"])
            draw_wrapped(draw, (190, y), item["label"], width=36, font_obj=font(24, True), fill=palette["ink"])
        save(included, "03-whats-included-mockup.png")

        volume = artwork.copy()
        draw = ImageDraw.Draw(volume, "RGBA")
        draw.rounded_rectangle((75, 75, 910, 760), radius=32, fill=(242, 245, 235, 242))
        draw.text((125, 125), "04  /  A QUICK VOLUME ESTIMATE", font=font(20, True), fill=palette["accent"])
        draw_wrapped(draw, (120, 190), "The calculator, in use", width=27, font_obj=font(46, True), fill=palette["ink"])
        draw_wrapped(draw, (125, 320), "A real browser preview of the included offline tool, paired with its companion Excel planner.", width=39, font_obj=font(24), fill=palette["ink"])
        if calculator_preview and calculator_preview.is_file():
            volume.paste(fit_preview(calculator_preview, (810, 456)), (1000, 470))
        else:
            draw.rounded_rectangle((124, 580, 785, 675), radius=17, fill=palette["accent"])
            draw.text((158, 606), "LENGTH x WIDTH x WATER DEPTH", font=font(19, True), fill=WHITE)
        save(volume, "04-tank-volume-tool-mockup.png")
    elif layout_id == "puppy_firstyear_editorial":
        cover = Image.new("RGB", CANVAS, palette["background"])
        hero = ImageOps.fit(artwork, (2000, 790), method=Image.Resampling.LANCZOS, centering=(0.5, 0.78))
        cover.paste(hero, (0, 0))
        draw = ImageDraw.Draw(cover, "RGBA")
        draw.rounded_rectangle((78, 65, 895, 680), radius=18, fill=(246, 241, 232, 242))
        draw.text((125, 112), "A FIRST YEAR TOGETHER", font=font(22, True), fill=palette["accent"])
        draw_wrapped(draw, (118, 175), title, width=24, font_obj=font(56, True), fill=palette["ink"], spacing=4)
        draw_wrapped(draw, (125, 405), "Make room for the puppy, the questions and the real costs.", width=41, font_obj=font(25), fill=palette["ink"], spacing=7)
        draw.rectangle((125, 585, 205, 591), fill=palette["secondary"])
        draw.text((125, 615), "PREPARE  /  PLAN  /  ADJUST", font=font(18, True), fill=palette["ink"])
        if calculator_preview and calculator_preview.is_file():
            workbook_frame = device_frame(fit_preview(screenshots[0], (820, 461)), "FIRST-YEAR EXCEL PLANNER", palette)
            calculator_frame = device_frame(fit_preview(calculator_preview, (820, 461)), "OFFLINE PUPPY COST CALCULATOR", palette)
            cover.paste(workbook_frame, (60, 830), workbook_frame)
            cover.paste(calculator_frame, (1060, 830), calculator_frame)
        else:
            cover.paste(device_frame(previews[0], "FIRST-YEAR EXCEL PLANNER", palette), (105, 830))
        draw = ImageDraw.Draw(cover)
        draw.rectangle((0, 1395, 2000, 1500), fill=palette["ink"])
        draw.text((100, 1432), "EXCEL + OFFLINE CALCULATOR + FIELD GUIDE + PREPARATION WORKSHEETS", font=font(19, True), fill=WHITE)
        save(cover, "00-product-cover.png")

        first_weeks = Image.new("RGB", CANVAS, palette["background"])
        first_weeks.paste(ImageOps.fit(artwork, (730, 1500), method=Image.Resampling.LANCZOS, centering=(0.5, 0.75)), (0, 0))
        draw = ImageDraw.Draw(first_weeks)
        draw.rectangle((730, 0, 2000, 1500), fill=palette["background"])
        draw.text((810, 105), "01  /  THE FIRST WEEK, MADE SMALLER", font=font(21, True), fill=palette["accent"])
        draw_wrapped(draw, (805, 165), "A few steady anchors beat a perfect checklist.", width=29, font_obj=font(44, True), fill=palette["ink"])
        for index, label in enumerate(("Choose a quiet home base", "Gather records and questions", "Agree on household hand-offs", "Leave room to adjust")):
            y = 505 + index * 205
            draw.ellipse((825, y, 888, y + 63), fill=palette["secondary"] if index % 2 == 0 else palette["accent"])
            draw.text((847, y + 16), str(index + 1), font=font(22, True), fill=WHITE)
            draw.text((925, y + 12), label, font=font(24, True), fill=palette["ink"])
            if index < 3:
                draw.line((856, y + 68, 856, y + 197), fill=palette["highlight"], width=4)
        save(first_weeks, "01-first-weeks-home-mockup.png")

        budget_board = Image.new("RGB", CANVAS, palette["ink"])
        draw = ImageDraw.Draw(budget_board)
        draw.text((92, 75), "02  /  COSTS YOU CAN ACTUALLY CHECK", font=font(22, True), fill=palette["highlight"])
        draw_wrapped(draw, (88, 135), "Start with your estimates. Update them as you learn.", width=42, font_obj=font(46, True), fill=WHITE)
        budget_board.paste(previews[1] if len(previews) > 1 else previews[0], (80, 390))
        if calculator_preview and calculator_preview.is_file():
            budget_board.paste(fit_preview(calculator_preview, (760, 428)), (1120, 410))
        draw = ImageDraw.Draw(budget_board)
        draw.rounded_rectangle((1120, 930, 1880, 1165), radius=16, fill=palette["background"])
        draw_wrapped(draw, (1160, 975), "One-time setup + 12 months of costs you enter yourself.", width=35, font_obj=font(25, True), fill=palette["ink"])
        draw_wrapped(draw, (1160, 1080), "A budgeting aid only — not veterinary or insurance advice.", width=45, font_obj=font(17), fill=palette["ink"])
        save(budget_board, "02-first-year-budget-mockup.png")

        questions = Image.new("RGB", CANVAS, palette["background"])
        questions.paste(ImageOps.fit(artwork, (670, 1500), method=Image.Resampling.LANCZOS, centering=(0.5, 0.75)), (1330, 0))
        draw = ImageDraw.Draw(questions)
        draw.rectangle((0, 0, 1370, 1500), fill=palette["background"])
        draw.text((88, 90), "04  /  TAKE BETTER QUESTIONS WITH YOU", font=font(22, True), fill=palette["accent"])
        draw_wrapped(draw, (82, 155), "A little preparation for a more useful conversation.", width=31, font_obj=font(45, True), fill=palette["ink"])
        for index, label in enumerate(("What records should I bring?", "What does this estimate include?", "Who should I call with follow-up questions?")):
            y = 465 + index * 245
            draw.rectangle((105, y, 1210, y + 170), fill=WHITE, outline=palette["secondary"], width=3)
            draw.text((145, y + 26), f"ASK  /  0{index + 1}", font=font(17, True), fill=palette["secondary"])
            draw_wrapped(draw, (145, y + 72), label, width=36, font_obj=font(28, True), fill=palette["ink"])
        save(questions, "04-provider-question-cards-mockup.png")
    elif layout_id == "garage_service_manual":
        # A service-ticket/ledger composition, deliberately unlike the other
        # products' route maps, room boards, specimen windows and pet editorial.
        cover = Image.new("RGB", CANVAS, palette["ink"])
        cover.paste(ImageOps.fit(artwork, (1180, 1500), method=Image.Resampling.LANCZOS, centering=(0.55, 0.5)), (820, 0))
        draw = ImageDraw.Draw(cover)
        draw.rectangle((0, 0, 900, 1500), fill=palette["background"])
        draw.rectangle((0, 0, 28, 1500), fill=palette["secondary"])
        draw.text((92, 85), "WORKSHOP LEDGER  /  FIELD EDITION", font=font(21, True), fill=palette["accent"])
        draw_wrapped(draw, (86, 155), title, width=20, font_obj=font(57, True), fill=palette["ink"], spacing=5)
        draw_wrapped(draw, (94, 400), "A clear record of service, quotes and the cost of keeping a vehicle.", width=36, font_obj=font(25), fill=palette["ink"], spacing=7)
        draw.line((95, 570, 700, 570), fill=palette["secondary"], width=5)
        for i, label in enumerate(("SERVICE HISTORY", "GARAGE QUOTES", "OWNERSHIP COSTS")):
            y = 610 + i * 65
            draw.text((100, y), f"0{i + 1}  /  {label}", font=font(18, True), fill=palette["accent"])
        if calculator_preview and calculator_preview.is_file():
            cover.paste(device_frame(fit_preview(calculator_preview, (740, 416)), "COST-PER-MILE TOOL", palette), (60, 930))
            cover.paste(device_frame(previews[0], "VEHICLE EXCEL PLANNER", palette), (1030, 930))
        else:
            cover.paste(device_frame(previews[0], "VEHICLE EXCEL PLANNER", palette), (175, 900))
        draw = ImageDraw.Draw(cover)
        draw.rectangle((0, 1390, 2000, 1500), fill=palette["ink"])
        draw.text((85, 1430), "EXCEL + OFFLINE CALCULATOR + FIELD GUIDE + SERVICE WORKSHEETS", font=font(18, True), fill=WHITE)
        save(cover, "00-product-cover.png")

        ticket = Image.new("RGB", CANVAS, palette["background"])
        draw = ImageDraw.Draw(ticket)
        draw.rectangle((0, 0, 2000, 260), fill=palette["ink"])
        draw.text((100, 62), "01  /  THE SERVICE TICKET", font=font(23, True), fill=palette["highlight"])
        draw_wrapped(draw, (95, 112), "Keep the work, mileage and paperwork connected.", width=40, font_obj=font(46, True), fill=WHITE)
        for i, label in enumerate(("DATE + ODOMETER", "WORK AUTHORIZED", "INVOICE + NEXT SOURCE")):
            y = 390 + i * 275
            draw.rectangle((105, y, 1885, y + 205), fill=WHITE, outline=palette["secondary"], width=3)
            draw.rectangle((105, y, 390, y + 205), fill=palette["accent"] if i % 2 == 0 else palette["secondary"])
            draw_wrapped(draw, (135, y + 66), f"0{i + 1}\n{label}", width=16, font_obj=font(22, True), fill=WHITE)
            draw.line((455, y + 74, 1780, y + 74), fill=palette["background"], width=3)
            draw.line((455, y + 140, 1780, y + 140), fill=palette["background"], width=3)
        save(ticket, "01-service-ticket-mockup.png")

        calculator = Image.new("RGB", CANVAS, palette["ink"])
        draw = ImageDraw.Draw(calculator)
        draw.text((100, 85), "02  /  COST PER MILE, IN CONTEXT", font=font(22, True), fill=palette["highlight"])
        draw_wrapped(draw, (94, 145), "Use your own costs and a matching time period.", width=38, font_obj=font(47, True), fill=WHITE)
        if calculator_preview and calculator_preview.is_file():
            calculator.paste(fit_preview(calculator_preview, (1160, 652)), (90, 450))
        calculator.paste(previews[1] if len(previews) > 1 else previews[0], (1280, 450))
        draw = ImageDraw.Draw(calculator)
        draw.text((100, 1170), "A planning ratio, not a repair forecast or safety assessment.", font=font(22, True), fill=palette["background"])
        save(calculator, "02-cost-per-mile-mockup.png")

        included = Image.new("RGB", CANVAS, palette["background"])
        included.paste(ImageOps.fit(artwork, (610, 1500), method=Image.Resampling.LANCZOS), (1390, 0))
        draw = ImageDraw.Draw(included)
        draw.text((90, 82), "03  /  WHAT'S IN THE KIT", font=font(22, True), fill=palette["accent"])
        draw_wrapped(draw, (84, 140), "A glovebox-to-desk record system", width=30, font_obj=font(47, True), fill=palette["ink"])
        for i, item in enumerate(deliverables[:6]):
            y = 400 + i * 165
            draw.rectangle((100, y, 1280, y + 125), fill=WHITE, outline=palette["secondary"], width=2)
            draw.text((125, y + 17), f"{i + 1:02d}  /  INCLUDED", font=font(16, True), fill=palette["accent"])
            draw_wrapped(draw, (355, y + 25), item["label"], width=38, font_obj=font(23, True), fill=palette["ink"])
        save(included, "03-whats-included-mockup.png")

    else:
        cover = Image.new("RGB", CANVAS, "#F7F2E9")
        room = ImageOps.fit(artwork, (1120, 1500), method=Image.Resampling.LANCZOS, centering=(0.68, 0.5))
        cover.paste(room, (880, 0))
        draw = ImageDraw.Draw(cover)
        draw.rectangle((0, 0, 950, 1500), fill="#F7F2E9")
        draw.text((105, 105), "THE RENOVATION EDIT", font=font(23, True), fill=palette["accent"])
        draw_wrapped(draw, (95, 175), title, width=20, font_obj=font(58, True), fill=palette["ink"], spacing=4)
        draw_wrapped(draw, (103, 395), subtitle, width=41, font_obj=font(25), fill=palette["ink"], spacing=7)
        draw.line((105, 570, 770, 570), fill=palette["accent"], width=5)
        draw.text((105, 605), "A ROOM-BY-ROOM PLANNING KIT", font=font(20, True), fill=palette["secondary"])
        frame = device_frame(previews[0], "PROJECT OVERVIEW", palette)
        cover.paste(frame, (85, 780), frame)
        if calculator_preview and calculator_preview.is_file():
            calc_frame = device_frame(fit_preview(calculator_preview, (650, 366)), "INTERACTIVE HTML TOOL", palette)
            cover.paste(calc_frame, (1130, 775), calc_frame)
        draw = ImageDraw.Draw(cover)
        draw_wrapped(draw, (105, 1350), "Includes: " + "  /  ".join(item["label"] for item in deliverables[:3]), width=62, font_obj=font(17, True), fill=palette["ink"])
        save(cover, "00-product-cover.png")

        spread = artwork.copy()
        draw = ImageDraw.Draw(spread, "RGBA")
        draw.rectangle((0, 0, 2000, 255), fill=(247, 242, 233, 238))
        draw.text((100, 55), "01  /  A WORKING EDITORIAL SPREAD", font=font(23, True), fill=palette["accent"])
        draw_wrapped(draw, (95, 105), "See the room. See the numbers.", width=40, font_obj=font(48, True), fill=palette["ink"])
        spread.paste(previews[0], (100, 450))
        draw = ImageDraw.Draw(spread)
        draw.rounded_rectangle((1000, 970, 1860, 1260), radius=10, fill="#F7F2E9", outline=palette["accent"], width=3)
        draw.text((1050, 1015), "SCOPE  /  COST  /  NEXT STEP", font=font(22, True), fill=palette["accent"])
        draw_wrapped(draw, (1050, 1070), scope[0] if scope else "Keep quotes, decisions and project costs connected.", width=45, font_obj=font(24, True), fill=palette["ink"])
        save(spread, "01-editorial-spread-mockup.png")

        board = Image.new("RGB", CANVAS, "#F7F2E9")
        board.paste(ImageOps.fit(artwork, (890, 1500), method=Image.Resampling.LANCZOS), (1110, 0))
        draw = ImageDraw.Draw(board)
        draw.text((95, 90), "02  /  MATERIAL & COST BOARD", font=font(23, True), fill=palette["accent"])
        draw_wrapped(draw, (90, 150), "Choose finishes with the budget in view", width=31, font_obj=font(48, True), fill=palette["ink"])
        for i, color in enumerate((palette["accent"], palette["secondary"], palette["highlight"], palette["ink"])):
            x = 110 + i * 190
            draw.rounded_rectangle((x, 395, x + 150, 515), radius=7, fill=color)
        if len(previews) > 1:
            board.paste(previews[1], (100, 610))
        draw = ImageDraw.Draw(board)
        draw.text((105, 1085), "COMPARE THE QUOTE", font=font(21, True), fill=palette["accent"])
        draw_wrapped(draw, (100, 1140), scope[1] if len(scope) > 1 else "Keep scope and cost beside the next decision.", width=39, font_obj=font(28, True), fill=palette["ink"])
        save(board, "02-material-cost-board-mockup.png")

        project = Image.new("RGB", CANVAS, palette["background"])
        project.paste(ImageOps.fit(artwork, (650, 1500), method=Image.Resampling.LANCZOS), (0, 0))
        draw = ImageDraw.Draw(project)
        draw.rectangle((650, 0, 2000, 1500), fill="#F7F2E9")
        draw.text((720, 90), "03  /  THE PROJECT WALL", font=font(23, True), fill=palette["accent"])
        draw_wrapped(draw, (710, 150), "One clear next action at a time", width=34, font_obj=font(47, True), fill=palette["ink"])
        for i, preview in enumerate(previews[:2]):
            y = 390 + i * 440
            project.paste(preview.resize((760, 430)), (730, y))
        save(project, "03-whats-included-mockup.png")

        stages = Image.new("RGB", CANVAS, palette["ink"])
        stages.paste(ImageOps.fit(artwork, (770, 1500), method=Image.Resampling.LANCZOS), (1230, 0))
        draw = ImageDraw.Draw(stages)
        draw.text((100, 95), "04  /  FROM IDEA TO HANDOVER", font=font(23, True), fill=palette["secondary"])
        draw_wrapped(draw, (95, 155), "A renovation plan that keeps decisions connected", width=34, font_obj=font(46, True), fill=WHITE)
        for i, label in enumerate(("Define the scope", "Compare contractors", "Track each cost", "Review what remains")):
            y = 480 + i * 205
            draw.rectangle((115, y, 175, y + 62), fill=palette["accent"] if i % 2 == 0 else palette["secondary"])
            draw.text((134, y + 14), str(i + 1), font=font(23, True), fill=WHITE)
            draw.text((215, y + 12), label, font=font(28, True), fill=WHITE)
        save(stages, "04-renovation-sequence-mockup.png")

    # These deliverable mockups retain required names but inherit the selected art direction.
    templates = Image.new("RGB", CANVAS, palette["background"])
    if layout_id == "moving_atlas":
        templates.paste(ImageOps.fit(artwork, (560, 1500), method=Image.Resampling.LANCZOS), (0, 0))
        panel_x = 560
        heading = "TWO COPIES FOR THE JOURNEY"
    elif layout_id == "aquarium_fieldguide":
        templates.paste(ImageOps.fit(artwork, (560, 1500), method=Image.Resampling.LANCZOS), (0, 0))
        panel_x = 560
        heading = "WORKED EXAMPLE + CLEAN TANK PLAN"
    elif layout_id == "puppy_firstyear_editorial":
        templates.paste(ImageOps.fit(artwork, (560, 1500), method=Image.Resampling.LANCZOS, centering=(0.5, 0.75)), (0, 0))
        panel_x = 560
        heading = "EXAMPLE PLAN + CLEAN COPY"
    else:
        templates.paste(ImageOps.fit(artwork, (760, 1500), method=Image.Resampling.LANCZOS), (1240, 0))
        panel_x = 0
        heading = "TWO EDITIONS, ONE PLAN"
    draw = ImageDraw.Draw(templates)
    draw.rectangle((panel_x, 0, panel_x + (1440 if panel_x == 560 else 1240), 1500), fill=palette["background"])
    draw.text((panel_x + 75, 100), heading, font=font(23, True), fill=palette["accent"])
    draw_wrapped(draw, (panel_x + 70, 160), "Choose the starting point that suits you", width=34, font_obj=font(46, True), fill=palette["ink"])
    workbooks = [item for item in deliverables if item["kind"] == "workbook"][:2]
    for i, item in enumerate(workbooks):
        y = 440 + i * 450
        preview = previews[min(i, len(previews) - 1)]
        templates.paste(preview, (panel_x + 75, y))
        draw = ImageDraw.Draw(templates)
        draw.text((panel_x + 80, y + 335), item["label"], font=font(25, True), fill=palette["ink"])
        draw.text((panel_x + 80, y + 375), item["name"][:64], font=font(16), fill=palette["ink"])
    save(templates, "05-excel-templates-mockup.png")

    pdf_items = [item for item in deliverables if item["kind"] in {"ebook", "pdf"}]
    if pdf_items:
        resources = Image.new("RGB", CANVAS, palette["background"])
        resources.paste(ImageOps.fit(artwork, (700, 1500), method=Image.Resampling.LANCZOS), (1300, 0))
        draw = ImageDraw.Draw(resources)
        reading_heading = "05  /  THE FIELD GUIDE SHELF" if layout_id == "aquarium_fieldguide" else "05  /  THE FIRST-YEAR READING SHELF" if layout_id == "puppy_firstyear_editorial" else "05  /  THE WORKSHOP FIELD GUIDE" if layout_id == "garage_service_manual" else "05  /  THE READING ROOM"
        draw.text((95, 100), reading_heading, font=font(23, True), fill=palette["accent"])
        draw_wrapped(draw, (90, 160), "Guides and ideas for the work ahead", width=33, font_obj=font(46, True), fill=palette["ink"])
        for i, item in enumerate(pdf_items[:4]):
            y = 400 + i * 245
            card_radius = 4 if layout_id in {"interiors_editorial", "puppy_firstyear_editorial"} else 28
            draw.rounded_rectangle((110, y, 1180, y + 190), radius=card_radius, fill=WHITE, outline=palette["secondary"], width=3)
            draw.text((145, y + 30), f"PDF  /  {i + 1:02d}", font=font(17, True), fill=palette["accent"])
            draw_wrapped(draw, (145, y + 67), item["label"], width=45, font_obj=font(25, True), fill=palette["ink"])
            draw.text((145, y + 138), item["name"][:68], font=font(15), fill=palette["ink"])
        save(resources, "06-pdf-resources-mockup.png")

    calculator_path = render_html_calculator_mockup(spec, package, manifest, identity, calculator_preview) if calculator_preview else None
    inventory_path = render_deliverables_inventory(spec, package, manifest, identity)
    render_ebook_highlight_mockup(spec, package, manifest, identity)
    extras = [path for path in (inventory_path, calculator_path) if path]
    return [paths[0], *extras, *[path for path in paths[1:] if path not in extras]]


def capture_ebook_cover(package: Path, manifest: dict[str, Any]) -> Path | None:
    """Rasterize the actual first page of the included e-book for its mockup."""
    ebook_ref = str(manifest.get("ebook_file") or "").strip()
    ebook_path = package / ebook_ref if ebook_ref else None
    if not ebook_path or not ebook_path.is_file():
        return None
    import pypdfium2 as pdfium

    document = pdfium.PdfDocument(str(ebook_path))
    try:
        if len(document) < 1:
            raise ValueError(f"The e-book has no cover page: {ebook_path}")
        capture = document[0].render(scale=1.6).to_pil().convert("RGB")
        output = package / "marketing_assets" / "ebook_cover_first_page.png"
        output.parent.mkdir(parents=True, exist_ok=True)
        capture.save(output, optimize=True)
    finally:
        document.close()
    return output


def render_ebook_highlight_mockup(
    spec: dict[str, Any], package: Path, manifest: dict[str, Any], identity: dict[str, Any],
) -> Path | None:
    """Give the real e-book cover its own prominent, product-styled marketing image."""
    ebook_ref = str(manifest.get("ebook_file") or "").strip()
    if not ebook_ref:
        return None
    palette = identity["palette"]
    copy = inventory_copy(spec)
    capture_ref = str(manifest.get("ebook_cover_capture_file") or "").strip()
    capture_path = package / capture_ref if capture_ref else None
    if not capture_path or not capture_path.is_file():
        # Pure-renderer tests and legacy manifests can use the declared cover art;
        # the production factory always writes the first-page capture above.
        fallback = str(manifest.get("ebook_cover_image") or "").strip()
        capture_path = package / fallback if fallback else None
    if not capture_path or not capture_path.is_file():
        raise FileNotFoundError(f"E-book cover capture is missing for {ebook_ref}")

    layout_id = str(identity.get("layout_id", ""))
    canvas = Image.new("RGB", CANVAS, palette["background"])
    art_ref = str(identity.get("artwork_file") or "")
    art_path = package / art_ref if art_ref else None
    if art_path and art_path.is_file():
        artwork = ImageOps.fit(Image.open(art_path).convert("RGB"), CANVAS, method=Image.Resampling.LANCZOS)
        if layout_id == "moving_atlas":
            canvas = Image.alpha_composite(artwork.convert("RGBA"), Image.new("RGBA", CANVAS, (248, 247, 235, 218))).convert("RGB")
        elif layout_id == "interiors_editorial":
            canvas.paste(ImageOps.fit(artwork, (620, 1500), method=Image.Resampling.LANCZOS, centering=(0.7, 0.5)), (1380, 0))
            ImageDraw.Draw(canvas).rectangle((0, 0, 1410, 1500), fill="#F7F2E9")
        elif layout_id == "puppy_firstyear_editorial":
            canvas.paste(ImageOps.fit(artwork, (650, 1500), method=Image.Resampling.LANCZOS, centering=(0.5, 0.76)), (1370, 0))
            ImageDraw.Draw(canvas).rectangle((0, 0, 1400, 1500), fill=palette["background"])

    draw = ImageDraw.Draw(canvas)
    draw.text((90, 80), copy["ebook_heading"], font=font(24, True), fill=palette["accent"])
    ebook_title = str(manifest.get("ebook_title") or Path(ebook_ref).stem)
    draw_wrapped(draw, (85, 135), ebook_title, width=35, font_obj=font(46, True), fill=palette["ink"], spacing=3)
    draw.text((95, 265), copy["ebook_subheading"], font=font(18, True), fill=palette["ink"])

    cover_box = (115, 350, 770, 1360)
    draw.rounded_rectangle((cover_box[0] + 24, cover_box[1] + 28, cover_box[2] + 24, cover_box[3] + 28), radius=18, fill="#A69B8D")
    draw.rounded_rectangle(cover_box, radius=18 if layout_id == "moving_atlas" else 5, fill=WHITE, outline=palette["accent"], width=4)
    with Image.open(capture_path) as source:
        cover_page = ImageOps.contain(source.convert("RGB"), (cover_box[2] - cover_box[0] - 34, cover_box[3] - cover_box[1] - 34), method=Image.Resampling.LANCZOS)
    canvas.paste(cover_page, (cover_box[0] + (cover_box[2] - cover_box[0] - cover_page.width) // 2, cover_box[1] + 17))

    pdf_items = [item for item in package_deliverables(manifest, package, spec) if item["kind"] in {"ebook", "pdf"} and item["name"] != Path(ebook_ref).name]
    draw = ImageDraw.Draw(canvas)
    draw.rounded_rectangle((850, 370, 1855, 1325), radius=24 if layout_id == "moving_atlas" else 5, fill=WHITE, outline=palette["secondary"], width=3)
    draw.text((900, 420), copy["also_in_kit"], font=font(21, True), fill=palette["accent"])
    for index, item in enumerate(pdf_items[:4]):
        y = 510 + index * 180
        draw.rounded_rectangle((900, y, 1800, y + 145), radius=18 if layout_id == "moving_atlas" else 4, fill=palette["background"])
        draw.rounded_rectangle((925, y + 28, 1005, y + 112), radius=16 if layout_id == "moving_atlas" else 3, fill=palette["secondary"])
        draw.text((944, y + 55), "PDF", font=font(17, True), fill=WHITE)
        draw_wrapped(draw, (1035, y + 22), inventory_item_label(item, copy), width=39, font_obj=font(20, True), fill=palette["ink"], spacing=3)
        draw_wrapped(draw, (1035, y + 80), item["name"], width=43, font_obj=font(14), fill=palette["ink"], spacing=2)

    draw_identity_signature(ImageDraw.Draw(canvas), identity)
    output = package / "marketing_mockups" / "06-pdf-resources-mockup.png"
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output, optimize=True)
    return output


def render_deliverables_inventory(
    spec: dict[str, Any], package: Path, manifest: dict[str, Any], identity: dict[str, Any],
) -> Path:
    """Render a complete, manifest-derived inventory of every customer file."""
    palette = identity["palette"]
    items = package_deliverables(manifest, package, spec)
    layout_id = str(identity.get("layout_id", ""))
    artwork_ref = str(identity.get("artwork_file") or "")
    artwork_path = package / artwork_ref if artwork_ref else None
    canvas = Image.new("RGB", CANVAS, palette["background"])
    draw = ImageDraw.Draw(canvas)

    if artwork_path and artwork_path.is_file():
        artwork = ImageOps.fit(Image.open(artwork_path).convert("RGB"), CANVAS, method=Image.Resampling.LANCZOS)
        if layout_id == "moving_atlas":
            canvas = artwork
            overlay = Image.new("RGBA", CANVAS, (248, 247, 235, 231))
            canvas = Image.alpha_composite(canvas.convert("RGBA"), overlay).convert("RGB")
            draw = ImageDraw.Draw(canvas)
        elif layout_id == "interiors_editorial":
            canvas.paste(ImageOps.fit(artwork, (520, 1500), method=Image.Resampling.LANCZOS, centering=(0.76, 0.5)), (1480, 0))
            draw.rectangle((0, 0, 1510, 1500), fill="#F7F2E9")
            draw = ImageDraw.Draw(canvas)
        elif layout_id == "puppy_firstyear_editorial":
            canvas.paste(ImageOps.fit(artwork, (560, 1500), method=Image.Resampling.LANCZOS, centering=(0.5, 0.76)), (1440, 0))
            draw.rectangle((0, 0, 1470, 1500), fill=palette["background"])
            draw = ImageDraw.Draw(canvas)

    archive_name = manifest.get("digital_kit_zip") or manifest.get("local_file")
    archive_path = package / str(archive_name or "")
    archive_count = 1 if archive_name and archive_path.is_file() else 0
    copy = inventory_copy(spec)
    totals: dict[str, int] = {copy["excel"]: 0, copy["pdf"]: 0, copy["guide_plural"]: 0, copy["html"]: 0, copy["other"]: 0}
    for item in items:
        kind = str(item.get("kind", "")).casefold()
        if kind == "workbook":
            totals[copy["excel"]] += 1
        elif kind in {"pdf", "ebook"}:
            totals[copy["pdf"]] += 1
        elif kind == "guide":
            totals[copy["guide_plural"]] += 1
        elif kind == "html":
            totals[copy["html"]] += 1
        else:
            totals[copy["other"]] += 1

    draw.text((82, 72), copy["heading"], font=font(25, True), fill=palette["accent"])
    draw_wrapped(draw, (78, 125), str(spec.get("title", "Digital kit")), width=39, font_obj=font(48, True), fill=palette["ink"], spacing=2)
    guide_count = totals[copy["guide_plural"]]
    summary = [
        f"{len(items)} {copy['files']}",
        f"{totals[copy['excel']]} {copy['excel_plural']}",
        f"{totals[copy['pdf']]} {copy['pdf_plural']}",
        f"{guide_count} {copy['guide_singular'] if guide_count == 1 else copy['guide_plural']}",
    ]
    if totals[copy["html"]]:
        summary.append(f"{totals[copy['html']]} {copy['html_plural']}")
    if archive_count:
        summary.append(f"{archive_count} {copy['zip']}")
    x = 82
    for label in summary:
        width = max(152, len(label) * 12 + 35)
        draw.rounded_rectangle((x, 242, x + width, 290), radius=18 if layout_id == "moving_atlas" else 4, fill=palette["ink"])
        draw.text((x + 16, 255), label, font=font(15, True), fill=WHITE)
        x += width + 16

    if not items:
        draw.rounded_rectangle((80, 360, 1915, 1310), radius=24, fill=WHITE, outline=palette["accent"], width=3)
        draw.text((125, 420), copy["empty"], font=font(30, True), fill=palette["ink"])
    else:
        cols = 4 if len(items) >= 10 else 3
        rows = (len(items) + cols - 1) // cols
        gap_x, gap_y = 24, 20
        left, top, right, bottom = 80, 340, 1920, 1395
        card_w = (right - left - (cols - 1) * gap_x) // cols
        card_h = min(430, (bottom - top - (rows - 1) * gap_y) // rows)
        icon_colors = {"workbook": palette["accent"], "ebook": palette["secondary"], "pdf": palette["secondary"], "guide": palette["ink"], "html": palette["highlight"]}
        icon_labels = {"workbook": "XLS", "ebook": "PDF", "pdf": "PDF", "guide": copy["guide_icon"], "html": "HTML"}
        for index, item in enumerate(items):
            row, col = divmod(index, cols)
            x0 = left + col * (card_w + gap_x)
            y0 = top + row * (card_h + gap_y)
            radius = 26 if layout_id == "moving_atlas" else 3 if layout_id == "puppy_firstyear_editorial" else 6
            draw.rounded_rectangle((x0, y0, x0 + card_w, y0 + card_h), radius=radius, fill=WHITE, outline=palette["secondary"], width=2)
            kind = str(item.get("kind", "other")).casefold()
            color = icon_colors.get(kind, palette["highlight"])
            icon = icon_labels.get(kind, "FILE")
            icon_radius = 18 if layout_id == "moving_atlas" else 3 if layout_id == "puppy_firstyear_editorial" else 4
            draw.rounded_rectangle((x0 + 24, y0 + 24, x0 + 114, y0 + 95), radius=icon_radius, fill=color)
            draw.text((x0 + 37, y0 + 47), icon, font=font(17, True), fill=WHITE)
            label_y = y0 + 120
            draw_wrapped(draw, (x0 + 25, label_y), inventory_item_label(item, copy), width=max(20, card_w // 17), font_obj=font(21, True), fill=palette["ink"], spacing=4)
            filename_y = y0 + min(card_h - 60, 230)
            draw_wrapped(draw, (x0 + 25, filename_y), str(item.get("name", "")), width=max(24, card_w // 14), font_obj=font(15), fill=palette["ink"], spacing=3)
            draw.rectangle((x0 + 25, y0 + card_h - 18, x0 + card_w - 25, y0 + card_h - 10), fill=color)

    draw_identity_signature(ImageDraw.Draw(canvas), identity)
    path = package / "marketing_mockups" / "03-whats-included-mockup.png"
    path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(path, optimize=True)
    return path


def render_html_calculator_mockup(
    spec: dict[str, Any], package: Path, manifest: dict[str, Any],
    identity: dict[str, Any], calculator_preview: Path | None,
) -> Path:
    """Give the actual browser-rendered HTML tool its own high-visibility mockup."""
    if not calculator_preview or not calculator_preview.is_file():
        raise FileNotFoundError("The HTML tool mockup requires its captured browser preview")
    palette = identity["palette"]
    layout_id = str(identity.get("layout_id", ""))
    copy = {
        "es": ("LA HERRAMIENTA INTERACTIVA", "Una utilidad práctica junto al Excel", "ABRE · INTRODUCE · REVISA", "Funciona offline en el navegador"),
        "en": ("THE INTERACTIVE HTML TOOL", "A practical companion to the Excel planner", "OPEN · ENTER · REVIEW", "Works offline in your browser"),
    }
    language = "es" if str(spec.get("design_language", "en")).casefold().startswith("es") else "en"
    heading, subtitle, steps, privacy = copy[language]
    tool = manifest.get("interactive_html_tool") or {}
    tool_title = str(tool.get("title") or Path(str(tool.get("filename") or "interactive-calculator.html")).stem.replace("-", " ").title())
    canvas = Image.new("RGB", CANVAS, palette["background"])
    artwork_ref = str(identity.get("artwork_file") or "")
    artwork_path = package / artwork_ref if artwork_ref else None
    if artwork_path and artwork_path.is_file():
        art = ImageOps.fit(Image.open(artwork_path).convert("RGB"), CANVAS, method=Image.Resampling.LANCZOS)
        if layout_id == "aquarium_fieldguide":
            canvas = art
            overlay = Image.new("RGBA", CANVAS, (242, 245, 235, 218))
            canvas = Image.alpha_composite(canvas.convert("RGBA"), overlay).convert("RGB")
        elif layout_id == "moving_atlas":
            canvas = art
            overlay = Image.new("RGBA", CANVAS, (248, 247, 235, 228))
            canvas = Image.alpha_composite(canvas.convert("RGBA"), overlay).convert("RGB")
        elif layout_id == "puppy_firstyear_editorial":
            canvas.paste(ImageOps.fit(art, (700, 1500), method=Image.Resampling.LANCZOS, centering=(0.5, 0.76)), (1300, 0))
            ImageDraw.Draw(canvas).rectangle((0, 0, 1330, 1500), fill=palette["background"])
    draw = ImageDraw.Draw(canvas)
    draw.text((95, 85), heading, font=font(24, True), fill=palette["accent"])
    draw_wrapped(draw, (90, 145), tool_title, width=32, font_obj=font(51, True), fill=palette["ink"])
    draw.text((100, 285), subtitle, font=font(24), fill=palette["ink"])
    frame = device_frame(fit_preview(calculator_preview, (1500, 844)), tool_title.upper(), palette)
    canvas.paste(frame, (220, 390), frame)
    draw = ImageDraw.Draw(canvas)
    draw.rounded_rectangle((260, 1300, 1740, 1400), radius=22 if layout_id != "interiors_editorial" else 5, fill=palette["ink"])
    draw.text((315, 1332), steps, font=font(22, True), fill=WHITE)
    draw.text((100, 1440), privacy, font=font(16), fill=palette["ink"])
    draw_identity_signature(ImageDraw.Draw(canvas), identity)
    output = package / "marketing_mockups" / "07-interactive-html-calculator-mockup.png"
    output.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output, optimize=True)
    return output


def inventory_copy(spec: dict[str, Any]) -> dict[str, str]:
    """Choose inventory labels in the same language declared for the workbook."""
    language = str(spec.get("design_language") or spec.get("language") or "en").strip().casefold()
    if language.startswith("es"):
        return {
            "heading": "TODO LO QUE INCLUYE", "files": "ARCHIVOS", "excel": "EXCEL",
            "excel_plural": "LIBROS EXCEL", "pdf": "PDF", "pdf_plural": "PDF",
            "guide_singular": "GUÍA", "guide_plural": "GUÍAS", "other": "OTROS",
            "zip": "ZIP DESCARGABLE", "guide_icon": "GUÍA",
            "empty": "No se han encontrado archivos declarados para mostrar.",
            "ebook_heading": "EL E-BOOK, POR DENTRO", "ebook_subheading": "Una guía práctica incluida en el kit",
            "also_in_kit": "TAMBIÉN INCLUIDO EN EL KIT",
            "html": "HTML", "html_plural": "HERRAMIENTAS HTML",
            "workbook_example": "Libro Excel con datos de ejemplo",
            "workbook_blank": "Libro Excel en blanco", "guide": "Guía de uso de Excel (Markdown)",
            "guide_pdf": "Guía de uso de Excel (PDF)", "ebook": "Libro electrónico temático (PDF)",
            "worksheet": "Plantillas de trabajo (PDF)", "index": "Índice del kit digital (PDF)",
            "pdf_generic": "Recurso en PDF", "file_generic": "Archivo digital",
        }
    return {
        "heading": "EVERYTHING INCLUDED", "files": "FILES", "excel": "EXCEL",
        "excel_plural": "EXCEL WORKBOOKS", "pdf": "PDF", "pdf_plural": "PDF FILES",
        "guide_singular": "GUIDE", "guide_plural": "GUIDES", "other": "OTHER",
        "zip": "DOWNLOAD ZIP", "guide_icon": "GUIDE",
        "empty": "No declared files were found to display.",
        "ebook_heading": "A CLOSER LOOK AT THE E-BOOK", "ebook_subheading": "A practical guide included in the kit",
        "also_in_kit": "ALSO INCLUDED IN THE KIT",
        "html": "HTML", "html_plural": "HTML TOOLS",
        "workbook_example": "Example-data workbook", "workbook_blank": "Blank workbook",
        "guide": "Excel usage guide (Markdown)", "guide_pdf": "Excel user guide (PDF)",
        "ebook": "Topic e-book (PDF)", "worksheet": "Printable worksheets (PDF)",
        "index": "Digital kit contents index (PDF)", "pdf_generic": "PDF resource", "file_generic": "Digital file",
    }


def inventory_item_label(item: dict[str, str], copy: dict[str, str]) -> str:
    """Translate descriptive labels while keeping the delivered filename exact."""
    if copy["heading"] == "EVERYTHING INCLUDED":
        return str(item.get("label") or copy["file_generic"])
    kind = str(item.get("kind", "")).casefold()
    name = str(item.get("name", "")).casefold()
    if kind == "workbook":
        return copy["workbook_blank"] if "blank" in name else copy["workbook_example"]
    if kind == "guide":
        return copy["guide"]
    if kind == "ebook":
        return copy["ebook"]
    if kind == "html":
        return "Calculadora interactiva offline (HTML)" if copy["heading"] != "EVERYTHING INCLUDED" else "Interactive offline tool (HTML)"
    if kind == "pdf":
        if "index" in name:
            return copy["index"]
        if any(token in name for token in ("worksheet", "template", "plantilla")):
            return copy["worksheet"]
        if any(token in name for token in ("guide", "guia")):
            return copy["guide_pdf"]
        return copy["pdf_generic"]
    return copy["file_generic"]


def generate_for_manifest(manifest_path: Path, *, reuse_captures: bool = False) -> dict[str, Any]:
    manifest_path = manifest_path.resolve()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    product_id = int(manifest["product_id"])
    product_dir = ROOT / "outputs" / "listings" / f"product-{product_id}"
    spec = json.loads((product_dir / "design_spec.json").read_text(encoding="utf-8"))
    assert_unique_active_identity(product_id, spec)
    workbook_path = manifest_path.parent / str(manifest.get("workbook_file", ""))
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Listing package has no workbook for marketing previews: {workbook_path}")

    video = manifest_path.parent / "marketing_video" / f"{spec.get('slug', f'product-{product_id}')}-excel-demo.mp4"
    html_demo_video = manifest_path.parent / "marketing_video" / f"{spec.get('slug', f'product-{product_id}')}-html-calculator-demo.mp4"
    if reuse_captures:
        previous = manifest.get("marketing_assets") or {}
        if previous.get("source_sha256") != sha256(workbook_path):
            raise RuntimeError("Cannot reuse marketing captures: the workbook hash has changed")
        screenshot_refs = [str(item) for item in manifest.get("screenshot_files", [])]
        screenshots = [manifest_path.parent / item for item in screenshot_refs]
        if not screenshots or any(not path.is_file() for path in screenshots):
            raise RuntimeError("Cannot reuse marketing captures: declared workbook captures are missing")
        video_record = {
            "sha256": previous.get("video_sha256"),
            "demo_input_changed": previous.get("demo_input_changed", False),
            "demo_input_sheet": previous.get("demo_input_sheet"),
            "captured_sheets": previous.get("captured_sheets", []),
            "screenshot_files": screenshot_refs,
        }
    else:
        from record_excel_demo import record_excel_demo

        video_record = record_excel_demo(manifest_path, video)
        screenshots = [manifest_path.parent / str(item) for item in video_record["screenshot_files"]]
    if not screenshots:
        raise RuntimeError("El demo de Excel no produjo capturas de hojas")
    calculator_preview = None
    html_demo_record = None
    html_tool = manifest.get("interactive_html_tool") or {}
    if html_tool.get("passed"):
        html_tool_path = manifest_path.parent / str(html_tool.get("file") or "")
        calculator_preview = manifest_path.parent / "marketing_assets" / "calculator-preview.png"
        from record_html_calculator_demo import record_html_calculator_demo

        html_demo_record = record_html_calculator_demo(
            html_tool_path,
            calculator_preview,
            html_demo_video,
            operation=str(html_tool.get("operation") or ""),
            language=str(html_tool.get("language") or spec.get("design_language") or "en"),
        )
        manifest["html_tool_preview_file"] = calculator_preview.relative_to(manifest_path.parent).as_posix()
        manifest["html_tool_video_file"] = html_demo_video.relative_to(manifest_path.parent).as_posix()
        manifest["html_tool_demo"] = html_demo_record
    identity = marketing_identity(spec, manifest_path.parent, manifest)
    ebook_capture = capture_ebook_cover(manifest_path.parent, manifest)
    if ebook_capture:
        manifest["ebook_cover_capture_file"] = ebook_capture.relative_to(manifest_path.parent).as_posix()
        manifest["ebook_cover_capture_sha256"] = sha256(ebook_capture)
    mockups = marketing_images(spec, screenshots, manifest_path.parent, manifest, calculator_preview)

    previous_cover = manifest.get("cover_image")
    cover = mockups[0].relative_to(manifest_path.parent).as_posix()
    existing_images = [str(name) for name in manifest.get("image_files", [])]
    # Keep the paired Excel + calculator cover first, then spotlight the actual
    # HTML interface before the remaining kit inventory and screenshots.
    mockup_refs = [p.relative_to(manifest_path.parent).as_posix() for p in mockups]
    html_mockup = next((name for name in mockup_refs if Path(name).name == "07-interactive-html-calculator-mockup.png"), None)
    inventory_mockup = next((name for name in mockup_refs if Path(name).name == "03-whats-included-mockup.png"), None)
    prioritized = [cover, *[name for name in (html_mockup, inventory_mockup) if name], *[name for name in mockup_refs if name not in {cover, html_mockup, inventory_mockup}]]
    ordered = list(dict.fromkeys([*prioritized, *[p.relative_to(manifest_path.parent).as_posix() for p in screenshots], *[name for name in existing_images if name != previous_cover]]))
    ordered = ordered[:10]  # Etsy allows ten listing photos; the cover is included.
    manifest.update({
        "image_files": ordered,
        "cover_image": cover,
        "marketing_identity": identity,
        "marketing_artwork_file": str((spec.get("visual_identity") or {}).get("marketing_artwork_file") or manifest.get("marketing_artwork_file") or ""),
        "ebook_cover_image": manifest.get("ebook_cover_image") or (previous_cover if previous_cover and "cover" in Path(str(previous_cover)).name.casefold() else None),
        "mockup_files": [p.relative_to(manifest_path.parent).as_posix() for p in mockups],
        "screenshot_files": [p.relative_to(manifest_path.parent).as_posix() for p in screenshots],
        "video_file": video.relative_to(manifest_path.parent).as_posix(),
        "video_files": [video.relative_to(manifest_path.parent).as_posix(), *([html_demo_video.relative_to(manifest_path.parent).as_posix()] if html_demo_record else [])],
        "marketing_assets": {
            "status": "ready",
            "source_workbook": str(manifest.get("workbook_file", "")),
            "source_sha256": sha256(workbook_path),
            "image_count": len(ordered),
            "mockup_count": len(mockups),
            "screenshot_count": len(screenshots),
            "video_sha256": video_record["sha256"],
            "demo_input_changed": video_record["demo_input_changed"],
            "demo_input_sheet": video_record["demo_input_sheet"],
            "captured_sheets": video_record["captured_sheets"],
            "generated_from_product_design_spec": True,
            "ebook_cover_capture_file": manifest.get("ebook_cover_capture_file"),
            "ebook_cover_capture_sha256": manifest.get("ebook_cover_capture_sha256"),
            "html_tool_demo": html_demo_record,
        },
    })
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest["marketing_assets"]


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="Genera previews del Excel, mockups y vídeo de marketing para un listing.")
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--product", type=int, help="Producto cuyo paquete draft ya existe")
    group.add_argument("--manifest", type=Path, help="Manifiesto de un paquete de listing existente")
    parser.add_argument("--reuse-captures", action="store_true", help="Regenera mockups con capturas existentes solo si la huella del XLSX coincide")
    args = parser.parse_args()
    if args.manifest:
        targets = [args.manifest]
    elif args.product is not None:
        targets = [ROOT / "outputs" / "listings" / f"product-{args.product}" / "draft_package" / "listing.json"]
    else:
        from factory_config import active_product_ids

        targets = [ROOT / "outputs" / "listings" / f"product-{product_id}" / "draft_package" / "listing.json" for product_id in active_product_ids()]
    reports = [generate_for_manifest(target, reuse_captures=args.reuse_captures) for target in targets]
    for report in reports:
        print(f"MARKETING_ASSETS={report['status']} images={report['image_count']} mockups={report['mockup_count']} screenshots={report['screenshot_count']}")
