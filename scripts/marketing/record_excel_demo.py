from __future__ import annotations

import argparse
import ctypes
import json
import re
import time
from pathlib import Path

import imageio.v2 as imageio
import mss
import numpy as np
import win32com.client
import win32gui
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[2]
FPS = 20
TARGET_SIZE = (1920, 1080)


def sha256(path: Path) -> str:
    import hashlib

    digest = hashlib.sha256(path.read_bytes())
    return digest.hexdigest()


def capture_frame(sct: mss.mss, monitor: dict[str, int], label: str) -> np.ndarray:
    shot = Image.frombytes("RGB", (monitor["width"], monitor["height"]), sct.grab(monitor).rgb)
    shot = shot.resize(TARGET_SIZE, Image.Resampling.LANCZOS)
    draw = ImageDraw.Draw(shot)
    draw.rectangle((0, TARGET_SIZE[1] - 72, TARGET_SIZE[0], TARGET_SIZE[1]), fill="#172C3B")
    try:
        label_font = ImageFont.truetype("C:/Windows/Fonts/segoeuib.ttf", 23)
    except OSError:
        label_font = ImageFont.load_default()
    draw.text((38, TARGET_SIZE[1] - 52), label[:105], font=label_font, fill="white")
    return np.asarray(shot)


def save_workbook_capture(sct: mss.mss, monitor: dict[str, int], path: Path, excel, sheet) -> None:
    foreground = win32gui.GetForegroundWindow()
    foreground_title = win32gui.GetWindowText(foreground)
    if foreground != int(excel.Hwnd) and "excel" not in foreground_title.casefold():
        raise RuntimeError(f"Refusing to capture a non-Excel window: {foreground_title or 'unknown window'}")
    shot = Image.frombytes("RGB", (monitor["width"], monitor["height"]), sct.grab(monitor).rgb)
    shot = shot.crop((0, 24, shot.width, shot.height - 42)).resize(TARGET_SIZE, Image.Resampling.LANCZOS)
    path.parent.mkdir(parents=True, exist_ok=True)
    shot.save(path, format="PNG", optimize=True)


def record_for(writer, sct: mss.mss, monitor: dict[str, int], seconds: float, label: str) -> None:
    frame = capture_frame(sct, monitor, label)
    for _ in range(round(seconds * FPS)):
        writer.append_data(frame)


def activate(excel, workbook, sheet_name: str, cell: str = "A1", column_count: int = 10) -> None:
    try:
        win32gui.ShowWindow(int(excel.Hwnd), 3)
        win32gui.SetForegroundWindow(int(excel.Hwnd))
    except Exception:
        pass
    sheet = workbook.Worksheets(sheet_name)
    sheet.Activate()
    sheet.Range(cell).Select()
    excel.ActiveWindow.ScrollRow = 1
    excel.ActiveWindow.ScrollColumn = 1
    excel.ActiveWindow.Zoom = max(85, min(140, int(1600 / (max(1, column_count) * 80 * 1.333) * 100)))
    time.sleep(1.5)
    foreground = win32gui.GetForegroundWindow()
    foreground_title = win32gui.GetWindowText(foreground)
    if foreground != int(excel.Hwnd) and "excel" not in foreground_title.casefold():
        raise RuntimeError(f"Excel could not be brought to the foreground; refusing to capture {foreground_title or 'unknown window'}")


def monitor_for_window(sct: mss.mss, hwnd: int) -> dict[str, int]:
    """Choose the display containing Excel, not whichever screen happens to be primary."""
    left, top, right, bottom = win32gui.GetWindowRect(hwnd)
    center_x, center_y = (left + right) // 2, (top + bottom) // 2
    monitors = sct.monitors[1:]
    for monitor in monitors:
        if (monitor["left"] <= center_x < monitor["left"] + monitor["width"]
                and monitor["top"] <= center_y < monitor["top"] + monitor["height"]):
            return monitor
    if not monitors:
        raise RuntimeError("No display is available for the Excel demo capture")
    return max(monitors, key=lambda item: item["width"] * item["height"])


def choose_demo_input(workbook_path: Path, spec: dict) -> tuple[str, str, float] | None:
    """Find an amount input declared in the contract; never guesses an arbitrary cell."""
    book = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        for table in spec.get("model_contract", {}).get("tables", []):
            headers = table.get("headers", [])
            if table.get("sheet") in {"Settings", "Start here", "Dashboard"}:
                continue
            min_col, header_row, _, _ = range_boundaries(table["ref"])
            for index, header in enumerate(headers):
                if header.get("role") != "input" or header.get("kind") not in {"amount", "number", "count"}:
                    continue
                row = header_row + 1
                col = min_col + index
                sheet_name = str(table["sheet"])
                if sheet_name not in book.sheetnames:
                    continue
                value = book[sheet_name].cell(row, col).value
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    address = book[sheet_name].cell(row, col).coordinate
                    increment = max(1.0, round(abs(float(value)) * 0.05, 2))
                    return sheet_name, address, float(value) + increment
    finally:
        book.close()
    return None


def record_excel_demo(manifest_path: Path, output_path: Path) -> dict:
    manifest_path = manifest_path.resolve()
    output_path = output_path.resolve()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    package = manifest_path.parent
    workbook_path = package / str(manifest.get("workbook_file", ""))
    if not workbook_path.is_file():
        raise FileNotFoundError(f"No se encuentra el XLSX de muestra: {workbook_path}")
    spec_path = ROOT / "outputs" / "listings" / f"product-{manifest['product_id']}" / "design_spec.json"
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    working_copy = output_path.parent / f"{workbook_path.stem}-demo-working-copy{workbook_path.suffix}"
    working_copy.parent.mkdir(parents=True, exist_ok=True)
    import shutil

    shutil.copy2(workbook_path, working_copy)
    excel = None
    book = None
    console = 0
    screen = None
    capture_paths: list[Path] = []
    changed = choose_demo_input(workbook_path, spec)
    try:
        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = True
        excel.DisplayAlerts = False
        excel.WindowState = -4137
        book = excel.Workbooks.Open(str(working_copy), ReadOnly=False, UpdateLinks=0)
        book.Activate()
        available_sheets = {str(sheet.Name) for sheet in book.Worksheets}
        column_counts = {
            str(table.get("sheet")): len(table.get("headers", []))
            for table in spec.get("model_contract", {}).get("tables", [])
        }
        try:
            win32gui.ShowWindow(int(excel.Hwnd), 3)
            win32gui.SetForegroundWindow(int(excel.Hwnd))
            excel.DisplayFullScreen = True
        except Exception:
            pass
        console = ctypes.windll.kernel32.GetConsoleWindow()
        if console:
            ctypes.windll.user32.ShowWindow(console, 0)
        excel.CalculateFull()
        time.sleep(1.5)
        with mss.mss() as sct:
            screen = monitor_for_window(sct, int(excel.Hwnd))
            with imageio.get_writer(output_path, fps=FPS, codec="libx264", quality=7, pixelformat="yuv420p", macro_block_size=1) as writer:
                for sheet_name in [name for name in spec.get("tabs", []) if name in available_sheets]:
                    if sheet_name in {"Start here", "Settings"}:
                        continue
                    activate(excel, book, sheet_name, "B3", column_counts.get(sheet_name, 10))
                    safe_name = re.sub(r"[^A-Za-z0-9]+", "-", sheet_name).strip("-").lower() or "sheet"
                    capture_path = package / "source_captures" / f"{len(capture_paths) + 1:02d}-{safe_name}.png"
                    save_workbook_capture(sct, screen, capture_path, excel, book.Worksheets(sheet_name))
                    capture_paths.append(capture_path)
                    record_for(writer, sct, screen, 1.4, f"{spec.get('title', 'Excel planner')} · {sheet_name}")
                if changed:
                    sheet_name, address, value = changed
                    activate(excel, book, sheet_name, address, column_counts.get(sheet_name, 10))
                    book.Worksheets(sheet_name).Range(address).Value2 = value
                    excel.CalculateFull()
                    record_for(writer, sct, screen, 1.8, f"Example input updated in {sheet_name} · formulas recalculate")
                    dashboard = "Dashboard" if "Dashboard" in book.Worksheets else sheet_name
                    activate(excel, book, dashboard, "B3", column_counts.get(dashboard, 10))
                    record_for(writer, sct, screen, 2.4, f"Updated overview · {dashboard}")
                else:
                    record_for(writer, sct, screen, 1.8, "Navigate the workbook sheets and review the connected views")
        book.Close(SaveChanges=False)
        book = None
    finally:
        if book is not None:
            book.Close(SaveChanges=False)
        if excel is not None:
            excel.DisplayAlerts = True
            excel.Quit()
        if console:
            ctypes.windll.user32.ShowWindow(console, 5)
        if working_copy.exists():
            working_copy.unlink()
    if not output_path.is_file() or output_path.stat().st_size < 100_000:
        raise RuntimeError("El vídeo demo no se ha generado correctamente")
    return {
        "file": output_path.relative_to(package).as_posix(),
        "sha256": sha256(output_path),
        "working_copy_used": True,
        "demo_input_changed": bool(changed),
        "demo_input_sheet": changed[0] if changed else None,
        "captured_sheets": [name for name in spec.get("tabs", []) if name in available_sheets and name not in {"Start here", "Settings"}],
        "screenshot_files": [path.relative_to(package).as_posix() for path in capture_paths],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Graba una demostración reutilizable del XLSX de un producto.")
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    manifest_path = args.manifest.resolve()
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    output = args.output or (manifest_path.parent / "marketing_video" / f"product-{manifest['product_id']}-excel-demo.mp4")
    result = record_excel_demo(manifest_path, output)
    print(f"EXCEL_DEMO={result['file']} changed_input={result['demo_input_changed']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
