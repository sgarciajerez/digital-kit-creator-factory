from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Border
from PIL import Image, ImageDraw, ImageOps


ROOT = Path(__file__).resolve().parent.parent
PRODUCT = ROOT / "outputs" / "listings" / "product-7"
BOOK = PRODUCT / "autopilot_3_Planificador_de_presupuesto_familiar_0.1.0.xlsx"
OVERLAY_DIR = PRODUCT / "visual_overlays"
NAVY = (20, 43, 69)
WHITE = (255, 255, 255)


def excel_width_to_pixels(width: float) -> int:
    return max(1, int(width * 7.0))


def row_height_to_pixels(height: float | None) -> int:
    return max(1, int((height or 15.0) * 1.333))


def range_size(ws, start_col: str, end_col: str, start_row: int, end_row: int) -> tuple[int, int]:
    width = sum(
        excel_width_to_pixels(ws.column_dimensions[chr(code)].width or 13.0)
        for code in range(ord(start_col), ord(end_col) + 1)
    )
    height = sum(row_height_to_pixels(ws.row_dimensions[row].height) for row in range(start_row, end_row + 1))
    return width, height


def rounded_overlay(path: Path, width: int, height: int, radius: int = 24) -> None:
    scale = 4
    canvas = (max(20, width * scale), max(20, height * scale))
    radius = min(radius * scale, canvas[0] // 3, canvas[1] // 3)
    shape_mask = Image.new("L", canvas, 0)
    mask_draw = ImageDraw.Draw(shape_mask)
    mask_draw.rounded_rectangle((2, 2, canvas[0] - 3, canvas[1] - 3), radius=radius, fill=255)

    overlay = Image.new("RGBA", canvas, (0, 0, 0, 0))
    outside = ImageOps.invert(shape_mask)
    white_layer = Image.new("RGBA", canvas, WHITE + (255,))
    overlay.paste(white_layer, (0, 0), outside)
    draw = ImageDraw.Draw(overlay)
    draw.rounded_rectangle(
        (2, 2, canvas[0] - 3, canvas[1] - 3),
        radius=radius,
        outline=NAVY + (255,),
        width=max(3, scale * 2),
    )
    overlay.resize((width, height), Image.Resampling.LANCZOS).save(path, format="PNG", optimize=True)


def add_overlay(ws, name: str, start_col: str, end_col: str, start_row: int, end_row: int) -> None:
    width, height = range_size(ws, start_col, end_col, start_row, end_row)
    output = OVERLAY_DIR / f"{name}.png"
    rounded_overlay(output, width, height)
    image = XLImage(str(output))
    image.width = width
    image.height = height
    ws.add_image(image, f"{start_col}{start_row}")


def clear_card_borders(ws, ranges: list[tuple[str, str, int, int]]) -> None:
    for start_col, end_col, start_row, end_row in ranges:
        for row in range(start_row, end_row + 1):
            for code in range(ord(start_col), ord(end_col) + 1):
                ws.cell(row, code - ord("A") + 1).border = Border()


def main() -> int:
    if not BOOK.exists():
        raise SystemExit(f"No existe el libro: {BOOK}")
    OVERLAY_DIR.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(BOOK)

    instructions = workbook["Instrucciones"]
    instruction_cards = [("A", "C", 5, 10), ("E", "G", 5, 10), ("I", "K", 5, 10)]
    clear_card_borders(instructions, instruction_cards)
    for index, (start, end, row_start, row_end) in enumerate(instruction_cards, 1):
        add_overlay(instructions, f"instructions_{index}", start, end, row_start, row_end)

    preview = workbook["Vista previa"]
    preview_cards = [("A", "B", 5, 7), ("D", "E", 5, 7), ("G", "H", 5, 7)]
    clear_card_borders(preview, preview_cards)
    for index, (start, end, row_start, row_end) in enumerate(preview_cards, 1):
        add_overlay(preview, f"preview_{index}", start, end, row_start, row_end)

    workbook.save(BOOK)
    print(f"Overlays applied: {BOOK}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
