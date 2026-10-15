from __future__ import annotations

"""Reusable Excel primitives for simple, formula-safe workbook templates."""

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from design_architecture import ROOT
from model_quality import number_format_for_header, validate_headers


LIBRARY_PATH = ROOT / "config" / "EXCEL_TEMPLATE_LIBRARY.json"


@dataclass(frozen=True)
class ExcelTheme:
    background: str
    ink: str
    accent: str
    secondary: str
    highlight: str
    surface: str = "FFFFFF"
    muted: str = "6D7D89"
    line: str = "DCE7E3"
    input_fill: str = "FFF9E9"
    formula_fill: str = "F4F8F6"

    @classmethod
    def from_spec(cls, spec: dict[str, Any]) -> "ExcelTheme":
        palette = spec.get("palette", {})
        return cls(**{key: str(palette.get(key, value)).replace("#", "").upper() for key, value in {
            "background": "FAF7F2",
            "ink": "263A35",
            "accent": "477A67",
            "secondary": "7A9B89",
            "highlight": "D0A542",
        }.items()})


class ExcelTemplateFactory:
    """Compile screen content into a small set of native worksheet primitives."""

    def __init__(self, workbook, spec: dict[str, Any]):
        self.workbook = workbook
        self.spec = spec
        self.library = json.loads(LIBRARY_PATH.read_text(encoding="utf-8"))
        self.grid = self.library["grid"]
        self.theme = ExcelTheme.from_spec(spec)
        self.thin = Side(style="thin", color=self.theme.line)
        self.border = Border(bottom=self.thin)
        self._table_counter = 0

    def prepare_workbook(self) -> None:
        for ws in self.workbook.worksheets:
            self.prepare_sheet(ws, ws.title)

    def prepare_sheet(self, ws, tab: str) -> None:
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = int(self.spec.get("native_excel_design_contract", {}).get("opening_zoom", 100))
        contract = self.spec.get("native_excel_design_contract", {})
        widths = (contract.get("screen_column_widths", {}) or {}).get(tab, contract.get("column_widths", {}))
        for column in range(1, self.grid["last_content_column"] + 1):
            letter = get_column_letter(column)
            ws.column_dimensions[letter].width = float(widths.get(letter, 15))
        for row in range(1, self.grid["navigation_row"] + 1):
            ws.row_dimensions[row].height = self.grid["body_row_height"]
        for row in self.grid["header_rows"]:
            ws.row_dimensions[row].height = self.grid["header_row_height"]
        ws.row_dimensions[self.grid["navigation_row"]].height = 24

    def write_header(self, ws, tab: str, title: str | None = None, subtitle: str | None = None) -> None:
        title = title or tab
        subtitle = subtitle or self.spec.get("subtitle", "")
        ws["B2"] = f"{self.spec.get('slug', 'excel').replace('-', ' ').upper()} · {tab.upper()}"
        ws["B3"] = title
        ws["B4"] = subtitle
        for cell in ("B2", "B3", "B4"):
            ws.merge_cells(start_row=ws[cell].row, start_column=2, end_row=ws[cell].row, end_column=13)
        ws["B2"].font = Font(name="Aptos", size=10, bold=True, color=self.theme.accent)
        ws["B3"].font = Font(name="Aptos Display", size=22, bold=True, color=self.theme.ink)
        ws["B4"].font = Font(name="Aptos", size=10, color=self.theme.muted)
        ws["B2"].alignment = Alignment(horizontal="left", vertical="center")
        ws["B3"].alignment = Alignment(horizontal="left", vertical="center")
        ws["B4"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def write_summary_row(self, ws, metrics: Iterable[tuple[str, Any]], row: int = 6) -> None:
        first = self.grid["first_content_column"]
        last = self.grid["last_content_column"]
        span = (last - first + 1) // 4
        for metric_index, (label, value) in enumerate(list(metrics)[:4]):
            start = first + metric_index * span
            end = start + span - 1
            for summary_row in (row, row + 1):
                ws.merge_cells(start_row=summary_row, start_column=start, end_row=summary_row, end_column=end)
                for column in range(start, end + 1):
                    cell = ws.cell(summary_row, column)
                    cell.fill = PatternFill("solid", fgColor=self.theme.formula_fill)
                    cell.border = self.border
            label_cell = ws.cell(row, start, str(label).upper())
            value_cell = ws.cell(row + 1, start, value)
            label_cell.font = Font(size=9, bold=True, color=self.theme.muted)
            value_cell.font = Font(size=15, bold=True, color=self.theme.ink)
            label_cell.alignment = Alignment(horizontal="left")
            value_cell.alignment = Alignment(horizontal="left")
            if isinstance(value, (int, float)) or (isinstance(value, str) and value.startswith("=")):
                value_cell.number_format = "#,##0.##"

    def write_table(
        self,
        ws,
        headers: list[str],
        rows: list[list[Any]],
        start_row: int = 9,
        editable_columns: set[int] | None = None,
        validations: list[dict[str, Any]] | None = None,
    ) -> None:
        editable_columns = editable_columns or set()
        if not isinstance(headers, list) or not headers:
            raise ValueError("Una tabla debe declarar una lista de encabezados")
        header_errors = validate_headers(headers)
        if header_errors:
            raise ValueError("; ".join(header_errors))
        if any(len(values) != len(headers) for values in rows):
            raise ValueError("Cada fila de una tabla debe tener el mismo número de columnas que sus encabezados")
        for column, header in enumerate(headers, self.grid["first_content_column"]):
            cell = ws.cell(start_row, column, header)
            cell.fill = PatternFill("solid", fgColor=self.theme.accent)
            cell.font = Font(size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = self.border
        ws.row_dimensions[start_row].height = self.grid["header_row_height"]
        for offset, values in enumerate(rows, start_row + 1):
            for column, value in enumerate(values, self.grid["first_content_column"]):
                cell = ws.cell(offset, column, value)
                is_formula = isinstance(value, str) and value.startswith("=")
                fill = self.theme.formula_fill if is_formula else (
                    self.theme.input_fill if column - self.grid["first_content_column"] in editable_columns else self.theme.surface
                )
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.font = Font(size=10, color=self.theme.ink)
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
                cell.border = self.border
                if isinstance(value, (int, float)) or (isinstance(value, str) and value.startswith("=")):
                    cell.number_format = number_format_for_header(
                        headers[column - self.grid["first_content_column"]], value
                    ) or "#,##0.##"
            ws.row_dimensions[offset].height = self.grid["body_row_height"]
        self._table_counter += 1
        start_column = self.grid["first_content_column"]
        end_column = start_column + len(headers) - 1
        end_row = start_row + max(len(rows), 1)
        table_ref = f"{get_column_letter(start_column)}{start_row}:{get_column_letter(end_column)}{end_row}"
        table_name = (
            f"tbl_{self.spec.get('slug', 'excel')}_{ws.title}_{self._table_counter}"
            .replace(" ", "_")
            .replace("-", "_")
        )
        table_name = "".join(char for char in table_name if char.isalnum() or char == "_")
        if table_name[:1].isdigit():
            table_name = f"tbl_{table_name}"
        ws.add_table(Table(displayName=table_name[:250], ref=table_ref))
        for validation in validations or []:
            index = int(validation["column"])
            if index < 0 or index >= len(headers):
                raise ValueError(f"La validación apunta a una columna inexistente: {index}")
            values = [str(value).replace('"', "'") for value in validation.get("values", [])]
            if not values:
                continue
            formula = '"' + ",".join(values) + '"'
            rule = DataValidation(type="list", formula1=formula, allow_blank=True)
            ws.add_data_validation(rule)
            first_row = start_row + 1
            last_row = start_row + max(len(rows), 1)
            column_letter = get_column_letter(self.grid["first_content_column"] + index)
            rule.add(f"{column_letter}{first_row}:{column_letter}{last_row}")

    def write_input_block(self, ws, rows: list[tuple[str, Any, str]], start_row: int = 6, validations: list[dict[str, Any]] | None = None) -> None:
        self.write_table(ws, ["Setting", "Value", "Why it matters"], [list(row) for row in rows], start_row, {1}, validations)

    def write_progress_bar(self, ws, value: float, row: int = 6) -> None:
        ws.cell(row, self.grid["first_content_column"], "Progress")
        ws.cell(row, self.grid["first_content_column"] + 1, value)
        if isinstance(value, str) and value.startswith("="):
            ws.cell(row, self.grid["first_content_column"] + 1).number_format = "0%"
        for column in range(self.grid["first_content_column"], self.grid["last_content_column"] + 1):
            cell = ws.cell(row, column)
            cell.fill = PatternFill("solid", fgColor=self.theme.accent if column <= self.grid["first_content_column"] + 2 else self.theme.formula_fill)
            cell.font = Font(size=10, bold=column == self.grid["first_content_column"], color=self.theme.ink)
            cell.border = self.border

    def write_note(self, ws, value: str, row: int = 28) -> None:
        ws.cell(row, self.grid["first_content_column"], value)
        ws.merge_cells(start_row=row, start_column=self.grid["first_content_column"], end_row=row, end_column=self.grid["last_content_column"])
        cell = ws.cell(row, self.grid["first_content_column"])
        cell.fill = PatternFill("solid", fgColor=self.theme.input_fill)
        cell.font = Font(size=10, bold=True, color=self.theme.ink)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def write_chart(self, ws, chart_spec: dict[str, Any]) -> None:
        """Add one small native Excel chart from explicit worksheet ranges."""
        chart_type = str(chart_spec.get("type", "")).lower()
        chart_classes = {
            "bar": BarChart,
            "column": BarChart,
            "line": LineChart,
            "pie": PieChart,
            "doughnut": DoughnutChart,
        }
        if chart_type not in chart_classes:
            raise ValueError(f"Tipo de grÃ¡fico no permitido: {chart_type}")

        def reference(raw: dict[str, Any]) -> Reference:
            required = {"min_col", "min_row", "max_col", "max_row"}
            if not required.issubset(raw):
                raise ValueError("El rango del grÃ¡fico debe declarar min/max row y column")
            return Reference(
                ws,
                min_col=int(raw["min_col"]),
                min_row=int(raw["min_row"]),
                max_col=int(raw["max_col"]),
                max_row=int(raw["max_row"]),
            )

        chart = chart_classes[chart_type]()
        if chart_type == "column":
            chart.type = "col"
        elif chart_type == "bar":
            chart.type = "bar"
        chart.title = chart_spec.get("title", "")
        chart.style = int(chart_spec.get("style", 10))
        chart.width = float(chart_spec.get("width", 12))
        chart.height = float(chart_spec.get("height", 7))
        values = reference(dict(chart_spec.get("values") or {}))
        chart.add_data(values, titles_from_data=bool(chart_spec.get("titles_from_data", True)))
        categories = chart_spec.get("categories")
        if categories:
            chart.set_categories(reference(dict(categories)))
        if chart_type in {"bar", "column", "line"} and chart_spec.get("legend") is False:
            chart.legend = None
        if chart.series:
            chart.series[0].graphicalProperties.solidFill = self.theme.secondary
            if len(chart.series) > 1:
                chart.series[1].graphicalProperties.solidFill = self.theme.highlight
        ws.add_chart(chart, str(chart_spec.get("anchor", "H9")))

    def write_navigation(self, ws, active: str, row: int | None = None) -> None:
        """Keep navigation in Excel's native sheet tabs, not in worksheet cells."""
        return None

    def apply_screen(self, ws, tab: str, content: dict[str, Any] | None = None) -> None:
        content = content or {}
        contract = (self.spec.get("screen_contracts") or {}).get(tab, {})
        template = contract.get("template", "table")
        self.write_header(ws, tab, content.get("title"), content.get("subtitle"))
        if template in {"dashboard", "analysis"}:
            self.write_summary_row(ws, content.get("summary", [("Main metric", ""), ("Items", ""), ("Due soon", ""), ("Status", "")]))
            self.write_table(ws, content.get("headers", ["Item", "Status", "Value"]), content.get("rows", []), 9, set(content.get("editable_columns", [])), content.get("validations"))
            if content.get("note"):
                self.write_note(ws, str(content["note"]))
        elif template == "checklist":
            self.write_progress_bar(ws, content.get("progress", 0), 6)
            self.write_table(ws, content.get("headers", ["Task", "Owner", "Due", "Status"]), content.get("rows", []), 9, set(content.get("editable_columns", [])), content.get("validations"))
            if content.get("note"):
                self.write_note(ws, str(content["note"]))
        elif template in {"guide", "settings"}:
            self.write_input_block(ws, content.get("rows", []), 6, content.get("validations"))
            self.write_note(ws, content.get("note", "Use highlighted cells for inputs."))
        else:
            self.write_table(ws, content.get("headers", ["Item", "Status", "Value"]), content.get("rows", []), 6, set(content.get("editable_columns", [])), content.get("validations"))
        if "chart" in contract.get("components", []):
            chart_spec = contract.get("chart") or content.get("chart")
            if chart_spec:
                self.write_chart(ws, chart_spec)


def create_template_workbook(spec: dict[str, Any], content_by_tab: dict[str, dict[str, Any]] | None = None) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    factory = ExcelTemplateFactory(workbook, spec)
    for tab in spec.get("tabs", []):
        ws = workbook.create_sheet(tab)
        factory.prepare_sheet(ws, tab)
        factory.apply_screen(ws, tab, (content_by_tab or {}).get(tab))
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    return workbook
