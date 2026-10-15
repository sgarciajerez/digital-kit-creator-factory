from __future__ import annotations

import json
from pathlib import Path
from copy import copy
from datetime import date

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from apply_excel_visual_overlays import main as apply_visual_overlays


ROOT = Path(__file__).resolve().parents[2]
BOOK = ROOT / "outputs" / "listings" / "product-7" / "autopilot_3_Planificador_de_presupuesto_familiar_0.1.0.xlsx"
SPEC_PATH = ROOT / "outputs" / "listings" / "product-7" / "design_spec.json"
REVIEW_PATH = ROOT / "outputs" / "listings" / "product-7" / "design_review.json"
SPEC: dict = {}

NAVY = "142B45"
TEAL = "2E9C9A"
MINT = "D8EFEB"
CORAL = "F18A72"
CREAM = "F5F1E8"
INK = "233044"
MUTED = "6B7785"
WHITE = "FFFFFF"
PALE_BLUE = "EAF1F7"
PALE_CORAL = "FCE7E1"
PALE_YELLOW = "FFF3C4"
GREEN = "237A63"

thin_gray = Side(style="thin", color="D7DEE7")
medium_navy = Side(style="medium", color=NAVY)


def sheet_spec(name: str) -> dict:
    return SPEC["sheets"][name]


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def set_title(ws, title: str, subtitle: str, end_col: str = "H", spacer_height: int = 80) -> None:
    ws.sheet_view.zoomScale = SPEC.get("typography", {}).get("zoom", 110)
    ws.sheet_properties.tabColor = TEAL
    ws.merge_cells(f"A1:{end_col}2")
    ws["A1"] = title
    typography = SPEC.get("typography", {})
    ws["A1"].font = Font(
        name=typography.get("title_font", "Aptos Display"),
        size=typography.get("title_size", 28),
        bold=True,
        color=WHITE,
    )
    ws["A1"].fill = fill(NAVY)
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=4)
    ws.row_dimensions[1].height = 38
    ws.row_dimensions[2].height = 42
    for row in ws.iter_rows(min_row=1, max_row=2, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.fill = fill(NAVY)
    ws.merge_cells(f"A3:{end_col}3")
    ws["A3"] = subtitle
    ws["A3"].font = Font(name=typography.get("body_font", "Aptos"), size=16, italic=True, color=MUTED)
    ws["A3"].alignment = Alignment(vertical="center", indent=4, shrink_to_fit=True)
    ws.row_dimensions[3].height = 30
    ws.row_dimensions[4].height = spacer_height


def style_header_row(ws, row: int, start: int, end: int) -> None:
    for cell in ws.iter_rows(min_row=row, max_row=row, min_col=start, max_col=end).__next__():
        typography = SPEC.get("typography", {})
        cell.font = Font(name=typography.get("body_font", "Aptos"), size=typography.get("body_size", 12), bold=True, color=WHITE)
        cell.fill = fill(TEAL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=medium_navy, bottom=medium_navy)
    ws.row_dimensions[row].height = 38


def style_body(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            typography = SPEC.get("typography", {})
            cell.font = Font(name=typography.get("body_font", "Aptos"), size=typography.get("body_size", 12), color=INK)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin_gray)


def reset_sheet(ws) -> None:
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    ws._charts = []
    ws._images = []
    ws.auto_filter.ref = None
    ws.data_validations.dataValidation = []
    ws.conditional_formatting = type(ws.conditional_formatting)()
    if ws.max_row:
        ws.delete_rows(1, ws.max_row)


def configure_instructions(ws) -> None:
    reset_sheet(ws)
    design = sheet_spec("Instrucciones")
    set_title(ws, design["title"], design["subtitle"], "K", spacer_height=112)
    for col in ["A", "B", "C"]:
        ws.column_dimensions[col].width = 22
    for col in ["E", "F", "G"]:
        ws.column_dimensions[col].width = 24
    for col in ["I", "J", "K"]:
        ws.column_dimensions[col].width = 21
    for col in ["D", "H"]:
        ws.column_dimensions[col].width = 4
    areas = [("A", "C"), ("E", "G"), ("I", "K")]
    for area, card in zip(areas, design["cards"], strict=True):
        number = card["number"]
        heading = card["heading"]
        text = card["text"]
        start_col, end_col = area
        for cell_range in [f"{start_col}5:{end_col}5", f"{start_col}6:{end_col}6", f"{start_col}7:{end_col}10"]:
            ws.merge_cells(cell_range)
        card_fill = fill(SPEC["theme"][card["color"]])
        for row in range(5, 11):
            for col in range(ws[f"{start_col}{row}"].column, ws[f"{end_col}{row}"].column + 1):
                ws.cell(row, col).fill = card_fill
        number_cell = ws[f"{start_col}5"]
        heading_cell = ws[f"{start_col}6"]
        body_cell = ws[f"{start_col}7"]
        number_cell.value = number
        heading_cell.value = heading
        body_cell.value = text
        number_cell.font = Font(name="Aptos", size=18, bold=True, color=NAVY)
        heading_cell.font = Font(name="Aptos", size=20, bold=True, color=INK)
        body_cell.font = Font(name="Aptos", size=16, color=INK)
        number_cell.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
        heading_cell.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
        body_cell.alignment = Alignment(vertical="top", wrap_text=True, indent=1)
    ws.row_dimensions[5].height = 50
    for r in range(6, 11):
        ws.row_dimensions[r].height = 50
    ws.merge_cells("A13:K13")
    ws["A13"] = design["tip"]
    ws["A13"].font = Font(name="Aptos", size=12, italic=True, color=MUTED)
    ws["A13"].alignment = Alignment(wrap_text=True)
    ws.sheet_view.showGridLines = False


def configure_budget(ws) -> None:
    reset_sheet(ws)
    design = sheet_spec("Presupuesto")
    set_title(ws, design["title"], design["subtitle"], "E", spacer_height=102)
    headers = design["headers"]
    for col, value in enumerate(headers, 1):
        ws.cell(5, col, value)
    style_header_row(ws, 5, 1, 5)
    rows = design["rows"]
    for r, (category, planned, actual) in enumerate(rows, 6):
        ws.cell(r, 1, category)
        ws.cell(r, 2, planned)
        ws.cell(r, 3, actual)
        ws.cell(r, 4, f"=B{r}-C{r}")
        ws.cell(r, 5, f'=IF(D{r}>=0,"Dentro del plan","Revisar")')
    total_row = 13
    ws.cell(total_row, 1, "TOTAL")
    ws.cell(total_row, 2, "=SUM(B6:B12)")
    ws.cell(total_row, 3, "=SUM(C6:C12)")
    ws.cell(total_row, 4, "=B13-C13")
    ws.cell(total_row, 5, '=IF(D13>=0,"Dentro del plan","Revisar")')
    style_body(ws, 6, total_row, 1, 5)
    for r in range(6, 13):
        ws.row_dimensions[r].height = 34
        if r % 2 == 0:
            for c in range(1, 6):
                ws.cell(r, c).fill = fill(CREAM)
        for c in [2, 3]:
            ws.cell(r, c).fill = fill(PALE_YELLOW)
        ws.cell(r, 4).fill = fill(MINT)
        for c in [2, 3, 4]:
            ws.cell(r, c).number_format = '€ #,##0;[Red]-€ #,##0'
            ws.cell(r, c).alignment = Alignment(horizontal="right", vertical="center")
    for c in range(1, 6):
        cell = ws.cell(total_row, c)
        cell.font = Font(name="Aptos", size=12, bold=True, color=NAVY)
        cell.fill = fill(MINT)
        cell.border = Border(top=medium_navy, bottom=medium_navy)
    for c in [2, 3, 4]:
        ws.cell(total_row, c).number_format = '€ #,##0;[Red]-€ #,##0'
    ws.conditional_formatting.add("D6:D13", DataBarRule(start_type="min", end_type="max", color=TEAL, showValue=True))
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 32
    ws.column_dimensions["D"].width = 32
    ws.column_dimensions["E"].width = 50
    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False


def configure_preview(ws) -> None:
    reset_sheet(ws)
    design = sheet_spec("Vista previa")
    set_title(ws, design["title"], design["subtitle"], "H", spacer_height=132)
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 17
    ws.column_dimensions["C"].width = 4
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 17
    ws.column_dimensions["F"].width = 4
    ws.column_dimensions["G"].width = 18
    ws.column_dimensions["H"].width = 17
    cards = [
        (("A5:B5", "A6:B7"), "Presupuesto total", "=Presupuesto!B13", MINT),
        (("D5:E5", "D6:E7"), "Gasto real total", "=Presupuesto!C13", PALE_BLUE),
        (("G5:H5", "G6:H7"), "Diferencia", "=Presupuesto!D13", PALE_CORAL),
    ]
    for (label_area, value_area), label, formula, color in cards:
        ws.merge_cells(label_area)
        ws.merge_cells(value_area)
        label_cell = ws[label_area.split(":")[0]]
        value_cell = ws[value_area.split(":")[0]]
        label_cell.value = label
        value_cell.value = formula
        label_cell.fill = fill(color)
        value_cell.fill = fill(color)
        label_cell.font = Font(name="Aptos", size=12, bold=True, color=NAVY)
        value_cell.font = Font(name="Aptos Display", size=24, bold=True, color=NAVY)
        label_cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
        value_cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
        for row in ws.iter_rows(min_row=ws[label_area.split(":")[0]].row, max_row=ws[value_area.split(":")[1]].row,
                                min_col=ws[label_area.split(":")[0]].column,
                                max_col=ws[value_area.split(":")[1]].column):
            for cell in row:
                cell.fill = fill(color)
                cell.border = Border(left=medium_navy, right=medium_navy, top=medium_navy, bottom=medium_navy)
        value_cell.number_format = '€ #,##0;[Red]-€ #,##0'
    ws.row_dimensions[5].height = 46
    ws.row_dimensions[6].height = 46
    ws.row_dimensions[7].height = 46
    ws.merge_cells("A9:H9")
    ws["A9"] = "Gasto real por categoría"
    ws["A9"].font = Font(name="Aptos", size=14, bold=True, color=NAVY)
    ws["A9"].fill = fill(CREAM)
    ws["A10"] = "Categoría"
    ws["B10"] = "Gasto real"
    style_header_row(ws, 10, 1, 2)
    categories = [row[0] for row in sheet_spec("Presupuesto")["rows"]]
    for r, category in enumerate(categories, 11):
        ws.row_dimensions[r].height = 30
        ws.cell(r, 1, category)
        ws.cell(r, 2, f"=Presupuesto!C{r-5}")
        ws.cell(r, 2).number_format = '€ #,##0;[Red]-€ #,##0'
        if r % 2 == 1:
            ws.cell(r, 1).fill = fill(CREAM)
            ws.cell(r, 2).fill = fill(CREAM)
    style_body(ws, 11, 17, 1, 2)
    ws.conditional_formatting.add("B11:B17", DataBarRule(start_type="min", end_type="max", color=TEAL, showValue=True))

    # Replantear el listado como barras horizontales, igual que el prototipo aprobado.
    ws["A10"] = None
    ws["B10"] = None
    for header_cell in (ws["A10"], ws["B10"]):
        header_cell.fill = fill(WHITE)
        header_cell.border = Border()
    ws.conditional_formatting = type(ws.conditional_formatting)()
    ws.conditional_formatting.add("B11:B17", DataBarRule(start_type="min", end_type="max", color=TEAL, showValue=False))
    for r, category in enumerate(categories, 11):
        ws.merge_cells(f"B{r}:G{r}")
        ws.cell(r, 2, f'=REPT("█",ROUND(Presupuesto!C{r-5}/MAX(Presupuesto!$C$6:$C$12)*35,0))')
        ws.cell(r, 8, f"=Presupuesto!C{r-5}")
        ws.cell(r, 8).number_format = 'â‚¬ #,##0;[Red]-â‚¬ #,##0'
        ws.cell(r, 1).font = Font(name="Aptos", size=12, color=INK)
        ws.cell(r, 2).font = Font(name="Aptos", size=11, bold=True, color=TEAL)
        ws.cell(r, 2).alignment = Alignment(vertical="center", horizontal="left")
        ws.cell(r, 8).font = Font(name="Aptos", size=12, bold=True, color=INK)
        ws.cell(r, 8).alignment = Alignment(vertical="center", horizontal="right")
        for col in range(2, 8):
            ws.cell(r, col).fill = fill("E8EFEC")
        if r % 2 == 1:
            ws.cell(r, 1).fill = fill(CREAM)
            ws.cell(r, 8).fill = fill(CREAM)
    euro_format = chr(0x20AC) + " #,##0;[Red]-" + chr(0x20AC) + " #,##0"
    for r in range(11, 18):
        ws.cell(r, 8).number_format = euro_format
    ws.sheet_view.showGridLines = False


def configure_settings(ws) -> None:
    reset_sheet(ws)
    design = sheet_spec("Configuración")
    set_title(ws, design["title"], design["subtitle"], "H", spacer_height=112)
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 58
    for col in ["C", "D", "E", "F", "G", "H"]:
        ws.column_dimensions[col].width = 8
    ws["A5"] = "Parámetro"
    ws["B5"] = "Valor"
    style_header_row(ws, 5, 1, 2)
    values = list(design["values"].items())
    for r, (key, value) in enumerate(values, 6):
        ws.cell(r, 1, key)
        ws.cell(r, 2, value)
    style_body(ws, 6, 10, 1, 2)
    for r in range(6, 11):
        ws.row_dimensions[r].height = 32
        ws.cell(r, 1).font = Font(name="Aptos", size=12, color=INK)
        if r % 2 == 0:
            ws.cell(r, 1).fill = fill(CREAM)
            ws.cell(r, 2).fill = fill(CREAM)
    dv = DataValidation(type="list", formula1='"EUR,USD,GBP"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(ws["B9"])
    ws.sheet_view.showGridLines = False


def configure_shared(ws) -> None:
    reset_sheet(ws)
    design = sheet_spec("Gastos compartidos")
    set_title(ws, design["title"], design["subtitle"], "G", spacer_height=112)
    headers = design["headers"]
    for col, value in enumerate(headers, 1):
        ws.cell(5, col, value)
    style_header_row(ws, 5, 1, 8)
    rows = design["rows"]
    for r, (when, desc, category, payer, amount) in enumerate(rows, 6):
        values = [date.fromisoformat(when), desc, category, payer, amount, f"=E{r}/2", f"=E{r}/2", ""]
        for c, value in enumerate(values, 1):
            ws.cell(r, c, value)
    style_body(ws, 6, 10, 1, 8)
    for r in range(6, 11):
        ws.row_dimensions[r].height = 32
        ws.cell(r, 1).number_format = "dd/mm/yyyy"
        for c in [5, 6, 7]:
            ws.cell(r, c).number_format = '€ #,##0.00;[Red]-€ #,##0.00'
        for c in [6, 7]:
            ws.cell(r, c).fill = fill(MINT)
        if r % 2 == 0:
            for c in range(1, 9):
                ws.cell(r, c).fill = fill(CREAM)
    for col, width in {"A": 14, "B": 23, "C": 18, "D": 15, "E": 14, "F": 14, "G": 14, "H": 22}.items():
        ws.column_dimensions[col].width = width
    ws.column_dimensions["B"].width = 32
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["H"].hidden = True
    payer_dv = DataValidation(type="list", formula1="=Miembros", allow_blank=False)
    cat_dv = DataValidation(type="list", formula1='"Vivienda,Alimentación,Transporte,Salud,Educación,Ocio,Otros"', allow_blank=False)
    ws.add_data_validation(payer_dv)
    ws.add_data_validation(cat_dv)
    payer_dv.add("D6:D1000")
    cat_dv.add("C6:C1000")
    ws.freeze_panes = "A6"
    ws.sheet_view.showGridLines = False


def main() -> None:
    if not BOOK.exists():
        raise FileNotFoundError(BOOK)
    if not SPEC_PATH.exists():
        raise FileNotFoundError(SPEC_PATH)
    if REVIEW_PATH.exists():
        review = json.loads(REVIEW_PATH.read_text(encoding="utf-8"))
        if review.get("status") != "approved":
            raise RuntimeError(
                "El diseño PNG está pendiente de revisión humana. "
                "Aprueba outputs/listings/product-7/design_review.json antes de construir el Excel."
            )
    global SPEC, NAVY, TEAL, MINT, CORAL, CREAM, INK, MUTED, WHITE, PALE_BLUE, PALE_CORAL, PALE_YELLOW, medium_navy
    SPEC = json.loads(SPEC_PATH.read_text(encoding="utf-8"))
    theme = SPEC["theme"]
    NAVY = theme["navy"]
    TEAL = theme["teal"]
    MINT = theme["mint"]
    CORAL = theme["coral"]
    CREAM = theme["cream"]
    INK = theme["ink"]
    MUTED = theme["muted"]
    WHITE = theme["white"]
    PALE_BLUE = theme["pale_blue"]
    PALE_CORAL = theme["pale_coral"]
    PALE_YELLOW = theme["pale_yellow"]
    medium_navy = Side(style="medium", color=NAVY)
    wb = load_workbook(BOOK)
    for sheet in ["Instrucciones", "Presupuesto", "Vista previa", "Configuración", "Gastos compartidos"]:
        if sheet not in wb.sheetnames:
            wb.create_sheet(sheet)
    configure_instructions(wb["Instrucciones"])
    configure_budget(wb["Presupuesto"])
    configure_preview(wb["Vista previa"])
    configure_settings(wb["Configuración"])
    configure_shared(wb["Gastos compartidos"])
    if "Miembros" in wb.defined_names:
        del wb.defined_names["Miembros"]
    wb.defined_names.add(DefinedName("Miembros", attr_text="'Configuración'!$B$6:$B$7"))
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.active = wb.sheetnames.index("Instrucciones")
    wb.save(BOOK)
    apply_visual_overlays()
    print(f"Updated {BOOK}")


if __name__ == "__main__":
    main()
