from __future__ import annotations

"""Semantic and functional quality gates for generated Excel models.

The visual contract describes how a workbook should look. This module is the
model contract: it checks that tabular data has a meaningful schema, that
number formats agree with the meaning of a column, that totals do not become
input rows, and (when declared) that a real input changes the expected output.
"""

import json
import gc
import re
import shutil
import tempfile
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries


GENERIC_HEADER_RE = re.compile(r"^(?:columna|column|campo|field)\s*\d+$", re.IGNORECASE)
CELL_REF_RE = re.compile(r"(?<![A-Z0-9_])\$?([A-Z]{1,3})\$?(\d+)", re.IGNORECASE)
SHEET_REF_RE = re.compile(r"(?:'((?:[^']|'')+)'|([^\s=!<>+\-*/(),:]+))!", re.IGNORECASE)
FORMULA_REF_RE = re.compile(
    r"(?<![A-Z0-9_:$!])(?:(?:'((?:[^']|'')+)'|([^\s=!<>+\-*/(),:]+))!)?"
    r"\$?([A-Z]{1,3})\$?(\d+)",
    re.IGNORECASE,
)
CELL_ADDRESS_RE = re.compile(r"^\$?[A-Z]{1,3}\$?\d+$", re.IGNORECASE)

PERCENT_WORDS = (
    "%", "percent", "percentage", "porcentaje", "ratio", "progress", "completion",
    "completed", "share", "proportion", "rate", "apr", "tasa", "margen", "avance", "progreso", "cumplimiento",
)
COUNT_WORDS = (
    "count", "number of", "quantity", "qty", "units", "items", "guests", "guest count",
    "invitados", "confirmados", "cantidad", "unidades", "personas", "mesas", "tables",
    "days", "días", "dias", "household size", "priority", "rating", "meal count", "party",
)
DATE_WORDS = (
    "date", "due", "deadline", "expiry", "start", "fecha", "vencimiento", "plazo",
    "caducidad", "inicio",
)
AMOUNT_WORDS = (
    "amount", "budget", "planned", "actual", "spent", "paid", "payment", "balance",
    "income", "expense", "cost", "price", "revenue", "opening", "closing", "difference",
    "remaining", "estimate", "deposit", "principal", "interest", "importe", "presupuesto",
    "gasto", "pagado", "pago", "saldo", "ingresos", "coste", "precio", "diferencia",
    "restante", "estimado", "anticipo",
)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def is_generic_header(value: Any) -> bool:
    return bool(GENERIC_HEADER_RE.fullmatch(_text(value)))


def validate_headers(headers: Iterable[Any]) -> list[str]:
    """Return schema errors for a table header row."""
    values = [_text(header) for header in headers]
    errors: list[str] = []
    if not values:
        return ["la tabla no tiene columnas"]
    if any(not value for value in values):
        errors.append("hay encabezados vacíos")
    generic = [value for value in values if is_generic_header(value)]
    if generic:
        errors.append(f"encabezados genéricos no permitidos: {generic}")
    normalized = [value.casefold() for value in values]
    duplicates = sorted({value for value in normalized if normalized.count(value) > 1 and value})
    if duplicates:
        errors.append(f"encabezados duplicados: {duplicates}")
    return errors


def semantic_kind(header: Any) -> str:
    """Classify a column header without requiring a product-specific schema."""
    value = _text(header).casefold()
    if any(token in value for token in PERCENT_WORDS):
        return "percent"
    if any(token in value for token in DATE_WORDS):
        return "date"
    if any(token in value for token in COUNT_WORDS):
        return "count"
    if any(token in value for token in AMOUNT_WORDS):
        return "amount"
    return "general"


def number_format_for_header(header: Any, value: Any = None) -> str | None:
    """Choose a conservative number format from the declared column meaning."""
    kind = semantic_kind(header)
    if kind == "percent":
        return "0%"
    if kind == "count":
        return "0"
    if kind == "date":
        return "dd/mm/yyyy"
    if kind == "amount":
        return "#,##0.00"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "#,##0.##"
    return None


def _is_percentage_format(number_format: str | None) -> bool:
    return "%" in (number_format or "")


def _is_date_format(number_format: str | None) -> bool:
    value = (number_format or "").casefold()
    return any(token in value for token in ("yy", "dd", "mm", "mmm")) and "%" not in value


def _is_integer_format(number_format: str | None) -> bool:
    value = (number_format or "").replace("\\", "")
    return "%" not in value and "." not in value and bool(re.search(r"[0#]", value))


def _looks_numeric(value: Any) -> bool:
    return (
        isinstance(value, (int, float)) and not isinstance(value, bool)
    ) or (isinstance(value, str) and value.startswith("="))


def _table_ranges(ws) -> list[tuple[str, str]]:
    """Return native table refs, or auto-filter refs for legacy table blocks."""
    ranges: list[tuple[str, str]] = []
    for name, table in getattr(ws, "tables", {}).items():
        ref = table.ref if hasattr(table, "ref") else table
        ranges.append((str(name), str(ref)))
    if not ranges and ws.auto_filter and ws.auto_filter.ref:
        ranges.append(("auto_filter", str(ws.auto_filter.ref)))
    return ranges


def _formula_sheet_errors(formula: str, sheet_names: set[str]) -> list[str]:
    errors: list[str] = []
    for quoted, plain in SHEET_REF_RE.findall(formula):
        sheet = (quoted or plain).replace("''", "'")
        if sheet not in sheet_names:
            errors.append(f"referencia a hoja inexistente: {sheet}")
    return errors


def _cell_has_internal_total(ws, min_col: int, max_col: int, min_row: int, max_row: int) -> list[str]:
    total_rows: list[str] = []
    for row in range(min_row + 1, max_row + 1):
        values = [
            _text(ws.cell(row, col).value).casefold()
            for col in range(min_col, min(max_col, min_col + 1) + 1)
        ]
        if any(value in {"total", "totals", "subtotal", "total general"} for value in values):
            total_rows.append(f"{get_column_letter(min_col)}{row}:{get_column_letter(max_col)}{row}")
    return total_rows


def _validate_tabular_range(
    ws,
    label: str,
    ref: str,
    errors: list[str],
    checks: list[str],
    details: dict[str, Any],
) -> None:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(ref)
    except ValueError:
        errors.append(f"{ws.title}/{label}: rango inválido {ref}")
        return
    if max_col < min_col or max_row < min_row:
        errors.append(f"{ws.title}/{label}: rango vacío")
        return
    headers = [ws.cell(min_row, col).value for col in range(min_col, max_col + 1)]
    header_errors = validate_headers(headers)
    if header_errors:
        errors.extend(f"{ws.title}/{label}: {error}" for error in header_errors)
    else:
        checks.append(f"{ws.title}/{label}: encabezados semánticos")
    total_rows = _cell_has_internal_total(ws, min_col, max_col, min_row, max_row)
    if total_rows:
        errors.append(f"{ws.title}/{label}: totales dentro del rango de datos: {total_rows}")
        details.setdefault("total_rows_inside_data", []).extend(total_rows)
    else:
        checks.append(f"{ws.title}/{label}: totales fuera del rango de datos")

    format_errors: list[str] = []
    for offset, header in enumerate(headers):
        kind = semantic_kind(header)
        if kind == "general":
            continue
        col = min_col + offset
        for row in range(min_row + 1, max_row + 1):
            cell = ws.cell(row, col)
            if not _looks_numeric(cell.value):
                continue
            if kind == "percent" and not _is_percentage_format(cell.number_format):
                format_errors.append(f"{cell.coordinate} ({header}) debe usar porcentaje")
            elif kind == "count" and (
                _is_percentage_format(cell.number_format) or not _is_integer_format(cell.number_format)
            ):
                format_errors.append(f"{cell.coordinate} ({header}) no puede usar porcentaje ni decimales")
            elif (
                kind == "date"
                and cell.value is not None
                and isinstance(cell.value, (int, float))
                and not _is_date_format(cell.number_format)
            ):
                format_errors.append(f"{cell.coordinate} ({header}) debe usar formato de fecha")
            elif kind == "amount" and _is_percentage_format(cell.number_format):
                format_errors.append(f"{cell.coordinate} ({header}) no puede usar porcentaje")
    if format_errors:
        errors.extend(f"{ws.title}/{label}: {error}" for error in format_errors)
        details.setdefault("semantic_format_errors", []).extend(format_errors)
    elif any(semantic_kind(header) != "general" for header in headers):
        checks.append(f"{ws.title}/{label}: formatos numéricos coherentes")


def validate_model_workbook(workbook, expected_sheets: Iterable[str] | None = None) -> dict[str, Any]:
    """Validate structure and formula references in an already opened workbook."""
    checks: list[str] = []
    errors: list[str] = []
    details: dict[str, Any] = {}
    sheet_names = set(workbook.sheetnames)
    expected = list(expected_sheets or [])
    missing = [sheet for sheet in expected if sheet not in sheet_names]
    if missing:
        errors.append(f"faltan hojas: {missing}")
    elif expected:
        checks.append("todas las hojas declaradas existen")

    formula_count = 0
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if not (isinstance(value, str) and value.startswith("=")):
                    continue
                formula_count += 1
                if "#REF!" in value.upper():
                    errors.append(f"{ws.title}!{cell.coordinate}: fórmula con #REF!")
                errors.extend(
                    f"{ws.title}!{cell.coordinate}: {error}"
                    for error in _formula_sheet_errors(value, sheet_names)
                )
                direct_self_reference = False
                for quoted_sheet, plain_sheet, column, row_number in FORMULA_REF_RE.findall(value):
                    referenced_sheet = (quoted_sheet or plain_sheet or "").replace("''", "'")
                    if (not referenced_sheet or referenced_sheet == ws.title) and (
                        f"{column}{row_number}".upper() == cell.coordinate.upper()
                    ):
                        direct_self_reference = True
                        break
                if direct_self_reference and "INDIRECT" not in value.upper():
                    errors.append(f"{ws.title}!{cell.coordinate}: posible referencia circular directa")
    if formula_count:
        checks.append(f"fórmulas revisadas: {formula_count}")
    else:
        checks.append("modelo sin fórmulas: no requiere propagación calculada")

    table_count = 0
    for ws in workbook.worksheets:
        ranges = _table_ranges(ws)
        table_count += len(ranges)
        for label, ref in ranges:
            _validate_tabular_range(ws, label, ref, errors, checks, details)
    if table_count:
        checks.append(f"rangos tabulares revisados: {table_count}")
    else:
        checks.append("sin rangos tabulares declarados")

    return {
        "passed": not errors,
        "checks": checks,
        "errors": errors,
        "details": details,
        "formula_count": formula_count,
        "table_count": table_count,
    }


def _native_table_records(workbook) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for ws in workbook.worksheets:
        for name, table in getattr(ws, "tables", {}).items():
            ref = str(getattr(table, "ref", table))
            min_col, min_row, max_col, max_row = range_boundaries(ref)
            headers = [ws.cell(min_row, col).value for col in range(min_col, max_col + 1)]
            records.append({
                "sheet": ws.title,
                "name": str(name),
                "ref": ref,
                "headers": headers,
                "min_row": min_row,
                "max_row": max_row,
            })
    return records


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _list_validations_from_xlsx(path: Path) -> dict[tuple[str, str], list[str]]:
    """Read list-validation values from XLSX XML, including Excel x12 extensions."""
    results: dict[tuple[str, str], list[str]] = {}
    named_ranges: dict[str, str] = {}
    workbook_values = load_workbook(path, data_only=False, read_only=False)
    with ZipFile(path) as archive:
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        for defined_name in workbook_root.iter():
            if _local_name(defined_name.tag) == "definedName":
                name = defined_name.attrib.get("name", "")
                if name and defined_name.text:
                    named_ranges[name] = defined_name.text.strip()
        rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_targets = {
            rel.attrib.get("Id"): rel.attrib.get("Target", "").lstrip("/")
            for rel in rels_root
        }
        for sheet in workbook_root.iter():
            if _local_name(sheet.tag) != "sheet":
                continue
            name = sheet.attrib.get("name", "")
            relationship_id = next((value for key, value in sheet.attrib.items() if _local_name(key) == "id"), "")
            target = rel_targets.get(relationship_id, "")
            sheet_path = target if target.startswith("xl/") else f"xl/{target}"
            if sheet_path not in archive.namelist():
                continue
            root = ET.fromstring(archive.read(sheet_path))
            for validation in root.iter():
                if _local_name(validation.tag) != "dataValidation" or validation.attrib.get("type") != "list":
                    continue
                sqref = validation.attrib.get("sqref", "").strip()
                formula = next((item.text or "" for item in validation.iter() if _local_name(item.tag) == "formula1"), "")
                formula = formula.strip().strip('"')
                values: list[str]
                named_range_key = formula.lstrip("=")
                if named_range_key in named_ranges:
                    reference = named_ranges[named_range_key]
                    reference = reference.lstrip("=")
                    match = re.match(r"^'((?:''|[^'])+)'!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)$", reference)
                    if not match:
                        match = re.match(r"^([^!]+)!\$([A-Z]+)\$(\d+):\$([A-Z]+)\$(\d+)$", reference)
                    if match:
                        sheet_name = match.group(1).replace("''", "'")
                        start_col, start_row, end_col, end_row = match.groups()[1:]
                        ws = workbook_values[sheet_name]
                        values = [
                            str(ws[f"{start_col}{row}"].value)
                            for row in range(int(start_row), int(end_row) + 1)
                            if ws[f"{start_col}{row}"].value not in (None, "")
                        ]
                    else:
                        values = [formula]
                else:
                    values = [formula] if formula.startswith("=") else [item.strip() for item in formula.split(",") if item.strip()]
                results[(name, sqref)] = values
    workbook_values.close()
    return results


def _validate_enum_slot_against_table(workbook, enum: dict[str, Any]) -> list[str]:
    """Validate that an enum targets the declared table column and seed data."""
    errors: list[str] = []
    sheet_name = str(enum.get("sheet", ""))
    address = str(enum.get("range", ""))
    expected_column = str(enum.get("column", "")).strip()
    expected_values = {str(value) for value in enum.get("values") or []}
    if sheet_name not in workbook.sheetnames:
        return [f"{sheet_name}!{address}: la hoja del enum no existe"]
    try:
        min_col, min_row, max_col, max_row = range_boundaries(address)
    except Exception:
        return [f"{sheet_name}!{address}: rango de enum inválido"]
    ws = workbook[sheet_name]
    matching_table = None
    for table_name, table in ws.tables.items():
        table_ref = getattr(table, "ref", table)
        table_min_col, table_min_row, table_max_col, table_max_row = range_boundaries(str(table_ref))
        if table_min_row == min_row - 1 and table_min_col <= min_col <= table_max_col:
            matching_table = (table_name, table_min_col, table_min_row, table_max_col, table_max_row)
            header = ws.cell(table_min_row, min_col).value
            if str(header).strip() != expected_column:
                errors.append(
                    f"{sheet_name}!{address}: apunta a {header!r}, pero el contrato declara la columna {expected_column!r}"
                )
            if max_col != min_col:
                errors.append(f"{sheet_name}!{address}: el enum debe cubrir una sola columna")
            break
    if matching_table is None:
        errors.append(f"{sheet_name}!{address}: no coincide con una columna de tabla declarada")
    else:
        _, _, _, _, table_max_row = matching_table
        for row in range(min_row, min(max_row, table_max_row) + 1):
            value = ws.cell(row, min_col).value
            if value not in (None, "") and str(value) not in expected_values:
                errors.append(
                    f"{sheet_name}!{ws.cell(row, min_col).coordinate}: valor mock {value!r} fuera del enum {sorted(expected_values)!r}"
                )
    return errors


def validate_model_behavior_markdown(path: Path, contract: dict[str, Any]) -> dict[str, Any]:
    """Validate the independent normative guide before it is used as QA."""
    if not path.exists():
        return {"passed": False, "checks": [], "errors": [f"falta guía funcional: {path}"], "details": {}}
    try:
        content = path.read_text(encoding="utf-8")
    except Exception as exc:
        return {"passed": False, "checks": [], "errors": [f"no se puede leer la guía funcional: {exc}"], "details": {}}
    errors: list[str] = []
    checks: list[str] = []
    required_sections = (
        "## Guía rápida para el usuario",
        "## Hojas que debe contener el libro",
        "## Tablas y columnas esperadas",
        "## Reglas de cálculo",
        "## Casos de verificación",
        "## Criterios de fallo",
        "## Regla de independencia",
    )
    required_sections = required_sections + ("## Listas y enums permitidos",)
    for section in required_sections:
        if section not in content:
            errors.append(f"la guía funcional no contiene la sección {section}")
    independent_markers = ("normativo", "no se debe redactar leyendo hacia atrás", "error observado")
    if not all(marker.casefold() in content.casefold() for marker in independent_markers):
        errors.append("la guía no declara su independencia del XLSX generado")
    for sheet in contract.get("sheets") or []:
        if f"`{sheet}`" not in content:
            errors.append(f"la guía no documenta la hoja esperada {sheet}")
    for case in contract.get("cases") or []:
        if case.get("name") and case["name"] not in content:
            errors.append(f"la guía no documenta el caso {case['name']!r}")
        for assertion in case.get("assertions") or []:
            expected = str(assertion.get("value", ""))
            if expected and expected not in content:
                errors.append(f"la guía no documenta el resultado esperado {expected!r}")
    for enum in contract.get("enum_contract") or []:
        for value in enum.get("values") or []:
            if f"`{value}`" not in content:
                errors.append(f"la guia no documenta el enum {value!r} de {enum.get('sheet')}!{enum.get('range')}")
    if not errors:
        checks.append("guía funcional normativa presente y alineada con el contrato declarado")
    return {"passed": not errors, "checks": checks, "errors": errors, "details": {"path": str(path)}}


def validate_model_contract(
    workbook,
    contract: dict[str, Any],
    markdown_path: Path | None = None,
    workbook_path: Path | None = None,
) -> dict[str, Any]:
    """Compare an expected model contract with an open workbook.

    The comparison is intentionally strict for declared ranges and formulas.
    A mismatch is an implementation failure; the contract is never adjusted
    from the observed workbook by this function.
    """
    from model_contracts import validate_contract_schema

    errors = list(validate_contract_schema(contract))
    checks: list[str] = []
    details: dict[str, Any] = {}
    if markdown_path is not None:
        markdown = validate_model_behavior_markdown(markdown_path, contract)
        details["markdown"] = markdown
        errors.extend(markdown.get("errors", []))
        checks.extend(markdown.get("checks", []))
    if contract.get("enum_contract"):
        if workbook_path is None:
            errors.append("el contrato declara enums pero falta workbook_path para validar las listas Excel")
        else:
            actual_validations = _list_validations_from_xlsx(workbook_path)
            for enum in contract["enum_contract"]:
                key = (str(enum["sheet"]), str(enum["range"]))
                expected_values = [str(value) for value in enum["values"]]
                actual_values = actual_validations.get(key)
                if actual_values is None:
                    errors.append(f"{enum['sheet']}!{enum['range']}: falta validacion de lista")
                elif actual_values != expected_values:
                    errors.append(f"{enum['sheet']}!{enum['range']}: enum real {actual_values!r}; se esperaba {expected_values!r}")
                else:
                    checks.append(f"{enum['sheet']}!{enum['range']}: enum alineado con el contrato")
                enum_slot_errors = _validate_enum_slot_against_table(workbook, enum)
                errors.extend(enum_slot_errors)
                if not enum_slot_errors:
                    checks.append(f"{enum['sheet']}!{enum['range']}: columna y datos mock alineados")
    expected_sheets = set(contract.get("sheets") or [])
    missing_sheets = sorted(expected_sheets - set(workbook.sheetnames))
    if missing_sheets:
        errors.append(f"contrato de modelo: faltan hojas {missing_sheets}")
    elif expected_sheets:
        checks.append("contrato de modelo: todas las hojas esperadas existen")

    actual_tables = _native_table_records(workbook)
    for index, expected in enumerate(contract.get("tables") or [], 1):
        candidates = [
            record for record in actual_tables
            if record["sheet"] == expected.get("sheet") and record["ref"] == expected.get("ref")
        ]
        if expected.get("name"):
            candidates = [record for record in candidates if record["name"] == expected["name"]]
        if not candidates:
            errors.append(
                f"tabla declarada {index}: no existe {expected.get('sheet')}!{expected.get('ref')}"
            )
            continue
        actual = candidates[0]
        expected_headers = [item.get("name") for item in expected.get("headers") or []]
        if actual["headers"] != expected_headers:
            errors.append(
                f"{expected['sheet']}!{expected['ref']}: cabeceras reales {actual['headers']!r}; "
                f"se esperaban {expected_headers!r}"
            )
        else:
            checks.append(f"{expected['sheet']}!{expected['ref']}: esquema declarado correcto")
            ws = workbook[expected["sheet"]]
            min_col, min_row, max_col, max_row = range_boundaries(expected["ref"])
            contract_format_errors: list[str] = []
            for offset, column in enumerate(expected.get("headers") or []):
                kind = column.get("kind")
                if kind not in {"amount", "count", "percent", "date"}:
                    continue
                for row_number in range(min_row + 1, max_row + 1):
                    cell = ws.cell(row_number, min_col + offset)
                    if not _looks_numeric(cell.value):
                        continue
                    if kind == "percent" and not _is_percentage_format(cell.number_format):
                        contract_format_errors.append(f"{cell.coordinate} ({column['name']}) debe ser porcentaje")
                    elif kind == "count" and (
                        _is_percentage_format(cell.number_format)
                        or (
                            not _is_integer_format(cell.number_format)
                            and not (isinstance(cell.value, int) and not isinstance(cell.value, bool))
                        )
                    ):
                        contract_format_errors.append(f"{cell.coordinate} ({column['name']}) debe ser entero")
                    elif kind == "amount" and _is_percentage_format(cell.number_format):
                        contract_format_errors.append(f"{cell.coordinate} ({column['name']}) no puede ser porcentaje")
                    elif kind == "date" and isinstance(cell.value, (int, float)) and not _is_date_format(cell.number_format):
                        contract_format_errors.append(f"{cell.coordinate} ({column['name']}) debe ser fecha")
            if contract_format_errors:
                errors.extend(f"{expected['sheet']}!{expected['ref']}: {error}" for error in contract_format_errors)
            else:
                checks.append(f"{expected['sheet']}!{expected['ref']}: formatos del contrato correctos")
        if expected.get("total_ref"):
            total_ref = str(expected["total_ref"])
            try:
                min_col, min_row, max_col, max_row = range_boundaries(total_ref)
                ws = workbook[expected["sheet"]]
                if min_row <= actual["max_row"]:
                    errors.append(f"{expected['sheet']}!{total_ref}: el total invade el rango de datos")
                label = _text(ws.cell(min_row, min_col).value).casefold()
                if any(token in label for token in ("total", "subtotal")):
                    checks.append(f"{expected['sheet']}!{total_ref}: total fuera de la tabla")
                else:
                    errors.append(f"{expected['sheet']}!{total_ref}: falta etiqueta TOTAL fuera de la tabla")
                formula_cells = [ws.cell(row, col) for row in range(min_row, max_row + 1) for col in range(min_col, max_col + 1) if isinstance(ws.cell(row, col).value, str) and ws.cell(row, col).value.startswith("=")]
                cramped_totals = [cell.coordinate for cell in formula_cells if not cell.alignment.shrinkToFit]
                if cramped_totals:
                    errors.append(f"{expected['sheet']}!{total_ref}: totales sin ajuste al ancho {cramped_totals}")
                else:
                    checks.append(f"{expected['sheet']}!{total_ref}: resultados ajustados al ancho de celda")
            except Exception as exc:
                errors.append(f"{expected['sheet']}!{total_ref}: rango de total inválido ({exc})")

    for rule in contract.get("formula_rules") or []:
        sheet = rule.get("sheet")
        cell_address = rule.get("cell")
        if sheet not in workbook.sheetnames:
            continue
        cell = workbook[sheet][cell_address]
        formula = str(cell.value or "")
        if not formula.startswith("="):
            errors.append(f"{sheet}!{cell_address}: falta fórmula declarada")
        elif str(rule.get("contains", "")).casefold() not in formula.casefold():
            errors.append(
                f"{sheet}!{cell_address}: la fórmula no contiene {rule.get('contains')!r}; fórmula {formula!r}"
            )
        else:
            checks.append(f"{sheet}!{cell_address}: regla de cálculo presente")
    return {
        "passed": not errors,
        "checks": checks,
        "errors": errors,
        "details": details,
        "table_count": len(actual_tables),
    }


def _values_equal(actual: Any, expected: Any, tolerance: float = 0.0001) -> bool:
    if isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
        return abs(float(actual) - float(expected)) <= tolerance
    if actual is None and expected in (None, ""):
        return True
    return str(actual) == str(expected)


def validate_functional_contract_schema(
    contract: dict[str, Any],
    sheet_names: Iterable[str],
) -> list[str]:
    """Check that functional cases point at real cells before opening Excel."""
    errors: list[str] = []
    known_sheets = set(sheet_names)
    if not contract.get("cases"):
        errors.append("faltan casos funcionales")
    operators = {"equals", "not_equals", "changed", "unchanged", "contains", "formula_contains", "not_error"}
    for case_index, case in enumerate(contract.get("cases") or [], 1):
        prefix = f"caso funcional {case_index}"
        if not case.get("assertions"):
            errors.append(f"{prefix}: faltan resultados esperados")
        for action in case.get("actions", []):
            if action.get("sheet") not in known_sheets or not action.get("table"):
                errors.append(f"{prefix}: acción sin hoja o tabla válida")
            if action.get("operation") not in {"append_rows", "delete_row"}:
                errors.append(f"{prefix}: operación no soportada")
            if action.get("operation") == "append_rows" and not 1 <= int(action.get("count", 1)) <= 10000:
                errors.append(f"{prefix}: número de filas fuera de rango")
        for item in list(case.get("inputs") or []) + list(case.get("assertions") or []):
            sheet = item.get("sheet")
            cell = item.get("cell")
            if sheet not in known_sheets:
                errors.append(f"{prefix}: hoja inexistente {sheet!r}")
            if not isinstance(cell, str) or not CELL_ADDRESS_RE.fullmatch(cell):
                errors.append(f"{prefix}: celda inválida {cell!r}")
        for assertion in case.get("assertions") or []:
            operator = assertion.get("operator", "equals")
            if operator not in operators:
                errors.append(f"{prefix}: operador no soportado {operator!r}")
            if operator in {"equals", "not_equals", "contains", "formula_contains"} and "value" not in assertion:
                errors.append(f"{prefix}: falta el valor esperado para {operator}")
    return errors


def run_functional_contract(path: Path, contract: dict[str, Any]) -> dict[str, Any]:
    """Run declared input/output cases through Excel's calculation engine.

    This is deliberately an optional execution gate: structural checks work
    everywhere openpyxl works, while this check uses installed desktop Excel
    so formulas are recalculated instead of merely inspected as strings.
    """
    cases = list(contract.get("cases") or [])
    if not cases:
        return {"status": "not_declared", "passed": False, "checks": [], "errors": ["Faltan casos funcionales"]}
    try:
        import win32com.client as win32  # type: ignore
        import pythoncom
    except Exception as exc:
        return {
            "status": "unavailable",
            "passed": False,
            "checks": [],
            "errors": [f"Excel funcional no disponible: {exc}"],
        }

    pythoncom.CoInitialize()
    temp_dir = Path(tempfile.mkdtemp(prefix="excel_model_qa_"))
    temp_path = temp_dir / path.name
    shutil.copy2(path, temp_path)
    excel = None
    workbook = None
    target = None
    table = None
    added = None
    checks: list[str] = []
    errors: list[str] = []
    try:
        excel = win32.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        for case in cases:
            target = table = added = None
            if workbook is not None:
                workbook.Close(SaveChanges=False)
            workbook = excel.Workbooks.Open(str(temp_path.resolve()), UpdateLinks=0, ReadOnly=False)
            workbook.Application.CalculateFullRebuild()
            name = str(case.get("name", "functional case"))
            for action in case.get("actions", []):
                table = workbook.Worksheets(action["sheet"]).ListObjects(action["table"])
                if action["operation"] == "append_rows":
                    for _ in range(int(action.get("count", 1))):
                        added = table.ListRows.Add()
                        for column, value in action.get("values", {}).items():
                            added.Range.Cells(1, table.ListColumns(column).Index).Value = value
                elif action["operation"] == "delete_row":
                    table.ListRows(int(action["index"])).Delete()
                else:
                    raise ValueError(f"Operación funcional no soportada: {action['operation']}")
            before: dict[tuple[str, str], Any] = {}
            for assertion in case.get("assertions", []):
                sheet = str(assertion["sheet"])
                address = str(assertion["cell"])
                before[(sheet, address)] = workbook.Worksheets(sheet).Range(address).Value
            for input_spec in case.get("inputs", []):
                workbook.Worksheets(str(input_spec["sheet"])).Range(str(input_spec["cell"])).Value = input_spec.get("value")
            try:
                workbook.Application.CalculateFullRebuild()
            except Exception:
                workbook.Application.CalculateFull()
            for assertion in case.get("assertions", []):
                sheet = str(assertion["sheet"])
                address = str(assertion["cell"])
                target = workbook.Worksheets(sheet).Range(address)
                actual = target.Value
                operator = str(assertion.get("operator", "equals"))
                expected = assertion.get("value")
                passed = True
                if operator == "equals":
                    passed = _values_equal(actual, expected, float(assertion.get("tolerance", 0.0001)))
                elif operator == "not_equals":
                    passed = not _values_equal(actual, expected)
                elif operator == "changed":
                    passed = not _values_equal(actual, before[(sheet, address)])
                elif operator == "unchanged":
                    passed = _values_equal(actual, before[(sheet, address)])
                elif operator == "contains":
                    passed = str(expected) in str(actual)
                elif operator == "formula_contains":
                    passed = str(expected).casefold() in str(target.Formula).casefold()
                elif operator == "not_error":
                    passed = not bool(excel.WorksheetFunction.IsError(target))
                else:
                    errors.append(f"{name}/{sheet}!{address}: operador no soportado {operator}")
                    continue
                if passed:
                    checks.append(f"{name}: {sheet}!{address} {operator}")
                else:
                    errors.append(
                        f"{name}: {sheet}!{address} esperaba {operator} {expected!r}, obtuvo {actual!r}"
                    )
        return {
            "status": "executed",
            "passed": not errors,
            "checks": checks,
            "errors": errors,
            "case_count": len(cases),
        }
    except Exception as exc:
        return {
            "status": "error",
            "passed": False,
            "checks": checks,
            "errors": [f"fallo en QA funcional: {exc}"],
        }
    finally:
        target = table = added = None
        if workbook is not None:
            try:
                workbook.Close(SaveChanges=False)
            except Exception:
                pass
            workbook = None
        target = None
        if excel is not None:
            try:
                excel.Quit()
            except Exception:
                pass
            excel = None
        gc.collect()
        pythoncom.CoUninitialize()
        shutil.rmtree(temp_dir, ignore_errors=True)


def validate_model_file(
    path: Path,
    expected_sheets: Iterable[str] | None = None,
    functional_contract: dict[str, Any] | None = None,
    run_functional: bool = False,
    model_contract: dict[str, Any] | None = None,
    model_contract_path: Path | None = None,
) -> dict[str, Any]:
    if not path.exists():
        return {"passed": False, "checks": [], "errors": ["XLSX no existe"], "details": {}}
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
    except Exception as exc:
        return {"passed": False, "checks": [], "errors": [f"No se puede abrir: {exc}"], "details": {}}
    try:
        actual_sheets = list(workbook.sheetnames)
        report = validate_model_workbook(workbook, expected_sheets)
        if model_contract:
            contract_report = validate_model_contract(workbook, model_contract, model_contract_path, path)
            report["model_contract"] = contract_report
            report["checks"].extend(contract_report.get("checks", []))
            report["errors"].extend(contract_report.get("errors", []))
            report["passed"] = report["passed"] and contract_report.get("passed", False)
        elif model_contract_path is not None:
            report["model_contract"] = {
                "passed": False,
                "checks": [],
                "errors": ["hay guía funcional pero falta el contrato declarativo model_contract"],
            }
            report["errors"].append("hay guía funcional pero falta el contrato declarativo model_contract")
            report["passed"] = False
    finally:
        workbook.close()
    if run_functional:
        functional_contract = functional_contract or {}
        contract_errors = validate_functional_contract_schema(functional_contract, actual_sheets)
        if contract_errors:
            functional = {
                "status": "invalid_contract",
                "passed": False,
                "checks": [],
                "errors": contract_errors,
            }
        else:
            functional = run_functional_contract(path, functional_contract)
        report["functional"] = functional
        report["checks"].extend(functional.get("checks", []))
        report["errors"].extend(functional.get("errors", []))
        report["passed"] = report["passed"] and functional.get("passed", False)
    elif functional_contract:
        contract_errors = validate_functional_contract_schema(functional_contract, actual_sheets)
        report["functional"] = {
            "status": "invalid_contract" if contract_errors else "declared_not_run",
            "passed": not contract_errors,
            "checks": [],
            "errors": contract_errors,
        }
        report["errors"].extend(contract_errors)
        report["passed"] = report["passed"] and not contract_errors
    return report


def load_contract(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))
