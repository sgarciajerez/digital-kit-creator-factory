from __future__ import annotations

"""Native XLSX preflight that checks the workbook directly."""

import argparse
import json
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from artifact_integrity import approval_is_current
from model_contracts import contract_for_product


ROOT = Path(__file__).resolve().parent.parent


def locate_workbook(product: Path, spec: dict) -> Path:
    expected = product / "app" / f"{spec['slug']}-excel-template.xlsx"
    if expected.is_file():
        return expected
    candidates = sorted((product / "app").glob("*.xlsx"))
    if not candidates:
        raise FileNotFoundError(f"No hay XLSX generado en {product / 'app'}")
    return candidates[-1]


def visual_identity_contract(spec: dict) -> dict:
    palette = spec.get("palette") or {}
    required = ["background", "ink", "accent", "secondary", "highlight"]
    errors = []
    missing = [key for key in required if not palette.get(key)]
    if missing:
        errors.append(f"Faltan colores de paleta: {', '.join(missing)}")
    elif len({str(palette[key]).casefold() for key in required}) != len(required):
        errors.append("La paleta debe contener cinco colores distintos")
    identity = spec.get("visual_identity") or {}
    if not identity.get("name") or not identity.get("motif"):
        errors.append("Falta la identidad visual declarada")
    return {"passed": not errors, "palette": {key: palette.get(key) for key in required}, "errors": errors}


def table_layout_contract(workbook, spec: dict | None = None) -> dict:
    errors = []
    tables = []
    for ws in workbook.worksheets:
        if ws.freeze_panes != (spec or {}).get("freeze_panes", {}).get(ws.title):
            errors.append(f"{ws.title}: freeze_panes_not_declared")
        for table in ws.tables.values():
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            headers = [ws.cell(min_row, col).value for col in range(min_col, max_col + 1)]
            if any(value in (None, "") for value in headers):
                errors.append(f"{ws.title}!{table.ref}: empty_table_header")
            if max_row <= min_row:
                errors.append(f"{ws.title}!{table.ref}: empty_data_range")
            if len(set(str(value).strip().casefold() for value in headers)) != len(headers):
                errors.append(f"{ws.title}!{table.ref}: duplicate_table_headers")
            tables.append({"sheet": ws.title, "name": table.name, "ref": table.ref, "headers": headers})
    return {"passed": not errors, "tables": tables, "errors": errors}


def native_geometry_contract(final_wb, master_wb, expected_tabs: list[str], spec: dict | None = None) -> dict:
    errors = []
    details = {"mapped_sheets": {}, "master_sheet_map": {}}
    adaptations = (spec or {}).get("native_surface_adaptations") or {}
    width_adaptation_tabs = set(adaptations.get("allow_column_width_changes", []))
    merged_slot_resizes = adaptations.get("merged_slot_resizes", {})
    if final_wb.sheetnames != expected_tabs:
        errors.append("final_workbook_tabs_do_not_match_spec")
    if not master_wb.sheetnames:
        errors.append("master_workbook_has_no_sheets")
    # The product may rename or select a subset of master screens, but every
    # mapped screen must retain the master's native frame. This is the check
    # that was missing when the product builder recreated sheets from scratch.
    mapping = getattr(final_wb, "_native_sheet_map", None) or {}
    if not mapping:
        errors.append("native_sheet_map_missing: no se puede validar geometría contra el master")
        mapping = {}
    elif list(mapping) != expected_tabs:
        errors.append("native_sheet_map_incomplete_or_out_of_order")
    elif len(set(mapping.values())) != len(mapping):
        errors.append("native_sheet_map_has_duplicate_master_sheets")
    details["master_sheet_map"] = mapping
    for name in expected_tabs:
        final_ws = final_wb[name]
        if final_ws.freeze_panes != (spec or {}).get("freeze_panes", {}).get(name):
            errors.append(f"{name}: freeze_panes_not_declared")
        if final_ws.sheet_view.showGridLines is not False:
            errors.append(f"{name}: gridlines_must_be_hidden")
        if final_ws.max_row < 4 or final_ws.max_column < 4:
            errors.append(f"{name}: native_grid_too_small")
        for col in "ABCDEFGHIJKLMNO":
            final_width = final_ws.column_dimensions[col].width
            if final_width is not None and final_width <= 0:
                errors.append(f"{name}!{col}: invalid_column_width")
        master_name = mapping.get(name)
        if not master_name:
            errors.append(f"{name}: missing_master_mapping")
            continue
        if master_name not in master_wb.sheetnames:
            errors.append(f"{name}: master_sheet_not_found:{master_name}")
            continue
        if master_name and master_name in master_wb.sheetnames:
            master_ws = master_wb[master_name]
            width_errors = []
            for col in "ABCDEFGHIJKLMNOPQ":
                final_width = final_ws.column_dimensions[col].width
                master_width = master_ws.column_dimensions[col].width
                if final_width is not None and master_width is not None and abs(float(final_width) - float(master_width)) > 0.05:
                    width_errors.append(col)
            if width_errors and name not in width_adaptation_tabs:
                errors.append(f"{name}: column_widths_do_not_match_master:{','.join(width_errors)}")
            final_merged = {str(item) for item in final_ws.merged_cells.ranges}
            master_merged = {str(item) for item in master_ws.merged_cells.ranges}
            expected_merged = set(master_merged)
            for old_range, new_range in (merged_slot_resizes.get(name) or {}).items():
                expected_merged.discard(old_range)
                expected_merged.add(new_range)
            if final_merged != expected_merged:
                errors.append(f"{name}: merged_cells_do_not_match_master")
            if len(final_ws._charts) != len(master_ws._charts):
                errors.append(f"{name}: chart_count_does_not_match_master")
            # The master carries one legacy identity image. New products
            # intentionally remove that logo; geometry must not require the
            # obsolete external identity asset.
            images_match = len(final_ws._images) == len(master_ws._images)
            intentional_master_identity_removal = len(final_ws._images) == 0 and len(master_ws._images) == 1
            if not images_match and not intentional_master_identity_removal:
                errors.append(f"{name}: image_count_does_not_match_master")
            master_table_refs = [table.ref for table in master_ws.tables.values()]
            final_table_refs = [table.ref for table in final_ws.tables.values()]
            table_ranges_match = final_table_refs == master_table_refs
            if (spec or {}).get("native_geometry_policy") == "exact_master_grid" and not table_ranges_match:
                errors.append(f"{name}: table_ranges_do_not_match_master")
            details["mapped_sheets"][name] = {
                "master_sheet": master_name,
                "column_widths_match": not width_errors or name in width_adaptation_tabs,
                "merged_cells_match": final_merged == expected_merged,
                "charts": len(final_ws._charts),
                "master_charts": len(master_ws._charts),
                "images": len(final_ws._images),
                "master_images": len(master_ws._images),
                "legacy_master_identity_removed": intentional_master_identity_removal,
                "table_ranges_match": table_ranges_match,
            }
    return {"passed": not errors, "errors": errors, "details": details}


def evaluate(product_id: int) -> tuple[bool, dict]:
    product = ROOT / "outputs" / "listings" / f"product-{product_id}"
    spec = json.loads((product / "design_spec.json").read_text(encoding="utf-8"))
    book_path = locate_workbook(product, spec)
    master_path = ROOT / str(spec.get("master_workbook", "templates/excel-first-base/master_workbook.xlsx"))
    workbook = load_workbook(book_path, read_only=False, data_only=False)
    master = load_workbook(master_path, read_only=False, data_only=False) if master_path.is_file() else None
    try:
        app_report_path = product / "app_generation_report.json"
        app_report = json.loads(app_report_path.read_text(encoding="utf-8")) if app_report_path.is_file() else {}
        app_guide_path = product / str(spec.get("app_user_guide_file", "app/README.md"))
        expected_language = str(spec.get("design_language") or "en").casefold()
        checks = {
            "workbook_exists": book_path.is_file(),
            "workbook_tabs_match": workbook.sheetnames == list(spec["tabs"]),
            "master_workbook_exists": master is not None,
            "native_contract_mode": spec.get("native_reference_mode") == "master_workbook_contract",
            "model_and_sources_declared": bool(spec.get("model_base")) and bool(spec.get("source_lineage")),
            "app_user_guide_exists": app_guide_path.is_file(),
            "app_user_guide_language_matches": app_guide_path.is_file() and f"**Language:** `{expected_language}`" in app_guide_path.read_text(encoding="utf-8"),
            "functional_validation_passed": approval_is_current(app_report, book_path, master_path, contract_for_product(product_id), spec),
            "table_layout_contract": table_layout_contract(workbook, spec),
            "visual_identity_contract": visual_identity_contract(spec),
        }
        workbook._native_sheet_map = spec.get("native_sheet_map") or {}
        checks["native_geometry_contract"] = native_geometry_contract(workbook, master, list(spec["tabs"]), spec) if master else {"passed": False, "errors": ["missing_master_workbook"]}
        passed = all((
            checks["workbook_exists"],
            checks["workbook_tabs_match"],
            checks["master_workbook_exists"],
            checks["native_contract_mode"],
            checks["model_and_sources_declared"],
            checks["functional_validation_passed"],
            checks["app_user_guide_exists"],
            checks["app_user_guide_language_matches"],
            checks["table_layout_contract"]["passed"],
            checks["visual_identity_contract"]["passed"],
            checks["native_geometry_contract"]["passed"],
        ))
    finally:
        workbook.close()
        if master is not None:
            master.close()
    report = {
        "schema_version": "2.0.0",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "product_id": product_id,
        "workbook": str(book_path),
        "status": "passed" if passed else "failed",
        "checks": checks,
        "next_stage": "functional_qa_and_contract_audit" if passed else "repair_native_xlsx_contract",
    }
    (product / "native_xlsx_preflight_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return passed, report


def main() -> int:
    parser = argparse.ArgumentParser(description="Preflight estructural nativo del XLSX; no usa Excel ni capturas")
    parser.add_argument("--product", type=int, required=True)
    args = parser.parse_args()
    passed, report = evaluate(args.product)
    print(f"product-{args.product}: NATIVE_XLSX_PREFLIGHT={report['status'].upper()}")
    print(f"next_stage={report['next_stage']}")
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
